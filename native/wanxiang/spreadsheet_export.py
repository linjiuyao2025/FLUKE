"""Create safe, atomic XLSX files for local data exports."""

from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal
import math
import os
from pathlib import Path
import tempfile
from typing import Any, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.cell.cell import TYPE_STRING


_MAX_ROWS = 1_048_576
_MAX_COLUMNS = 16_384
_MAX_CELL_TEXT = 32_767
_INVALID_SHEET_CHARS = set("[]:*?/\\")


class SpreadsheetExportError(ValueError):
    """A readable error raised when an XLSX export cannot be completed."""


def _validate_sheet_name(sheet_name: str) -> str:
    if not isinstance(sheet_name, str) or not sheet_name.strip():
        raise SpreadsheetExportError("工作表名称不能为空。")
    if len(sheet_name) > 31:
        raise SpreadsheetExportError("工作表名称不能超过 31 个字符。")
    if any(char in _INVALID_SHEET_CHARS for char in sheet_name):
        raise SpreadsheetExportError("工作表名称不能包含 [ ] : * ? / 或反斜杠。")
    if sheet_name.startswith("'") or sheet_name.endswith("'"):
        raise SpreadsheetExportError("工作表名称不能以单引号开头或结尾。")
    return sheet_name


def _validate_text(value: str, location: str) -> None:
    if len(value) > _MAX_CELL_TEXT:
        raise SpreadsheetExportError(
            f"{location}超过 Excel 单元格可保存的 {_MAX_CELL_TEXT} 个字符上限。"
        )


def _validate_value(value: Any, location: str) -> None:
    if isinstance(value, str):
        _validate_text(value, location)
        return
    if value is None or isinstance(value, bool):
        return
    if isinstance(value, Decimal):
        if not value.is_finite():
            raise SpreadsheetExportError(f"{location}不是有效的有限数字。")
        return
    if isinstance(value, int):
        return
    if isinstance(value, float):
        if not math.isfinite(value):
            raise SpreadsheetExportError(f"{location}不是有效的有限数字。")
        return
    if isinstance(value, datetime):
        if value.tzinfo is not None and value.utcoffset() is not None:
            raise SpreadsheetExportError(f"{location}含时区；Excel 日期不支持时区。")
        return
    if isinstance(value, (date, time)):
        if (
            isinstance(value, time)
            and value.tzinfo is not None
            and value.utcoffset() is not None
        ):
            raise SpreadsheetExportError(f"{location}含时区；Excel 时间不支持时区。")
        return
    raise SpreadsheetExportError(
        f"{location}的数据类型不受支持：请使用文本、数字、布尔值、日期/时间或空值。"
    )


def _write_value(cell: Any, value: Any) -> None:
    cell.value = value
    # openpyxl normally treats strings beginning with '=' as formulas. The
    # exported workbook is a data file, so every caller-provided string stays
    # explicitly marked as text, including header strings.
    if isinstance(value, str):
        cell.data_type = TYPE_STRING


def export_xlsx(
    target_path: str | os.PathLike[str],
    sheet_name: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
) -> Path:
    """Write an XLSX workbook and return the resolved target path.

    ``headers`` becomes the first row. Each data row must have the same number
    of cells. Text is always serialized as text, never as a spreadsheet formula.
    The workbook is saved to a temporary file beside the destination and then
    atomically replaces the destination when the save succeeds.
    """

    try:
        target = Path(target_path).expanduser().resolve(strict=False)
    except (TypeError, OSError, RuntimeError) as exc:
        raise SpreadsheetExportError(f"导出路径无效：{exc}") from exc

    if target.suffix.lower() != ".xlsx":
        raise SpreadsheetExportError("文件名必须以 .xlsx 结尾。")
    if not target.parent.is_dir():
        raise SpreadsheetExportError("导出文件夹不存在或不可访问。")
    if target.exists() and target.is_dir():
        raise SpreadsheetExportError("目标路径是文件夹，请选择一个 .xlsx 文件名。")

    sheet_name = _validate_sheet_name(sheet_name)
    if isinstance(headers, (str, bytes)) or not isinstance(headers, Sequence):
        raise SpreadsheetExportError("标题行必须是字符串列表或元组。")
    if not headers:
        raise SpreadsheetExportError("标题行至少需要一列。")
    if len(headers) > _MAX_COLUMNS:
        raise SpreadsheetExportError(f"列数不能超过 {_MAX_COLUMNS} 列。")
    for index, header in enumerate(headers, start=1):
        if not isinstance(header, str):
            raise SpreadsheetExportError(f"标题行第 {index} 列必须是文本。")
        _validate_text(header, f"标题行第 {index} 列")

    temporary_path: Path | None = None
    try:
        fd, temp_name = tempfile.mkstemp(
            dir=target.parent,
            prefix=f".{target.stem}.",
            suffix=".tmp.xlsx",
        )
        temporary_path = Path(temp_name)
        os.close(fd)

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name
        for column, header in enumerate(headers, start=1):
            _write_value(worksheet.cell(row=1, column=column), header)

        try:
            row_iterator = iter(rows)
        except TypeError as exc:
            raise SpreadsheetExportError("数据行必须是可迭代的行列表。") from exc

        for excel_row, row in enumerate(row_iterator, start=2):
            if excel_row > _MAX_ROWS:
                raise SpreadsheetExportError(f"数据行数不能超过 {_MAX_ROWS - 1} 行。")
            if isinstance(row, (str, bytes)) or not isinstance(row, Sequence):
                raise SpreadsheetExportError(f"第 {excel_row} 行必须是单元格列表或元组。")
            if len(row) != len(headers):
                raise SpreadsheetExportError(
                    f"第 {excel_row} 行有 {len(row)} 列，标题行有 {len(headers)} 列。"
                )
            for column, value in enumerate(row, start=1):
                location = f"第 {excel_row} 行第 {column} 列"
                _validate_value(value, location)
                _write_value(worksheet.cell(row=excel_row, column=column), value)

        workbook.save(temporary_path)
        os.replace(temporary_path, target)
        temporary_path = None
        return target
    except SpreadsheetExportError:
        raise
    except Exception as exc:
        raise SpreadsheetExportError(f"XLSX 导出失败，目标文件未被替换：{exc}") from exc
    finally:
        if temporary_path is not None:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass


__all__ = ["SpreadsheetExportError", "export_xlsx"]
