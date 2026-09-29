from __future__ import annotations

import csv
from collections import Counter
from contextlib import ExitStack
from dataclasses import replace
from html import escape
from html.parser import HTMLParser
import io
import json
import math
import os
from pathlib import Path
from pathlib import PurePosixPath
import posixpath
import re
import shutil
import stat
import struct
import subprocess
import sys
import tempfile
import time
import uuid
import zlib
from typing import Any, Callable
from urllib.parse import unquote, urlsplit
import zipfile
from functools import lru_cache
from datetime import datetime, timezone

from PySide6.QtCore import (
    QByteArray,
    QBuffer,
    QCoreApplication,
    QIODevice,
    QObject,
    Property,
    QSettings,
    QMarginsF,
    QRunnable,
    QThreadPool,
    QTimer,
    QRectF,
    QSize,
    QSizeF,
    QUrl,
    Signal,
    Slot,
    Qt,
)
from PySide6.QtGui import (
    QImage,
    QImageReader,
    QImageWriter,
    QPageLayout,
    QPageSize,
    QPainter,
    QPdfWriter,
)
from PySide6.QtQml import QJSValue
from PySide6.QtPdf import QPdfDocument

from .media_conversion import (
    AUDIO_INPUT_EXTENSIONS,
    MEDIA_INPUT_EXTENSIONS,
    MAX_MEDIA_BYTES,
    VIDEO_TARGETS,
    VIDEO_INPUT_EXTENSIONS,
    available_targets as media_targets_for,
    available_video_codecs,
    convert_media,
    media_target_label,
)
from .converter_engines import engine_executable
from .pdf_ocr import (
    MAX_OCR_DOCUMENT_SECONDS,
    MAX_OCR_PAGE_PIXELS,
    MAX_OCR_PAGE_SECONDS,
    MAX_OCR_PAGES,
    MAX_OCR_TOTAL_PIXELS,
    OCR_DPI,
    available_ocr_languages,
    create_searchable_pdf,
    recognize_tesseract_tsv,
    searchable_pdf_ocr_available,
    tesseract_engine,
)


TABULAR_FORMATS = {
    "csv": {"label": "CSV", "value": "csv"},
    "tsv": {"label": "TSV", "value": "tsv"},
    "json": {"label": "JSON", "value": "json"},
    "xlsx": {"label": "Excel 工作簿（XLSX）", "value": "xlsx"},
}
IMAGE_FORMATS = {
    "bmp": {"label": "BMP", "value": "bmp", "qt": "bmp"},
    "cur": {"label": "CUR", "value": "cur", "qt": "cur"},
    "gif": {"label": "GIF 首帧", "value": "gif", "qt": "gif"},
    "jpg": {"label": "JPEG", "value": "jpg", "qt": "jpg"},
    "jpeg": {"label": "JPEG", "value": "jpeg", "qt": "jpeg"},
    "png": {"label": "PNG", "value": "png", "qt": "png"},
    "icns": {"label": "ICNS", "value": "icns", "qt": "icns"},
    "ico": {"label": "ICO（多尺寸）", "value": "ico", "qt": "ico"},
    "svg": {"label": "SVG", "value": "svg", "qt": "svg"},
    "webp": {"label": "WebP", "value": "webp", "qt": "webp"},
    "tif": {"label": "TIFF", "value": "tif", "qt": "tif"},
    "tiff": {"label": "TIFF", "value": "tiff", "qt": "tiff"},
    "tga": {"label": "TGA", "value": "tga", "qt": "tga"},
    "heic": {"label": "HEIC", "value": "heic", "qt": "heic"},
    "heif": {"label": "HEIF", "value": "heif", "qt": "heif"},
    "avif": {"label": "AVIF", "value": "avif", "qt": "avif"},
    "jp2": {"label": "JPEG 2000", "value": "jp2", "qt": "jp2"},
    "j2k": {"label": "JPEG 2000", "value": "j2k", "qt": "jp2"},
    "jpf": {"label": "JPEG 2000", "value": "jpf", "qt": "jp2"},
    "mng": {"label": "MNG", "value": "mng", "qt": "mng"},
    "wbmp": {"label": "WBMP", "value": "wbmp", "qt": "wbmp"},
    "3fr": {"label": "相机 RAW（3FR）", "value": "3fr", "qt": "3fr"},
    "arw": {"label": "相机 RAW（ARW）", "value": "arw", "qt": "arw"},
    "cr2": {"label": "相机 RAW（CR2）", "value": "cr2", "qt": "cr2"},
    "cr3": {"label": "相机 RAW（CR3）", "value": "cr3", "qt": "cr3"},
    "dng": {"label": "相机 RAW（DNG）", "value": "dng", "qt": "dng"},
    "nef": {"label": "相机 RAW（NEF）", "value": "nef", "qt": "nef"},
    "orf": {"label": "相机 RAW（ORF）", "value": "orf", "qt": "orf"},
    "pef": {"label": "相机 RAW（PEF）", "value": "pef", "qt": "pef"},
    "raf": {"label": "相机 RAW（RAF）", "value": "raf", "qt": "raf"},
    "rw2": {"label": "相机 RAW（RW2）", "value": "rw2", "qt": "rw2"},
    "srw": {"label": "相机 RAW（SRW）", "value": "srw", "qt": "srw"},
    "srf": {"label": "相机 RAW（SRF）", "value": "srf", "qt": "srf"},
    "sr2": {"label": "相机 RAW（SR2）", "value": "sr2", "qt": "sr2"},
    "nrw": {"label": "相机 RAW（NRW）", "value": "nrw", "qt": "nrw"},
    "kdc": {"label": "相机 RAW（KDC）", "value": "kdc", "qt": "kdc"},
    "mos": {"label": "相机 RAW（MOS）", "value": "mos", "qt": "mos"},
    "mef": {"label": "相机 RAW（MEF）", "value": "mef", "qt": "mef"},
    "erf": {"label": "相机 RAW（ERF）", "value": "erf", "qt": "erf"},
    "raw": {"label": "相机 RAW", "value": "raw", "qt": "raw"},
}
READ_ONLY_IMAGE_EXTENSIONS = {
    "heic", "heif", "avif", "jp2", "j2k", "jpf", "mng", "wbmp",
    "3fr", "arw", "cr2", "cr3", "dng", "nef", "orf", "pef", "raf", "rw2",
    "srw", "srf", "sr2", "nrw", "kdc", "mos", "mef", "erf", "raw",
}
RAW_CAMERA_IMAGE_EXTENSIONS = frozenset({
    "3fr", "arw", "cr2", "cr3", "dng", "nef", "orf", "pef", "raf", "rw2",
    "srw", "srf", "sr2", "nrw", "kdc", "mos", "mef", "erf", "raw",
})
RAW_IMAGE_OUTPUTS = ("jpg", "png", "tif", "webp", "bmp")
TEXT_EXTENSIONS = {
    "txt", "md", "markdown", "html", "htm", "log", "xml", "yaml", "yml",
}
DOCUMENT_FORMATS = {
    "txt": {"label": "纯文本（TXT）", "value": "txt"},
    "md": {"label": "Markdown（MD）", "value": "md"},
    "html": {"label": "HTML 网页", "value": "html"},
    "docx": {"label": "Word 文档（DOCX）", "value": "docx"},
    "pdf": {"label": "PDF 文档", "value": "pdf"},
    "epub": {"label": "电子书（EPUB）", "value": "epub"},
}
WORD_OFFICE_EXTENSIONS = {"doc", "docm", "dot", "dotm", "dotx", "odt", "rtf"}
EXCEL_OFFICE_EXTENSIONS = {"ods", "xls", "xlsb", "xlsm", "xlt", "xltm", "xltx"}
OFFICE_TABULAR_OUTPUTS = {
    "xls": {"label": "Excel 97–2003 工作簿（XLS）", "value": "xls"},
    "ods": {"label": "OpenDocument 电子表格（ODS）", "value": "ods"},
}
POWERPOINT_OFFICE_EXTENSIONS = {"potx", "pps", "ppsx", "ppt", "pptm"}
WPS_WORD_EXTENSIONS = {"wps", "wpt"}
WPS_EXCEL_EXTENSIONS = {"et", "ett"}
WPS_PRESENTATION_EXTENSIONS = {"dps", "dpt"}
MAX_IMAGE_BYTES = 512 * 1024 * 1024
MAX_IMAGE_MERGE_FILES = 500
MAX_ANIMATED_IMAGE_FRAMES = 500
MAX_ANIMATED_IMAGE_TOTAL_PIXELS = 64_000_000
MAX_ANIMATED_IMAGE_OUTPUT_BYTES = 512 * 1024 * 1024
MAX_TABULAR_BYTES = 100 * 1024 * 1024
MAX_IMAGE_PIXELS = 40_000_000
MAX_DOCUMENT_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_EXPANDED_BYTES = 256 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 20_000
MAX_PDF_PAGES = 2_000
MAX_PDF_TABLE_PAGES = 500
MAX_PDF_TABLES = 5_000
MAX_PDF_TABLE_CELLS = 250_000
MAX_PDF_TABLE_TEXT_BYTES = 64 * 1024 * 1024
MAX_PDF_TABLE_XLSX_BYTES = 128 * 1024 * 1024
MAX_PDF_TABLE_PAGE_CHARS = 150_000
MAX_PDF_TABLE_TOTAL_CHARS = 2_000_000
MAX_PDF_TABLE_PAGE_EDGES = 100_000
MAX_PDF_TABLE_OCR_PAGES = MAX_OCR_PAGES
MAX_PDF_TABLE_OCR_CELLS = 25_000
MAX_PDF_TABLE_OCR_TEXT_BYTES = 32 * 1024 * 1024
MAX_PDF_LAYOUT_PAGES = 500
MAX_PDF_LAYOUT_PAGE_CHARS = 150_000
MAX_PDF_LAYOUT_TOTAL_CHARS = 2_000_000
MAX_PDF_LAYOUT_PIXELS = MAX_OCR_TOTAL_PIXELS
MAX_PDF_LAYOUT_DOCX_BYTES = 128 * 1024 * 1024
PDF_PAGE_IMAGE_DPI = 144
MAX_PDF_IMAGE_PAGE_PIXELS = MAX_IMAGE_PIXELS
MAX_PDF_IMAGE_TOTAL_PIXELS = 500_000_000
MAX_MERGED_PDF_BYTES = 512 * 1024 * 1024
MAX_EXTRACTED_TEXT_BYTES = 32 * 1024 * 1024
MAX_EBOOK_SOURCE_BYTES = 64 * 1024 * 1024
MAX_EBOOK_OUTPUT_BYTES = 256 * 1024 * 1024
MAX_OFFICE_OUTPUT_BYTES = 512 * 1024 * 1024
MAX_OFD_OUTPUT_BYTES = 512 * 1024 * 1024
MAX_OFD_XML_BYTES = 16 * 1024 * 1024
MAX_OFD_PAGES = 500
OFD_TIMEOUT_SECONDS = 180
MAX_TEXT_PREVIEW_BYTES = 1024 * 1024
MAX_PDF_PASSWORD_LENGTH = 512
MIN_PDF_PASSWORD_LENGTH = 8
OFFICE_TIMEOUT_SECONDS = 180
EBOOK_TIMEOUT_SECONDS = 600

_OFFICE_AUTOMATION_SCRIPT = r"""
param([string]$Operation, [string]$Source, [string]$Destination)
$ErrorActionPreference = 'Stop'
$application = $null
$document = $null
try {
    switch ($Operation) {
        'word-pdf' {
            $application = New-Object -ComObject Word.Application
            $application.Visible = $false
            $application.DisplayAlerts = 0
            $application.AutomationSecurity = 3
            $application.Options.UpdateLinksAtOpen = $false
            $document = $application.Documents.Open($Source, $false, $true, $false)
            $document.ExportAsFixedFormat($Destination, 17)
        }
        'word-docx' {
            $application = New-Object -ComObject Word.Application
            $application.Visible = $false
            $application.DisplayAlerts = 0
            $application.AutomationSecurity = 3
            $application.Options.UpdateLinksAtOpen = $false
            $document = $application.Documents.Open($Source, $false, $true, $false)
            $document.SaveAs2($Destination, 16)
        }
        'wps-word-pdf' {
            $application = New-Object -ComObject KWPS.Application
            $application.Visible = $false
            $application.DisplayAlerts = 0
            $application.AutomationSecurity = 3
            $application.Options.UpdateLinksAtOpen = $false
            $document = $application.Documents.Open($Source, $false, $true, $false)
            $document.ExportAsFixedFormat($Destination, 17)
        }
        'wps-word-docx' {
            $application = New-Object -ComObject KWPS.Application
            $application.Visible = $false
            $application.DisplayAlerts = 0
            $application.AutomationSecurity = 3
            $application.Options.UpdateLinksAtOpen = $false
            $document = $application.Documents.Open($Source, $false, $true, $false)
            $document.SaveAs2($Destination, 16)
        }
        'excel-pdf' {
            $application = New-Object -ComObject Excel.Application
            $application.Visible = $false
            $application.DisplayAlerts = $false
            $application.AutomationSecurity = 3
            $application.EnableEvents = $false
            $document = $application.Workbooks.Open($Source, 0, $true)
            $document.ExportAsFixedFormat(0, $Destination)
        }
        'excel-xlsx' {
            $application = New-Object -ComObject Excel.Application
            $application.Visible = $false
            $application.DisplayAlerts = $false
            $application.AutomationSecurity = 3
            $application.EnableEvents = $false
            $document = $application.Workbooks.Open($Source, 0, $true)
            $document.SaveAs($Destination, 51)
        }
        'excel-xls' {
            $application = New-Object -ComObject Excel.Application
            $application.Visible = $false
            $application.DisplayAlerts = $false
            $application.AutomationSecurity = 3
            $application.EnableEvents = $false
            $document = $application.Workbooks.Open($Source, 0, $true)
            $document.SaveAs($Destination, 56)
        }
        'excel-csv' {
            $application = New-Object -ComObject Excel.Application
            $application.Visible = $false
            $application.DisplayAlerts = $false
            $application.AutomationSecurity = 3
            $application.EnableEvents = $false
            $document = $application.Workbooks.Open($Source, 0, $true)
            $document.SaveAs($Destination, 62)
        }
        'excel-ods' {
            $application = New-Object -ComObject Excel.Application
            $application.Visible = $false
            $application.DisplayAlerts = $false
            $application.AutomationSecurity = 3
            $application.EnableEvents = $false
            $document = $application.Workbooks.Open($Source, 0, $true)
            $document.SaveAs($Destination, 60)
        }
        'wps-excel-pdf' {
            $application = New-Object -ComObject KET.Application
            $application.Visible = $false
            $application.DisplayAlerts = $false
            $application.AutomationSecurity = 3
            $application.EnableEvents = $false
            $document = $application.Workbooks.Open($Source, 0, $true)
            $document.ExportAsFixedFormat(0, $Destination)
        }
        'wps-excel-xlsx' {
            $application = New-Object -ComObject KET.Application
            $application.Visible = $false
            $application.DisplayAlerts = $false
            $application.AutomationSecurity = 3
            $application.EnableEvents = $false
            $document = $application.Workbooks.Open($Source, 0, $true)
            $document.SaveAs($Destination, 51)
        }
        'wps-excel-csv' {
            $application = New-Object -ComObject KET.Application
            $application.Visible = $false
            $application.DisplayAlerts = $false
            $application.AutomationSecurity = 3
            $application.EnableEvents = $false
            $document = $application.Workbooks.Open($Source, 0, $true)
            $document.SaveAs($Destination, 62)
        }
        'powerpoint-pdf' {
            $application = New-Object -ComObject PowerPoint.Application
            $application.DisplayAlerts = 1
            $application.AutomationSecurity = 3
            $document = $application.Presentations.Open($Source, $true, $false, $false)
            $document.SaveAs($Destination, 32)
        }
        'wps-powerpoint-pdf' {
            $application = New-Object -ComObject KWPP.Application
            $application.DisplayAlerts = 0
            $application.AutomationSecurity = 3
            $document = $application.Presentations.Open($Source, $true, $false, $false)
            $document.SaveAs($Destination, 32)
        }
        'wps-powerpoint-pptx' {
            $application = New-Object -ComObject KWPP.Application
            $application.DisplayAlerts = 0
            $application.AutomationSecurity = 3
            $document = $application.Presentations.Open($Source, $true, $false, $false)
            $document.SaveAs($Destination, 24)
        }
        default { throw 'Unsupported Office conversion operation.' }
    }
    if ($document) {
        if ($Operation -like 'word-*' -or $Operation -like 'excel-*' -or $Operation -like 'wps-word-*' -or $Operation -like 'wps-excel-*') { $document.Close(0) }
        else { $document.Close() }
    }
} catch {
    [Console]::Error.WriteLine('Office automation failed: ' + $_.Exception.GetType().FullName)
    exit 2
} finally {
    if ($document) {
        try {
            if ($Operation -like 'word-*' -or $Operation -like 'excel-*' -or $Operation -like 'wps-word-*' -or $Operation -like 'wps-excel-*') { $document.Close(0) }
            else { $document.Close() }
        } catch {}
        [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($document)
    }
    if ($application) {
        try { $application.Quit() } catch {}
        [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($application)
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
"""


class ConversionError(ValueError):
    """An input file cannot be converted by the built-in converter."""


def _qt_formats(values: list[bytes]) -> set[str]:
    return {bytes(value).decode("ascii", errors="ignore").lower() for value in values}


@lru_cache(maxsize=1)
def _rawpy_module() -> Any | None:
    try:
        import rawpy
    except (ImportError, OSError):
        return None
    return rawpy


@lru_cache(maxsize=1)
def _pillow_image_modules() -> tuple[Any, Any] | None:
    try:
        from PIL import Image, ImageOps
    except (ImportError, OSError):
        return None
    Image.init()
    return Image, ImageOps


def _pillow_can_save(image_format: str) -> bool:
    modules = _pillow_image_modules()
    return bool(modules and image_format.upper() in modules[0].SAVE)


def _pillow_frame_count(source: Path) -> int:
    modules = _pillow_image_modules()
    if not modules:
        return 0
    try:
        with modules[0].open(source) as image:
            return max(1, int(getattr(image, "n_frames", 1)))
    except Exception:
        return 0


def _write_pillow_animation(
    source: Path,
    destination: Path,
    image_format: str,
    progress_callback: Callable[[int], None] | None = None,
) -> None:
    modules = _pillow_image_modules()
    if not modules or not _pillow_can_save(image_format):
        raise ConversionError("当前设备没有可用的动图读写组件。")
    image_module, image_ops = modules
    frames: list[Any] = []
    durations: list[int] = []
    total_pixels = 0
    try:
        with image_module.open(source) as image:
            frame_count = max(1, int(getattr(image, "n_frames", 1)))
            if frame_count > MAX_ANIMATED_IMAGE_FRAMES:
                raise ConversionError("动图帧数超过本机安全处理上限。")
            original_info = dict(image.info)
            for index in range(frame_count):
                image.seek(index)
                frame = image_ops.exif_transpose(image.copy()).convert("RGBA")
                pixels = frame.width * frame.height
                if frame.width <= 0 or frame.height <= 0 or pixels > MAX_IMAGE_PIXELS:
                    raise ConversionError("动图单帧尺寸超过本机安全处理上限。")
                total_pixels += pixels
                if total_pixels > MAX_ANIMATED_IMAGE_TOTAL_PIXELS:
                    raise ConversionError("动图累计像素超过本机安全处理上限。")
                frames.append(frame)
                duration = int(image.info.get("duration", 100) or 100)
                durations.append(max(10, min(duration, 60_000)))
                if progress_callback:
                    progress_callback(int((index + 1) * 80 / frame_count))

        save_options: dict[str, Any] = {
            "format": image_format.upper(),
            "save_all": True,
            "append_images": frames[1:],
            "duration": durations,
        }
        if "loop" in original_info:
            save_options["loop"] = int(original_info["loop"])
        if image_format.upper() == "GIF":
            save_options["disposal"] = 2
        else:
            save_options.setdefault("loop", 1)
            save_options.update({"quality": 90, "method": 4})
        frames[0].save(destination, **save_options)
        if destination.stat().st_size > MAX_ANIMATED_IMAGE_OUTPUT_BYTES:
            raise ConversionError("动图输出超过本机安全体积上限。")
        if progress_callback:
            progress_callback(100)
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError("动图文件无法读取或输出。") from exc
    finally:
        for frame in frames:
            try:
                frame.close()
            except Exception:
                pass


def supported_image_extensions() -> set[str]:
    """Image sources available from Qt's installed image reader plugins."""
    readers = _qt_formats(QImageReader.supportedImageFormats())
    return {
        extension
        for extension, info in IMAGE_FORMATS.items()
        if info["qt"] in readers
    }


def writable_image_extensions() -> set[str]:
    writers = _qt_formats(QImageWriter.supportedImageFormats())
    return {
        extension for extension, info in IMAGE_FORMATS.items()
        if info["qt"] in writers and extension not in READ_ONLY_IMAGE_EXTENSIONS
    }


def input_extensions() -> set[str]:
    extensions = set(TABULAR_FORMATS) | set(TEXT_EXTENSIONS)
    extensions |= supported_image_extensions()
    if _rawpy_module() is not None:
        extensions |= RAW_CAMERA_IMAGE_EXTENSIONS
    if supported_image_extensions():
        extensions.add("zip")
    if media_targets_for("wav"):
        extensions |= MEDIA_INPUT_EXTENSIONS
    if _has_docx_support():
        extensions.add("docx")
    if _has_pdf_text_support() or "pdf" in _qt_formats(QImageReader.supportedImageFormats()):
        extensions.add("pdf")
    if _has_epub_support():
        extensions.add("epub")
    if ofd_engine_available():
        extensions.add("ofd")
    if _ebook_convert_path():
        extensions.add("mobi")
    if "word" in _office_engines():
        extensions |= WORD_OFFICE_EXTENSIONS
    if "excel" in _office_engines():
        extensions |= EXCEL_OFFICE_EXTENSIONS
    if "powerpoint" in _office_engines():
        extensions |= POWERPOINT_OFFICE_EXTENSIONS | {"pptx"}
    wps_engines = _wps_engines()
    if "word" in wps_engines:
        extensions |= WPS_WORD_EXTENSIONS | {"doc", "odt", "rtf"}
    if "excel" in wps_engines:
        extensions |= WPS_EXCEL_EXTENSIONS
    if "powerpoint" in wps_engines:
        extensions |= WPS_PRESENTATION_EXTENSIONS | {"pptx"}
    return extensions


def _looks_like_pdf(source: str | Path) -> bool:
    path = Path(source)
    try:
        size = path.stat().st_size
        if not path.is_file() or size < 8 or size > MAX_DOCUMENT_BYTES:
            return False
        with path.open("rb") as stream:
            header = stream.read(1024)
            stream.seek(max(0, size - 1024))
            trailer = stream.read(1024)
        return b"%PDF-" in header and b"%%EOF" in trailer
    except OSError:
        return False


def _ofd_engine_root() -> Path:
    if getattr(sys, "frozen", False):
        base = Path(getattr(sys, "_MEIPASS", Path(sys.executable).resolve().parent))
    else:
        base = Path(__file__).resolve().parent.parent
    return base / "engine" / "ofd" / "build"


def _ofd_java_path() -> str | None:
    bundled = _ofd_engine_root() / "java-runtime" / "bin" / "java.exe"
    if bundled.is_file():
        return str(bundled)
    return shutil.which("java")


def ofd_engine_available() -> bool:
    root = _ofd_engine_root()
    return bool(
        _ofd_java_path()
        and (root / "bridge.jar").is_file()
        and (root / "lib").is_dir()
        and any((root / "lib").glob("*.jar"))
    )


def _looks_like_ofd(source: str | Path) -> bool:
    path = Path(source)
    try:
        if not path.is_file() or path.stat().st_size < 8 or path.stat().st_size > MAX_DOCUMENT_BYTES:
            return False
        if not zipfile.is_zipfile(path):
            return False
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            return (
                len(members) <= MAX_ARCHIVE_MEMBERS
                and any(member.filename == "OFD.xml" for member in members)
                and sum(member.file_size for member in members) <= MAX_ARCHIVE_EXPANDED_BYTES
            )
    except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile):
        return False


def _validate_ofd_archive(source: Path) -> None:
    from lxml import etree

    try:
        with zipfile.ZipFile(source) as archive:
            members = archive.infolist()
            if not members or len(members) > MAX_ARCHIVE_MEMBERS:
                raise ConversionError("OFD 文件内部条目数量异常，已停止处理。")
            total_expanded = 0
            names: set[str] = set()
            for member in members:
                name = member.filename
                normalized = PurePosixPath(name)
                if (
                    not name
                    or "\\" in name
                    or "\x00" in name
                    or normalized.is_absolute()
                    or ".." in normalized.parts
                    or (normalized.parts and ":" in normalized.parts[0])
                ):
                    raise ConversionError("OFD 文件包含无效的内部路径，已停止处理。")
                if stat.S_ISLNK(member.external_attr >> 16):
                    raise ConversionError("OFD 文件包含不安全的符号链接条目。")
                if name in names:
                    raise ConversionError("OFD 文件包含重复的内部路径，已停止处理。")
                names.add(name)
                if member.flag_bits & 0x1:
                    raise ConversionError("当前不能读取带密码保护的 OFD 文件。")
                if member.is_dir():
                    continue
                total_expanded += member.file_size
                if total_expanded > MAX_ARCHIVE_EXPANDED_BYTES:
                    raise ConversionError("OFD 文件展开后的内容超过 256 MiB 安全上限。")
                if name.lower().endswith(".xml") and member.file_size > MAX_OFD_XML_BYTES:
                    raise ConversionError("OFD XML 条目超过本机 16 MiB 安全上限。")
            if "OFD.xml" not in names:
                raise ConversionError("这不是有效的 OFD 文档包。")
            if archive.testzip() is not None:
                raise ConversionError("OFD 文件校验失败，压缩包可能已损坏。")
            xml_members = [member for member in members if not member.is_dir() and member.filename.lower().endswith(".xml")]
            for member in xml_members:
                xml_content = archive.read(member)
                # XML markup keywords are ASCII even when the document is UTF-16/32;
                # remove their interleaved NUL bytes before checking the Java parser's input.
                upper_content = xml_content.replace(b"\x00", b"").upper()
                if b"<!DOCTYPE" in upper_content or b"<!ENTITY" in upper_content:
                    raise ConversionError("OFD XML 不允许包含 DTD 或实体声明。")
                if member.filename == "OFD.xml":
                    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False)
                    root = etree.fromstring(xml_content, parser=parser)
                    if etree.QName(root).localname != "OFD":
                        raise ConversionError("OFD 主索引文件结构无效。")
    except ConversionError:
        raise
    except (OSError, ValueError, RuntimeError, etree.LxmlError, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise ConversionError("OFD 文件损坏或无法安全读取。") from exc


def _run_ofdrw_bridge(source: Path, target: str, output: Path) -> None:
    root = _ofd_engine_root()
    java = _ofd_java_path()
    bridge = root / "bridge.jar"
    libraries = root / "lib"
    if not java or not bridge.is_file() or not libraries.is_dir():
        raise ConversionError("本机没有可用的 OFD 转换引擎。")
    classpath = str(bridge) + os.pathsep + str(libraries / "*")
    work_directory = output.parent
    command = [
        java,
        "-Djava.awt.headless=true",
        "-Dfile.encoding=UTF-8",
        "-Xms32m",
        "-Xmx512m",
        "-cp",
        classpath,
        "com.fluke.ofd.OfdBridge",
        str(source),
        target,
        str(output),
    ]
    startup = None
    creation_flags = 0
    if os.name == "nt":
        startup = subprocess.STARTUPINFO()
        startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        startup.wShowWindow = subprocess.SW_HIDE
        creation_flags = subprocess.CREATE_NO_WINDOW
    try:
        process = subprocess.Popen(
            command,
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            cwd=str(work_directory),
            startupinfo=startup,
            creationflags=creation_flags,
        )
    except OSError as exc:
        raise ConversionError("无法启动 OFD 转换引擎。") from exc
    deadline = time.monotonic() + OFD_TIMEOUT_SECONDS
    timed_out = False
    oversized = False
    too_many_pages = False
    try:
        while process.poll() is None:
            if target == "png":
                output_size = 0
                image_count = 0
                if output.is_dir():
                    for candidate in output.glob("*"):
                        if candidate.is_file():
                            output_size += candidate.stat().st_size
                            image_count += candidate.suffix.lower() == ".png"
                            if image_count > MAX_OFD_PAGES:
                                too_many_pages = True
                                process.terminate()
                                break
            else:
                output_size = output.stat().st_size if output.is_file() else 0
            if too_many_pages:
                break
            if output_size > MAX_OFD_OUTPUT_BYTES:
                oversized = True
                try:
                    process.terminate()
                except OSError:
                    pass
                break
            if time.monotonic() >= deadline:
                timed_out = True
                try:
                    process.terminate()
                except OSError:
                    pass
                break
            time.sleep(0.2)
        try:
            return_code = process.wait(timeout=5)
        except subprocess.TimeoutExpired:
            process.kill()
            return_code = process.wait()
    finally:
        if process.poll() is None:
            try:
                process.kill()
                process.wait(timeout=5)
            except (OSError, subprocess.TimeoutExpired):
                pass
    if timed_out:
        raise ConversionError("OFD 转换超时，已停止处理。")
    if oversized:
        raise ConversionError("OFD 输出超过本机 512 MiB 上限，已停止处理。")
    if too_many_pages:
        raise ConversionError("OFD 页数超过本机 500 页安全上限。")
    if return_code != 0:
        raise ConversionError("OFD 引擎无法读取此文件，或转换过程中遇到不支持的内容。")
    if target == "png" and output.is_dir():
        image_count = 0
        output_size = 0
        for candidate in output.glob("*"):
            if candidate.is_file():
                output_size += candidate.stat().st_size
                image_count += candidate.suffix.lower() == ".png"
        if image_count > MAX_OFD_PAGES:
            raise ConversionError("OFD 页数超过本机 500 页安全上限。")
        if output_size > MAX_OFD_OUTPUT_BYTES:
            raise ConversionError("OFD 输出超过本机 512 MiB 上限。")


def _render_ofd_pages(source: Path, page_directory: Path) -> list[Path]:
    _run_ofdrw_bridge(source, "png", page_directory)
    images = sorted(
        (path for path in page_directory.glob("*.png") if path.is_file()),
        key=lambda path: (int(path.stem) if path.stem.isdigit() else MAX_OFD_PAGES + 1, path.name.casefold()),
    )
    if not images:
        raise ConversionError("OFD 引擎没有导出页面图片。")
    if len(images) > MAX_OFD_PAGES:
        raise ConversionError("OFD 页数超过本机 500 页安全上限。")
    total_pixels = 0
    for image_path in images:
        if image_path.stat().st_size > MAX_OFD_OUTPUT_BYTES:
            raise ConversionError("OFD 页面图片超过本机输出上限。")
        reader = QImageReader(str(image_path))
        image = reader.read()
        if image.isNull():
            raise ConversionError("OFD 引擎生成的页面图片无法读取。")
        pixels = image.width() * image.height()
        if pixels > MAX_IMAGE_PIXELS:
            raise ConversionError("OFD 单页图片超过本机 4,000 万像素上限。")
        total_pixels += pixels
        if total_pixels > MAX_PDF_IMAGE_TOTAL_PIXELS:
            raise ConversionError("OFD 页面图片总像素超过本机 5 亿像素上限。")
        del image
    return images


def _write_ofd_pages_zip(source: Path, destination: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="fluke-ofd-pages-") as temporary_directory:
        images = _render_ofd_pages(source, Path(temporary_directory))
        try:
            with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
                for index, image_path in enumerate(images, start=1):
                    archive.write(image_path, f"page-{index:04d}.png")
        except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
            raise ConversionError("无法打包 OFD 页面图片。") from exc
        if destination.stat().st_size > MAX_OFD_OUTPUT_BYTES:
            raise ConversionError("OFD 页面图片 ZIP 超过本机 512 MiB 输出上限。")


def _write_ofd_pdf(source: Path, destination: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="fluke-ofd-pdf-pages-") as temporary_directory:
        images = _render_ofd_pages(source, Path(temporary_directory))
        writer = QPdfWriter(str(destination))
        resolution = 300
        writer.setResolution(resolution)
        writer.setPageMargins(QMarginsF(0, 0, 0, 0), QPageLayout.Unit.Millimeter)
        first_image = QImageReader(str(images[0])).read()
        if first_image.isNull():
            raise ConversionError("OFD 首页图片无法读取。")
        first_width_mm = first_image.width() * 25.4 / resolution
        first_height_mm = first_image.height() * 25.4 / resolution
        first_scale = min(1.0, 5080 / first_width_mm, 5080 / first_height_mm)
        writer.setPageSize(
            QPageSize(
                QSizeF(first_width_mm * first_scale, first_height_mm * first_scale),
                QPageSize.Unit.Millimeter,
            )
        )
        painter = QPainter()
        if not painter.begin(writer):
            raise ConversionError("无法创建 OFD PDF 输出。")
        try:
            for index, image_path in enumerate(images):
                image = first_image if index == 0 else QImageReader(str(image_path)).read()
                if image.isNull():
                    raise ConversionError("OFD 页面图片在 PDF 输出时无法读取。")
                width_mm = image.width() * 25.4 / resolution
                height_mm = image.height() * 25.4 / resolution
                if width_mm <= 0 or height_mm <= 0:
                    raise ConversionError("OFD 页面尺寸无效。")
                scale = min(1.0, 5080 / width_mm, 5080 / height_mm)
                page_size = QPageSize(
                    QSizeF(width_mm * scale, height_mm * scale),
                    QPageSize.Unit.Millimeter,
                )
                if index:
                    writer.setPageSize(page_size)
                    if not writer.newPage():
                        raise ConversionError("无法继续写入 OFD PDF 页面。")
                else:
                    image = first_image
                rect = writer.pageLayout().paintRectPixels(resolution)
                painter.drawImage(rect, image)
                if destination.is_file() and destination.stat().st_size > MAX_OFD_OUTPUT_BYTES:
                    raise ConversionError("OFD PDF 输出超过本机 512 MiB 上限。")
                del image
        finally:
            painter.end()
        if not destination.is_file() or destination.stat().st_size == 0:
            raise ConversionError("OFD PDF 输出为空。")
        if destination.stat().st_size > MAX_OFD_OUTPUT_BYTES:
            raise ConversionError("OFD PDF 输出超过本机 512 MiB 上限。")
        try:
            from pypdf import PdfReader

            with PdfReader(str(destination), strict=False) as reader:
                if len(reader.pages) != len(images):
                    raise ConversionError("OFD PDF 页面数量校验失败。")
        except ConversionError:
            raise
        except Exception as exc:
            raise ConversionError("OFD PDF 输出无法读取。") from exc


def _convert_ofd(source: Path, target: str, destination: Path) -> None:
    _validate_ofd_archive(source)
    if target == "ofd-png-zip":
        _write_ofd_pages_zip(source, destination)
        return
    if target == "pdf":
        _write_ofd_pdf(source, destination)
        return
    bridge_target = {"html": "html", "txt": "text"}.get(target)
    if not bridge_target:
        raise ConversionError("尚不支持所选 OFD 输出格式。")
    _run_ofdrw_bridge(source, bridge_target, destination)
    if not destination.is_file() or destination.stat().st_size == 0:
        raise ConversionError("OFD 引擎没有生成有效输出文件。")
    if destination.stat().st_size > MAX_OFD_OUTPUT_BYTES:
        raise ConversionError("OFD 输出超过本机 512 MiB 上限。")
    if target == "txt" and destination.stat().st_size > MAX_EXTRACTED_TEXT_BYTES:
        raise ConversionError("OFD 文本输出超过本机 32 MiB 安全上限。")
    if target == "html":
        try:
            content = destination.read_bytes()
            if b"<html" not in content.lower() or b"</html>" not in content.lower():
                raise ConversionError("OFD 引擎生成的 HTML 文件校验失败。")
        except OSError as exc:
            raise ConversionError("无法读取 OFD HTML 输出。") from exc


def formats_for(source: str | Path) -> list[dict[str, str]]:
    extension = Path(source).suffix.lower().lstrip(".")
    source_path = Path(source)
    if extension == "ofd":
        if not ofd_engine_available() or not _looks_like_ofd(source):
            return []
        return [
            {"label": "PDF 文档（页面图像版，文字不可搜索）", "value": "pdf"},
            {"label": "HTML 网页", "value": "html"},
            {"label": "纯文本（TXT）", "value": "txt"},
            {"label": "页面图片（PNG ZIP）", "value": "ofd-png-zip"},
        ]
    if extension == "zip" and supported_image_extensions():
        return [{"label": "ZIP 图片包转 PDF", "value": "zip-images-pdf"}]
    if extension in MEDIA_INPUT_EXTENSIONS:
        return [
            {"label": media_target_label(target), "value": target}
            for target in media_targets_for(extension)
        ]
    if extension in TABULAR_FORMATS:
        formats = [
            dict(info)
            for key, info in TABULAR_FORMATS.items()
            if key != extension
        ]
        formats.extend(dict(DOCUMENT_FORMATS[target]) for target in ("html",))
        pdf_engine = (
            "excel" if extension == "xlsx" and "excel" in _office_engines()
            else "wps-excel" if extension == "xlsx" and "excel" in _wps_engines()
            else ""
        )
        formats.append(_document_format("pdf", pdf_engine))
        if _has_docx_support():
            formats.append(dict(DOCUMENT_FORMATS["docx"]))
        if extension in {"csv", "tsv", "xlsx"} and "excel" in _office_engines():
            formats.append(_tabular_format("xls", "excel"))
            formats.append(_tabular_format("ods", "excel"))
        return formats
    if extension in TEXT_EXTENSIONS:
        targets = ["txt", "md", "html"]
        if _has_docx_support():
            targets.append("docx")
        targets.append("pdf")
        if _has_epub_support():
            targets.append("epub")
        same_format = {"markdown": "md", "htm": "html"}.get(extension, extension)
        return [dict(DOCUMENT_FORMATS[target]) for target in targets if target != same_format]
    if extension == "mobi" and _ebook_convert_path():
        return [dict(DOCUMENT_FORMATS[target]) for target in ("epub", "pdf", "docx", "txt")]
    if extension == "docx":
        formats: list[dict[str, str]] = []
        if _has_docx_support():
            formats.extend(dict(DOCUMENT_FORMATS[target]) for target in ("txt", "md", "html"))
            if _has_epub_support():
                formats.append(dict(DOCUMENT_FORMATS["epub"]))
        if _has_docx_support() or "word" in _office_engines() or "word" in _wps_engines():
            pdf_engine = "word" if "word" in _office_engines() else "wps-word" if "word" in _wps_engines() else ""
            formats.append(_document_format("pdf", pdf_engine))
        return formats
    if extension == "epub":
        if not _has_epub_support():
            return []
        formats = [dict(DOCUMENT_FORMATS[target]) for target in ("txt", "md", "html")]
        if _has_docx_support():
            formats.append(dict(DOCUMENT_FORMATS["docx"]))
        if _has_pdf_text_support():
            formats.append(dict(DOCUMENT_FORMATS["pdf"]))
        return formats
    if extension in WPS_WORD_EXTENSIONS and "word" in _wps_engines():
        return [
            _document_format("docx", "wps-word"),
            _document_format("pdf", "wps-word"),
        ]
    if extension in {"doc", "odt", "rtf"} and "word" in _wps_engines():
        return [
            _document_format("docx", "wps-word"),
            _document_format("pdf", "wps-word"),
        ]
    if extension in WORD_OFFICE_EXTENSIONS and "word" in _office_engines():
        return [
            _document_format("docx", "word"),
            _document_format("pdf", "word"),
        ]
    if extension in WPS_EXCEL_EXTENSIONS and "excel" in _wps_engines():
        return [
            _tabular_format("xlsx", "wps-excel"),
            _tabular_format("csv", "wps-excel"),
            _document_format("pdf", "wps-excel"),
        ]
    if extension in EXCEL_OFFICE_EXTENSIONS and "excel" in _office_engines():
        formats = [
            _tabular_format("xlsx", "excel"),
            _tabular_format("csv", "excel"),
        ]
        if extension != "xls":
            formats.append(_tabular_format("xls", "excel"))
        if extension != "ods":
            formats.append(_tabular_format("ods", "excel"))
        formats.append(_document_format("pdf", "excel"))
        return formats
    if extension in WPS_PRESENTATION_EXTENSIONS and "powerpoint" in _wps_engines():
        return [
            _document_format("pptx", "wps-powerpoint"),
            _document_format("pdf", "wps-powerpoint"),
        ]
    if extension == "pptx" and "powerpoint" in _wps_engines() and "powerpoint" not in _office_engines():
        return [_document_format("pdf", "wps-powerpoint")]
    if extension in POWERPOINT_OFFICE_EXTENSIONS | {"pptx"} and "powerpoint" in _office_engines():
        return [_document_format("pdf", "powerpoint")]
    if extension == "pdf":
        if not _looks_like_pdf(source):
            return []
        encrypted = _pdf_encryption_state_for_path(source)
        if encrypted is True:
            return [{"label": "PDF 解密（需原密码）", "value": "pdf-decrypt"}]
        formats: list[dict[str, str]] = []
        if searchable_pdf_ocr_available():
            formats.append({"label": "可搜索 PDF（OCR）", "value": "pdf-ocr"})
        if _has_pdf_text_support():
            formats.extend(dict(DOCUMENT_FORMATS[target]) for target in ("txt", "md", "html"))
            if _has_docx_support():
                formats.append(dict(DOCUMENT_FORMATS["docx"]))
            if _has_epub_support():
                formats.append(dict(DOCUMENT_FORMATS["epub"]))
            formats.append({"label": "按页或每 N 页拆分（ZIP）", "value": "zip"})
            formats.append({"label": "PDF 密码加密（AES-256）", "value": "pdf-encrypt"})
            if _has_pdf_table_support():
                formats.append({"label": "PDF 表格提取（XLSX）", "value": "xlsx"})
        formats.append({"label": "全部页面图片（ZIP）", "value": "pdf-images-zip"})
        if "pdf" in _qt_formats(QImageReader.supportedImageFormats()):
            labels: set[str] = set()
            for key, info in IMAGE_FORMATS.items():
                if key in writable_image_extensions() and info["label"] not in labels:
                    formats.append(
                        {"label": f"{info['label']}（仅第一页）", "value": key, "scope": "first-page"}
                    )
                    labels.add(info["label"])
        return formats
    if extension in READ_ONLY_IMAGE_EXTENSIONS:
        try:
            _check_source(source_path, MAX_IMAGE_BYTES)
            if extension not in RAW_CAMERA_IMAGE_EXTENSIONS or _rawpy_module() is None:
                reader = QImageReader(str(source_path))
                reader.setAutoTransform(True)
                if not reader.canRead():
                    return []
        except (OSError, RuntimeError, ConversionError):
            return []
    rawpy_can_read = extension in RAW_CAMERA_IMAGE_EXTENSIONS and _rawpy_module() is not None
    readable = supported_image_extensions()
    if rawpy_can_read:
        readable.add(extension)
    writable = writable_image_extensions()
    if extension in readable:
        formats: list[dict[str, str]] = []
        labels: set[str] = set()
        source_codec = {"jpg": "jpeg", "jpeg": "jpeg", "tif": "tiff", "tiff": "tiff"}.get(extension, extension)
        for key, info in IMAGE_FORMATS.items():
            if rawpy_can_read and key not in RAW_IMAGE_OUTPUTS:
                continue
            codec = {"jpg": "jpeg", "jpeg": "jpeg", "tif": "tiff", "tiff": "tiff"}.get(key, key)
            if key in writable and codec != source_codec and info["label"] not in labels:
                formats.append(dict(info))
                labels.add(info["label"])
        if extension != "gif" and _pillow_can_save("GIF"):
            frame_count = _pillow_frame_count(source_path)
            if frame_count:
                formats.append({
                    "label": "GIF（保留动画帧）" if frame_count > 1 else "GIF 图像",
                    "value": "gif-animation",
                })
        formats.append(dict(DOCUMENT_FORMATS["pdf"]))
        if searchable_pdf_ocr_available():
            formats.append({"label": "提取文字（OCR/TXT）", "value": "image-ocr-txt"})
        return formats
    return []


def _read_raw_camera_image(source: Path) -> Any:
    rawpy = _rawpy_module()
    if rawpy is None:
        raise ConversionError("当前版本没有可用的相机 RAW 解码器。")
    try:
        with rawpy.imread(str(source)) as raw:
            width = int(raw.sizes.width)
            height = int(raw.sizes.height)
            if width <= 0 or height <= 0 or width * height > MAX_IMAGE_PIXELS:
                raise ConversionError("相机 RAW 解码尺寸超过本机 4000 万像素安全上限。")
            rgb = raw.postprocess(output_bps=8)
        if rgb.ndim != 3 or rgb.shape[2] != 3 or rgb.dtype.name != "uint8":
            raise ConversionError("相机 RAW 解码器返回了不支持的像素数据。")
        if rgb.shape[0] * rgb.shape[1] > MAX_IMAGE_PIXELS:
            raise ConversionError("相机 RAW 解码尺寸超过本机 4000 万像素安全上限。")
        if not rgb.flags.c_contiguous:
            rgb = rgb.copy(order="C")
        image = QImage(
            rgb.data,
            int(rgb.shape[1]),
            int(rgb.shape[0]),
            int(rgb.strides[0]),
            QImage.Format.Format_RGB888,
        ).copy()
        if image.isNull():
            raise ConversionError("无法建立相机 RAW 的输出图像。")
        return image
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"无法解码相机 RAW 文件：{exc}") from exc


def _write_multi_size_ico(image: QImage, destination: Path) -> None:
    icon_sizes = (16, 24, 32, 48, 64, 128, 256)
    payloads: list[tuple[int, bytes]] = []
    for size in icon_sizes:
        canvas = QImage(size, size, QImage.Format.Format_ARGB32)
        canvas.fill(Qt.GlobalColor.transparent)
        scaled = image.scaled(
            size,
            size,
            Qt.AspectRatioMode.KeepAspectRatio,
            Qt.TransformationMode.SmoothTransformation,
        )
        painter = QPainter(canvas)
        painter.setRenderHint(QPainter.RenderHint.SmoothPixmapTransform, True)
        painter.drawImage((size - scaled.width()) // 2, (size - scaled.height()) // 2, scaled)
        painter.end()

        encoded = QByteArray()
        buffer = QBuffer(encoded)
        if not buffer.open(QIODevice.OpenModeFlag.WriteOnly):
            raise ConversionError("无法创建 ICO 图标数据。")
        try:
            if not canvas.save(buffer, "PNG"):
                raise ConversionError("无法编码 ICO 多尺寸图标。")
        finally:
            buffer.close()
        payloads.append((size, bytes(encoded)))

    header_size = 6 + 16 * len(payloads)
    offset = header_size
    entries: list[bytes] = []
    for size, payload in payloads:
        dimension = 0 if size == 256 else size
        entries.append(struct.pack(
            "<BBBBHHII",
            dimension,
            dimension,
            0,
            0,
            0,
            0,
            len(payload),
            offset,
        ))
        offset += len(payload)
    destination.write_bytes(struct.pack("<HHH", 0, 1, len(entries)) + b"".join(entries) + b"".join(
        payload for _, payload in payloads
    ))


def _write_image_ocr_text(
    source: Path,
    destination: Path,
    language: str = "",
    progress_callback: Callable[[int], None] | None = None,
    is_raw_image: bool = False,
) -> None:
    engine = tesseract_engine()
    if engine is None:
        raise ConversionError("图片文字识别需要可用的 Tesseract OCR 引擎和语言模型。")
    selected_language = str(language or engine.language).strip()
    language_parts = selected_language.split("+")
    if (
        len(selected_language) > 128
        or not language_parts
        or not all(part in engine.languages for part in language_parts)
    ):
        raise ConversionError("所选 OCR 语言不受当前可用引擎支持，请重新选择。")
    engine = replace(engine, language=selected_language)
    if progress_callback:
        progress_callback(5)

    with tempfile.TemporaryDirectory(prefix="fluke-image-ocr-") as temp_name:
        if is_raw_image:
            image = _read_raw_camera_image(source)
        else:
            reader = QImageReader(str(source))
            reader.setAutoTransform(True)
            size = reader.size()
            if size.isValid() and size.width() * size.height() > MAX_OCR_PAGE_PIXELS:
                raise ConversionError("图片尺寸超过 OCR 每页 4000 万像素安全上限。")
            image = reader.read()
            if image.isNull():
                raise ConversionError(f"无法解码 OCR 输入图片：{reader.errorString()}")
            if image.width() * image.height() > MAX_OCR_PAGE_PIXELS:
                raise ConversionError("图片尺寸超过 OCR 每页 4000 万像素安全上限。")
        image_path = Path(temp_name) / "ocr-input.png"
        if not image.save(str(image_path), "PNG"):
            raise ConversionError("无法暂存 OCR 输入图片。")
        del image
        if progress_callback:
            progress_callback(20)
        try:
            raw_lines = recognize_tesseract_tsv(engine, image_path, MAX_OCR_PAGE_SECONDS)
        except RuntimeError as exc:
            raise ConversionError(str(exc)) from None

    lines = [str(line.get("text", "")).strip() for line in raw_lines]
    text = "\n".join(line for line in lines if line)
    if not text:
        raise ConversionError("图片 OCR 未识别出文字。")
    if len(text.encode("utf-8")) > MAX_EXTRACTED_TEXT_BYTES:
        raise ConversionError("OCR 提取出的文字超过本机 32 MiB 安全上限。")
    destination.write_text(text + "\n", encoding="utf-8", newline="")
    if progress_callback:
        progress_callback(100)


def _document_format(value: str, engine: str = "") -> dict[str, str]:
    result = dict(DOCUMENT_FORMATS[value])
    return _label_format_for_engine(result, engine)


def _tabular_format(value: str, engine: str = "") -> dict[str, str]:
    source = TABULAR_FORMATS[value] if value in TABULAR_FORMATS else OFFICE_TABULAR_OUTPUTS[value]
    result = dict(source)
    return _label_format_for_engine(result, engine)


def _label_format_for_engine(result: dict[str, str], engine: str) -> dict[str, str]:
    if engine:
        result["engine"] = engine
        names = {
            "word": "Microsoft Word", "excel": "Microsoft Excel", "powerpoint": "Microsoft PowerPoint",
            "wps-word": "WPS 文字", "wps-excel": "WPS 表格", "wps-powerpoint": "WPS 演示",
        }
        result["label"] += f"（{names[engine]}）"
    return result


def _has_docx_support() -> bool:
    try:
        import docx  # noqa: F401
    except ImportError:
        return False
    return True


def _has_pdf_text_support() -> bool:
    try:
        import pypdf  # noqa: F401
    except ImportError:
        return False
    return True


@lru_cache(maxsize=1)
def _has_pdf_table_support() -> bool:
    try:
        import pdfplumber  # noqa: F401
    except ImportError:
        return False
    return _has_pdf_text_support()


@lru_cache(maxsize=128)
def _pdf_encryption_state(path_text: str, size: int, modified_ns: int) -> bool | None:
    try:
        from pypdf import PdfReader

        with PdfReader(path_text, strict=False) as reader:
            return reader.is_encrypted
    except Exception:
        return None


def _pdf_encryption_state_for_path(source: str | Path) -> bool | None:
    path = Path(source)
    try:
        stat_result = path.stat()
    except OSError:
        return None
    return _pdf_encryption_state(str(path.resolve()), stat_result.st_size, stat_result.st_mtime_ns)


@lru_cache(maxsize=1)
def _has_epub_support() -> bool:
    try:
        from lxml import etree  # noqa: F401
        return True
    except ImportError:
        return False


def _ebook_convert_path() -> str | None:
    configured = os.environ.get("FLUKE_EBOOK_CONVERT_PATH", "").strip().strip('"')
    if configured:
        path = Path(configured).expanduser()
        return str(path.resolve()) if path.is_file() else None
    bundled = engine_executable("calibre")
    if bundled:
        return bundled
    return shutil.which("ebook-convert.exe") or shutil.which("ebook-convert")


@lru_cache(maxsize=1)
def _office_engines() -> frozenset[str]:
    if os.name != "nt":
        return frozenset()
    try:
        import winreg
    except ImportError:
        return frozenset()
    engines: set[str] = set()
    for engine, prog_id in (
        ("word", "Word.Application"),
        ("excel", "Excel.Application"),
        ("powerpoint", "PowerPoint.Application"),
    ):
        try:
            with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, prog_id + "\\CLSID"):
                engines.add(engine)
        except OSError:
            continue
    return frozenset(engines)


@lru_cache(maxsize=1)
def _wps_engines() -> frozenset[str]:
    if os.name != "nt":
        return frozenset()
    try:
        import winreg
    except ImportError:
        return frozenset()
    engines: set[str] = set()
    for engine, prog_id in (
        ("word", "KWPS.Application"),
        ("excel", "KET.Application"),
        ("powerpoint", "KWPP.Application"),
    ):
        try:
            with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, prog_id + "\\CLSID"):
                engines.add(engine)
        except OSError:
            continue
    return frozenset(engines)


def _office_operation(extension: str, target: str) -> str | None:
    engines = _office_engines()
    wps_engines = _wps_engines()
    if extension in WPS_WORD_EXTENSIONS | {"doc", "odt", "rtf"} and "word" in wps_engines:
        return {"docx": "wps-word-docx", "pdf": "wps-word-pdf"}.get(target)
    if extension in WPS_EXCEL_EXTENSIONS and "excel" in wps_engines:
        return {"xlsx": "wps-excel-xlsx", "csv": "wps-excel-csv", "pdf": "wps-excel-pdf"}.get(target)
    if extension in WPS_PRESENTATION_EXTENSIONS and "powerpoint" in wps_engines:
        return {"pptx": "wps-powerpoint-pptx", "pdf": "wps-powerpoint-pdf"}.get(target)
    if extension in WORD_OFFICE_EXTENSIONS and "word" in engines:
        return {"docx": "word-docx", "pdf": "word-pdf"}.get(target)
    if extension == "docx" and target == "pdf":
        if "word" in engines:
            return "word-pdf"
        if "word" in wps_engines:
            return "wps-word-pdf"
    if extension == "xlsx" and target == "pdf":
        if "excel" in engines:
            return "excel-pdf"
        if "excel" in wps_engines:
            return "wps-excel-pdf"
    if extension in {"csv", "tsv", "xlsx"} and target in {"xls", "ods"} and "excel" in engines:
        return f"excel-{target}"
    if extension in EXCEL_OFFICE_EXTENSIONS and "excel" in engines:
        return {
            "xlsx": "excel-xlsx", "xls": "excel-xls", "csv": "excel-csv",
            "ods": "excel-ods", "pdf": "excel-pdf",
        }.get(target)
    if extension in POWERPOINT_OFFICE_EXTENSIONS | {"pptx"} and target == "pdf":
        if "powerpoint" in engines:
            return "powerpoint-pdf"
        if "powerpoint" in wps_engines:
            return "wps-powerpoint-pdf"
    return None


def _run_office_conversion(operation: str, source: Path, destination: Path) -> None:
    executable = shutil.which("powershell.exe")
    if not executable:
        raise ConversionError("找不到 Windows PowerShell，无法启动本机 Office/WPS 转换。")
    handle, script_name = tempfile.mkstemp(prefix="fluke-office-", suffix=".ps1")
    script_path = Path(script_name)
    try:
        with os.fdopen(handle, "w", encoding="utf-8-sig", newline="\r\n") as stream:
            stream.write(_OFFICE_AUTOMATION_SCRIPT)
        startup = subprocess.STARTUPINFO()
        startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        startup.wShowWindow = subprocess.SW_HIDE
        try:
            process = subprocess.Popen(
                [
                    executable,
                    "-NoProfile",
                    "-NonInteractive",
                    "-ExecutionPolicy",
                    "Bypass",
                    "-File",
                    str(script_path),
                    "-Operation",
                    operation,
                    "-Source",
                    str(source),
                    "-Destination",
                    str(destination),
                ],
                stdin=subprocess.DEVNULL,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                startupinfo=startup,
                creationflags=subprocess.CREATE_NO_WINDOW,
            )
        except OSError as exc:
            raise ConversionError("无法启动本机 Office/WPS 转换程序。") from exc
        deadline = time.monotonic() + OFFICE_TIMEOUT_SECONDS
        timed_out = False
        oversized = False
        try:
            while process.poll() is None:
                if destination.is_file() and destination.stat().st_size > MAX_OFFICE_OUTPUT_BYTES:
                    oversized = True
                    try:
                        process.terminate()
                    except OSError:
                        pass
                    break
                if time.monotonic() >= deadline:
                    timed_out = True
                    try:
                        process.terminate()
                    except OSError:
                        pass
                    break
                time.sleep(0.2)
            try:
                return_code = process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                process.kill()
                return_code = process.wait()
        finally:
            if process.poll() is None:
                try:
                    process.kill()
                except OSError:
                    pass
                try:
                    process.wait(timeout=5)
                except subprocess.TimeoutExpired:
                    pass
        if timed_out:
            raise ConversionError("本机 Office/WPS 转换超时，已停止处理。")
        if oversized:
            raise ConversionError("Office/WPS 输出超过本机 512 MiB 上限，已提前停止转换。")
        if return_code != 0:
            raise ConversionError("本机 Office/WPS 转换失败，请确认程序已激活且文件未加密，然后重试。")
        if not destination.is_file() or destination.stat().st_size == 0:
            raise ConversionError("本机 Office/WPS 没有生成有效的转换文件。")
        if destination.stat().st_size > MAX_OFFICE_OUTPUT_BYTES:
            raise ConversionError("Office/WPS 输出超过本机 512 MiB 上限。")
    finally:
        script_path.unlink(missing_ok=True)


def _unique_destination(
    source: Path,
    extension: str,
    output_directory: str | Path | None,
    *,
    stem_override: str = "",
) -> Path:
    directory = Path(output_directory).expanduser() if output_directory else source.parent
    directory.mkdir(parents=True, exist_ok=True)
    stem = stem_override or source.stem or "converted"
    candidate = directory / f"{stem}_converted.{extension}"
    suffix = 2
    while True:
        try:
            with candidate.open("xb"):
                return candidate
        except FileExistsError:
            candidate = directory / f"{stem}_converted_{suffix}.{extension}"
            suffix += 1


def _temporary_path(destination: Path, *, preserve_extension: bool = False) -> Path:
    suffix = destination.suffix if preserve_extension else ".tmp"
    handle, name = tempfile.mkstemp(
        prefix=f".{destination.stem}_", suffix=suffix, dir=destination.parent
    )
    os.close(handle)
    return Path(name)


def _save_staged_result(source: Path, staged: Path, output_directory: str | Path | None) -> Path:
    if not staged.is_file() or staged.stat().st_size <= 0:
        raise ConversionError("转换结果已不存在或为空，请重新转换。")
    extension = staged.suffix.lower().lstrip(".")
    if not extension:
        raise ConversionError("转换结果缺少文件扩展名，无法保存。")
    destination = _unique_destination(source, extension, output_directory)
    temporary: Path | None = None
    try:
        temporary = _temporary_path(destination, preserve_extension=True)
        shutil.copyfile(staged, temporary)
        if not temporary.is_file() or temporary.stat().st_size != staged.stat().st_size:
            raise ConversionError("保存时文件长度校验失败，请重试。")
        os.replace(temporary, destination)
        return destination
    except Exception:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        destination.unlink(missing_ok=True)
        raise


_STAGING_OWNER_FILE = ".fluke-session.json"


def _process_is_running(pid: int) -> bool:
    if pid <= 0:
        return False
    if os.name == "nt":
        try:
            import ctypes
            from ctypes import wintypes

            kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
            kernel32.OpenProcess.argtypes = (wintypes.DWORD, wintypes.BOOL, wintypes.DWORD)
            kernel32.OpenProcess.restype = wintypes.HANDLE
            kernel32.GetExitCodeProcess.argtypes = (wintypes.HANDLE, ctypes.POINTER(wintypes.DWORD))
            kernel32.GetExitCodeProcess.restype = wintypes.BOOL
            kernel32.CloseHandle.argtypes = (wintypes.HANDLE,)
            kernel32.CloseHandle.restype = wintypes.BOOL
            handle = kernel32.OpenProcess(0x1000, False, pid)  # PROCESS_QUERY_LIMITED_INFORMATION
            if not handle:
                # Access denied means a process exists but cannot be queried.
                # Unknown errors conservatively preserve its temporary files.
                return ctypes.get_last_error() not in (87, 1168)  # invalid parameter / not found
            try:
                exit_code = wintypes.DWORD()
                if not kernel32.GetExitCodeProcess(handle, ctypes.byref(exit_code)):
                    return True
                return exit_code.value == 259  # STILL_ACTIVE
            finally:
                kernel32.CloseHandle(handle)
        except Exception:
            return True
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    except PermissionError:
        return True
    except OSError:
        return True
    return True


def _cleanup_abandoned_staging_roots(temp_root: str | Path | None = None) -> None:
    try:
        root = Path(temp_root or tempfile.gettempdir()).resolve(strict=True)
        candidates = tuple(root.glob("fluke-converter-*"))
    except OSError:
        return
    for candidate in candidates:
        if candidate.is_symlink() or not candidate.is_dir():
            continue
        try:
            if candidate.resolve(strict=True).parent != root:
                continue
            owner_file = candidate / _STAGING_OWNER_FILE
            if owner_file.is_symlink() or not owner_file.is_file() or owner_file.stat().st_size > 1024:
                continue
            owner = json.loads(owner_file.read_text(encoding="utf-8"))
            pid = owner.get("pid") if isinstance(owner, dict) else None
            if not isinstance(pid, int) or isinstance(pid, bool) or pid <= 0:
                continue
        except (OSError, ValueError, TypeError):
            continue
        if not _process_is_running(pid):
            shutil.rmtree(candidate, ignore_errors=True)


def _run_ebook_convert(source: Path, destination: Path, target: str) -> None:
    executable = _ebook_convert_path()
    if not executable:
        raise ConversionError("未找到本机 Calibre ebook-convert 程序。")
    startup = None
    creation_flags = 0
    if os.name == "nt":
        startup = subprocess.STARTUPINFO()
        startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        startup.wShowWindow = subprocess.SW_HIDE
        creation_flags = subprocess.CREATE_NO_WINDOW
    try:
        process = subprocess.Popen(
            [executable, str(source), str(destination)],
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            startupinfo=startup,
            creationflags=creation_flags,
        )
    except OSError as exc:
        raise ConversionError("无法启动本机 Calibre ebook-convert 程序。") from exc

    deadline = time.monotonic() + EBOOK_TIMEOUT_SECONDS
    timed_out = False
    oversized = False
    try:
        while process.poll() is None:
            if destination.is_file() and destination.stat().st_size > MAX_EBOOK_OUTPUT_BYTES:
                oversized = True
                try:
                    process.terminate()
                except OSError:
                    pass
                break
            if time.monotonic() >= deadline:
                timed_out = True
                try:
                    process.terminate()
                except OSError:
                    pass
                break
            time.sleep(0.2)
        try:
            return_code = process.wait(timeout=5)
        except subprocess.TimeoutExpired:
            process.kill()
            return_code = process.wait()
    finally:
        if process.poll() is None:
            try:
                process.kill()
            except OSError:
                pass
            try:
                process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                pass
    if timed_out:
        raise ConversionError("MOBI 转换超时，已停止处理。")
    if oversized:
        raise ConversionError("电子书输出超过本机 256 MiB 上限，已提前停止转换。")
    if return_code != 0:
        raise ConversionError("本机 Calibre 无法读取 MOBI 文件；文件可能损坏或受 DRM 保护。")
    try:
        if not destination.is_file() or destination.stat().st_size == 0:
            raise ConversionError("本机电子书转换未生成有效文件。")
        if destination.stat().st_size > MAX_EBOOK_OUTPUT_BYTES:
            raise ConversionError("电子书输出超过本机 256 MiB 上限。")
        if target == "pdf":
            with destination.open("rb") as stream:
                header = stream.read(1024)
                stream.seek(max(0, destination.stat().st_size - 1024))
                trailer = stream.read(1024)
            if b"%PDF-" not in header or b"%%EOF" not in trailer:
                raise ConversionError("电子书输出完整性检查未通过。")
        elif target in {"epub", "docx"}:
            with zipfile.ZipFile(destination) as archive:
                names = set(archive.namelist())
                if sum(info.file_size for info in archive.infolist()) > MAX_ARCHIVE_EXPANDED_BYTES:
                    raise ConversionError("电子书输出展开后超过本机安全处理上限。")
                if archive.testzip() is not None:
                    raise ConversionError("电子书输出完整性检查未通过。")
                if target == "epub":
                    members = archive.infolist()
                    if (
                        not members
                        or members[0].filename != "mimetype"
                        or archive.read("mimetype") != b"application/epub+zip"
                        or "META-INF/container.xml" not in names
                    ):
                        raise ConversionError("电子书输出完整性检查未通过。")
                elif "[Content_Types].xml" not in names or "word/document.xml" not in names:
                    raise ConversionError("电子书输出完整性检查未通过。")
    except (OSError, zipfile.BadZipFile, KeyError) as exc:
        raise ConversionError("电子书输出完整性检查未通过。") from exc


def _check_source(path: Path, limit: int) -> None:
    if not path.is_file():
        raise ConversionError("找不到这个文件，或它不是普通文件。")
    size = path.stat().st_size
    if size == 0:
        raise ConversionError("文件是空的。")
    if size > limit:
        raise ConversionError("文件超过本机转换器的大小上限。")


def _read_table(path: Path, extension: str) -> tuple[list[str], list[list[Any]]]:
    if extension == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ConversionError("当前环境没有安装 XLSX 支持组件 openpyxl。") from exc
        try:
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                if not workbook.worksheets:
                    raise ConversionError("工作簿中没有可读取的工作表。")
                sheet = workbook.worksheets[0]
                rows = list(sheet.iter_rows(values_only=True))
            finally:
                workbook.close()
        except ConversionError:
            raise
        except Exception as exc:
            raise ConversionError(f"无法读取 XLSX 工作簿：{exc}") from exc
        if not rows:
            return [], []
        headers, body = _headers_and_body(rows)
        return headers, body

    if extension in ("csv", "tsv"):
        payload = path.read_bytes()
        try:
            content = payload.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                content = payload.decode("gb18030")
            except UnicodeDecodeError as exc:
                raise ConversionError("无法识别文本编码，请先另存为 UTF-8 或 GB18030。") from exc
        delimiter = "\t" if extension == "tsv" else ","
        try:
            rows = list(csv.reader(io.StringIO(content, newline=""), delimiter=delimiter))
        except csv.Error as exc:
            raise ConversionError(f"无法读取分隔文本：{exc}") from exc
        if not rows:
            return [], []
        headers, body = _headers_and_body(rows)
        return headers, body

    try:
        document = json.loads(path.read_text(encoding="utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ConversionError(f"无法读取 JSON：{exc}") from exc
    if isinstance(document, list):
        records = document
    elif isinstance(document, dict):
        records = [document]
    else:
        raise ConversionError("JSON 顶层需要是对象或数组，才能转换为表格。")
    headers: list[str] = []
    for record in records:
        if isinstance(record, dict):
            for key in record:
                name = str(key)
                if name not in headers:
                    headers.append(name)
        elif "value" not in headers:
            headers.append("value")
    body = []
    for record in records:
        values = record if isinstance(record, dict) else {"value": record}
        body.append([_table_value(values.get(header, "")) for header in headers])
    return headers, body


def _headers_and_body(rows: list[tuple[Any, ...]] | list[list[Any]]) -> tuple[list[str], list[list[Any]]]:
    if not rows:
        return [], []
    raw_headers = list(rows[0])
    column_count = max((len(row) for row in rows), default=len(raw_headers))
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index in range(1, column_count + 1):
        value = raw_headers[index - 1] if index <= len(raw_headers) else ""
        base = str(value).strip() if value is not None else ""
        base = base or f"column_{index}"
        count = seen.get(base, 0) + 1
        seen[base] = count
        headers.append(base if count == 1 else f"{base}_{count}")
    body: list[list[Any]] = []
    for row in rows[1:]:
        values = list(row)
        if len(values) < len(headers):
            values.extend([""] * (len(headers) - len(values)))
        body.append(values)
    return headers, body


def _table_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def _json_value(value: Any) -> Any:
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return value


def _write_table(destination: Path, extension: str, headers: list[str], rows: list[list[Any]]) -> None:
    if extension == "json":
        records = [
            {
                header: _json_value(values[index]) if index < len(values) else ""
                for index, header in enumerate(headers)
            }
            for values in rows
        ]
        payload = json.dumps(records, ensure_ascii=False, indent=2) + "\n"
        destination.write_text(payload, encoding="utf-8")
        return
    if extension in ("csv", "tsv"):
        delimiter = "\t" if extension == "tsv" else ","
        with destination.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, delimiter=delimiter)
            if headers:
                writer.writerow([_safe_spreadsheet_text(value) for value in headers])
            writer.writerows([
                [_safe_spreadsheet_text(value) for value in row]
                for row in rows
            ])
        return
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise ConversionError("当前环境没有安装 XLSX 支持组件 openpyxl。") from exc
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    if headers:
        sheet.append(headers)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for row in rows:
        normalized = [_json_value(value) for value in row]
        sheet.append(normalized)
        # Source text beginning with '=' must remain text, not become a formula.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    workbook.save(destination)


def _safe_spreadsheet_text(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    candidate = value.lstrip(" \t\r\n")
    if candidate.startswith(("=", "@")):
        return "'" + value
    if candidate.startswith(("+", "-")):
        try:
            float(candidate)
            return value
        except ValueError:
            return "'" + value
    return value


class _HTMLTextExtractor(HTMLParser):
    _BLOCK_TAGS = {
        "address", "article", "aside", "blockquote", "br", "dd", "div", "dl",
        "dt", "fieldset", "figcaption", "figure", "footer", "form", "h1", "h2",
        "h3", "h4", "h5", "h6", "header", "hr", "li", "main", "nav", "ol",
        "p", "pre", "section", "table", "tr", "ul",
    }
    _HIDDEN_TAGS = {"script", "style", "noscript", "template"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._HIDDEN_TAGS:
            self.hidden_depth += 1
        if not self.hidden_depth and tag in self._BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._HIDDEN_TAGS and self.hidden_depth:
            self.hidden_depth -= 1
        if not self.hidden_depth and tag in self._BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.hidden_depth:
            self.parts.append(data)

    def text(self) -> str:
        lines = [re.sub(r"[ \t\xa0]+", " ", line).strip() for line in "".join(self.parts).splitlines()]
        compact: list[str] = []
        for line in lines:
            if line or (compact and compact[-1]):
                compact.append(line)
        return "\n".join(compact).strip()


def _read_text_file(path: Path) -> str:
    payload = path.read_bytes()
    if payload.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        raise ConversionError("当前不支持 UTF-32 文本，请先另存为 UTF-8、UTF-16 或 GB18030。")
    if payload.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return payload.decode("utf-16")
        except UnicodeDecodeError as exc:
            raise ConversionError("UTF-16 文本损坏，请检查原文件编码。") from exc
    try:
        return payload.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            return payload.decode("gb18030")
        except UnicodeDecodeError as exc:
            raise ConversionError("无法识别文本编码，请先另存为 UTF-8 或 GB18030。") from exc


def _docx_blocks(path: Path) -> list[tuple[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise ConversionError("当前环境没有安装 DOCX 支持组件 python-docx。") from exc
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise ConversionError("DOCX 内部文件数量过多，已停止处理。")
            total_size = sum(member.file_size for member in members)
            if total_size > MAX_ARCHIVE_EXPANDED_BYTES:
                raise ConversionError("DOCX 解压后的内容超过本机安全处理上限。")
            if any(member.file_size > 0 and member.file_size / max(1, member.compress_size) > 250 for member in members):
                raise ConversionError("DOCX 内部压缩比例异常，已停止处理。")
            names = {member.filename for member in members}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise ConversionError("这个文件不是有效的 DOCX 文档。")
        document = Document(str(path))
        blocks: list[tuple[str, Any]] = []
        for item in document.iter_inner_content():
            if hasattr(item, "rows"):
                blocks.append(("table", [[cell.text for cell in row.cells] for row in item.rows]))
            else:
                text = item.text.strip()
                if text:
                    style = getattr(getattr(item, "style", None), "name", "") or ""
                    match = re.match(r"Heading\s+(\d+)", style, re.IGNORECASE)
                    blocks.append(("heading" if match else "paragraph", (match.group(1), text) if match else text))
        return blocks
    except ConversionError:
        raise
    except (OSError, zipfile.BadZipFile, ValueError, KeyError) as exc:
        raise ConversionError(f"无法读取 DOCX 文档：{exc}") from exc


def _epub_member_path(reference: str, base: str = "") -> str:
    parsed_reference = urlsplit(reference)
    if parsed_reference.scheme or parsed_reference.netloc:
        raise ConversionError("EPUB 包内不允许引用外部资源。")
    reference_path = unquote(parsed_reference.path)
    if not reference_path or reference_path.startswith("/") or "\\" in reference_path or "\x00" in reference_path:
        raise ConversionError("EPUB 内部路径无效，已停止处理。")
    parts = list(PurePosixPath(base).parts) if base else []
    for part in reference_path.split("/"):
        if part in ("", "."):
            continue
        if part == "..":
            if not parts:
                raise ConversionError("EPUB 内部路径超出压缩包范围，已停止处理。")
            parts.pop()
        else:
            parts.append(part)
    if not parts:
        raise ConversionError("EPUB 内部路径无效，已停止处理。")
    return "/".join(parts)


def _epub_blocks(path: Path) -> tuple[str, list[tuple[str, Any]]]:
    if not _has_epub_support():
        raise ConversionError("当前环境没有安装 EPUB 支持组件 lxml。")
    from lxml import etree

    def parse_xml(payload: bytes) -> Any:
        parser = etree.XMLParser(
            resolve_entities=False,
            no_network=True,
            load_dtd=False,
            huge_tree=False,
            recover=False,
        )
        return etree.fromstring(payload, parser=parser)

    def local_name(node: Any) -> str:
        if not isinstance(getattr(node, "tag", None), str):
            return ""
        return etree.QName(node).localname

    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if not members or members[0].filename != "mimetype" or members[0].compress_type != zipfile.ZIP_STORED:
                raise ConversionError("这个 EPUB 不符合电子书容器结构要求。")
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise ConversionError("EPUB 内部文件数量过多，已停止处理。")
            total_size = 0
            names: set[str] = set()
            casefolded_names: set[str] = set()
            for member in members:
                name = member.filename
                pure_name = PurePosixPath(name)
                if name.startswith("/") or "\\" in name or "\x00" in name or ".." in pure_name.parts:
                    raise ConversionError("EPUB 含有不安全的内部路径，已停止处理。")
                if name in names or name.casefold() in casefolded_names:
                    raise ConversionError("EPUB 内部文件名重复或大小写冲突，已停止处理。")
                if (member.external_attr >> 16) & 0o170000 == 0o120000:
                    raise ConversionError("EPUB 含有符号链接，已停止处理。")
                names.add(name)
                casefolded_names.add(name.casefold())
                total_size += member.file_size
                if total_size > MAX_ARCHIVE_EXPANDED_BYTES:
                    raise ConversionError("EPUB 解压后的内容超过本机安全处理上限。")
                if member.file_size and member.file_size / max(1, member.compress_size) > 250:
                    raise ConversionError("EPUB 内部压缩比例异常，已停止处理。")
                if member.flag_bits & 0x1:
                    raise ConversionError("EPUB 内部文件已加密，当前版本无法读取。")
            if archive.read("mimetype") != b"application/epub+zip":
                raise ConversionError("这个文件不是有效的 EPUB 电子书。")
            if "META-INF/encryption.xml" in names:
                raise ConversionError("EPUB 含有加密资源，当前版本无法读取。")
            if "META-INF/container.xml" not in names:
                raise ConversionError("EPUB 缺少容器目录信息。")

            container = parse_xml(archive.read("META-INF/container.xml"))
            rootfile = next(
                (node for node in container.iter() if local_name(node) == "rootfile"),
                None,
            )
            if rootfile is None or not rootfile.get("full-path"):
                raise ConversionError("EPUB 没有指定有效的书籍目录文件。")
            package_path = _epub_member_path(rootfile.get("full-path", ""))
            if package_path not in names:
                raise ConversionError("EPUB 书籍目录文件不存在。")
            package = parse_xml(archive.read(package_path))

            title = ""
            for node in package.iter():
                if local_name(node) == "title" and (node.text or "").strip():
                    title = (node.text or "").strip()
                    break
            package_directory = posixpath.dirname(package_path)
            manifest: dict[str, tuple[str, str]] = {}
            for node in package.iter():
                if local_name(node) != "item":
                    continue
                item_id = node.get("id")
                href = node.get("href")
                media_type = node.get("media-type", "").lower()
                if item_id and href:
                    manifest[item_id] = (href, media_type)
            spine = next(
                (node for node in package.iter() if local_name(node) == "spine"),
                None,
            )
            if spine is None:
                raise ConversionError("EPUB 缺少章节顺序信息。")

            blocks: list[tuple[str, Any]] = []
            total_text_bytes = 0
            chapter_count = 0
            for reference in spine:
                if local_name(reference) != "itemref" or reference.get("linear", "yes") == "no":
                    continue
                item = manifest.get(reference.get("idref", ""))
                if item is None:
                    raise ConversionError("EPUB 章节目录引用了缺失的内容文件。")
                href, media_type = item
                if media_type not in ("application/xhtml+xml", "application/html+xml"):
                    continue
                chapter_path = _epub_member_path(href, package_directory)
                if chapter_path not in names:
                    raise ConversionError("EPUB 章节内容文件不存在。")
                chapter = parse_xml(archive.read(chapter_path))
                body = next(
                    (node for node in chapter.iter() if local_name(node) == "body"),
                    chapter,
                )
                chapter_title = ""
                for node in body.iter():
                    if local_name(node) in {"h1", "h2", "h3"}:
                        chapter_title = " ".join("".join(node.itertext()).split())
                        if chapter_title:
                            break
                html_body = etree.tostring(body, encoding="unicode", method="html")
                text_parser = _HTMLTextExtractor()
                text_parser.feed(html_body)
                chapter_text = text_parser.text()
                if not chapter_text:
                    continue
                chapter_count += 1
                if chapter_count > MAX_PDF_PAGES:
                    raise ConversionError("EPUB 章节数量超过本机处理上限。")
                total_text_bytes += len(chapter_text.encode("utf-8"))
                if total_text_bytes > MAX_EXTRACTED_TEXT_BYTES:
                    raise ConversionError("EPUB 提取出的文字超过本机处理上限。")
                if not chapter_title:
                    chapter_title = f"第 {chapter_count} 章"
                blocks.append(("heading", ("1", chapter_title)))
                for line in chapter_text.splitlines():
                    if line and line != chapter_title:
                        blocks.append(("paragraph", line))

            if not chapter_count:
                raise ConversionError("EPUB 中没有可读取的章节文字。")
            return title or path.stem, blocks
    except ConversionError:
        raise
    except (OSError, zipfile.BadZipFile, KeyError, ValueError, etree.XMLSyntaxError, RuntimeError) as exc:
        raise ConversionError("EPUB 文件损坏或结构不受支持，无法读取。") from exc


def _pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ConversionError("当前环境没有安装 PDF 文本支持组件 pypdf。") from exc
    try:
        with PdfReader(str(path), strict=False) as reader:
            if reader.is_encrypted:
                raise ConversionError("PDF 已加密；当前版本不会尝试密码破解或解锁。")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ConversionError("PDF 页数超过本机文本提取上限。")
            pages: list[str] = []
            size = 0
            for page in reader.pages:
                content = page.extract_text() or ""
                size += len(content.encode("utf-8"))
                if size > MAX_EXTRACTED_TEXT_BYTES:
                    raise ConversionError("PDF 提取出的文字过多，超过本机处理上限。")
                if content.strip():
                    pages.append(content.strip())
        result = "\n\n".join(pages)
        if not result:
            raise ConversionError("PDF 中没有可提取的文字。扫描件需要 OCR，本阶段尚未接入。")
        return result
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"无法提取 PDF 文字：{exc}") from exc


def _pdf_docx_line(raw_line: dict[str, Any]) -> dict[str, Any] | None:
    text = str(raw_line.get("text", "")).strip()
    if not text:
        return None
    chars = raw_line.get("chars") or []
    sizes = sorted(
        float(char["size"])
        for char in chars
        if char.get("size") is not None and float(char["size"]) > 0
    )
    size = sizes[len(sizes) // 2] if sizes else max(6.0, float(raw_line.get("bottom", 0)) - float(raw_line.get("top", 0)))
    font_names = [str(char.get("fontname", "")) for char in chars if char.get("fontname")]
    font_name = Counter(font_names).most_common(1)[0][0] if font_names else "Arial"
    font_name = font_name.rsplit("+", 1)[-1]
    lower_font = font_name.casefold()
    for source_name, target_name in (
        ("times", "Times New Roman"),
        ("courier", "Courier New"),
        ("arial", "Arial"),
        ("calibri", "Calibri"),
        ("simsun", "SimSun"),
        ("songti", "SimSun"),
        ("yahei", "Microsoft YaHei"),
    ):
        if source_name in lower_font:
            font_name = target_name
            break
    else:
        font_name = "Arial"
    return {
        "text": text,
        "x0": max(0.0, float(raw_line.get("x0", 0))),
        "top": max(0.0, float(raw_line.get("top", 0))),
        "x1": max(0.0, float(raw_line.get("x1", 0))),
        "bottom": max(0.0, float(raw_line.get("bottom", 0))),
        "font_size": min(72.0, max(4.0, size)),
        "font_name": font_name,
        "bold": any("bold" in str(char.get("fontname", "")).casefold() for char in chars),
        "italic": any(
            token in str(char.get("fontname", "")).casefold()
            for char in chars for token in ("italic", "oblique")
        ),
        "confidence": None,
    }


def _write_pdf_layout_docx(
    source: Path,
    destination: Path,
    progress_callback: Callable[[int], None] | None = None,
) -> None:
    try:
        import pdfplumber
    except ImportError as exc:
        raise ConversionError("当前环境没有安装 PDF 坐标提取组件 pdfplumber。") from exc
    try:
        from docx import Document
        from docx.enum.section import WD_SECTION_START
        from docx.oxml import OxmlElement
        from docx.shared import Pt
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise ConversionError("当前环境没有安装 DOCX 支持组件 python-docx。") from exc

    pages: list[dict[str, Any]] = []
    total_chars = 0
    total_text_bytes = 0
    needs_ocr: list[int] = []
    try:
        with pdfplumber.open(str(source)) as pdf:
            page_count = len(pdf.pages)
            if page_count < 1:
                raise ConversionError("PDF 中没有可转换的页面。")
            if page_count > MAX_PDF_LAYOUT_PAGES:
                raise ConversionError("PDF 页数超过本机版式重建上限。")
            for page_index, page in enumerate(pdf.pages):
                chars = page.chars
                char_count = len(chars)
                total_chars += char_count
                if char_count > MAX_PDF_LAYOUT_PAGE_CHARS:
                    raise ConversionError("PDF 单页字符数量过大，已停止版式重建。")
                if total_chars > MAX_PDF_LAYOUT_TOTAL_CHARS:
                    raise ConversionError("PDF 字符总量过大，已停止版式重建。")
                text = page.extract_text() or ""
                total_text_bytes += len(text.encode("utf-8"))
                if total_text_bytes > MAX_EXTRACTED_TEXT_BYTES:
                    raise ConversionError("PDF 提取出的文字超过本机 32 MiB 安全上限。")
                raw_lines = page.extract_text_lines(strip=True, return_chars=True)
                lines = [
                    line for raw in raw_lines
                    if (line := _pdf_docx_line(raw)) is not None
                ]
                has_page_image = bool(page.images)
                if not text.strip() and has_page_image:
                    needs_ocr.append(page_index)
                elif text.strip():
                    expected = Counter(char for char in text if not char.isspace())
                    observed = Counter(char for line in lines for char in line["text"] if not char.isspace())
                    total_expected = sum(expected.values())
                    matched = sum(min(count, observed[char]) for char, count in expected.items())
                    if total_expected and matched / total_expected < 0.90:
                        raise ConversionError("PDF 文字覆盖不足，已拒绝生成可能漏字的 Word 文件。")
                pages.append({
                    "width": float(page.width),
                    "height": float(page.height),
                    "lines": lines,
                    "source_text": text,
                })
                page.close()
                if progress_callback is not None:
                    progress_callback(round((page_index + 1) * 30 / page_count))
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError("无法分析 PDF 页面文字与坐标，不能安全重建 Word 版式。") from exc

    if not any(page["lines"] for page in pages) and not needs_ocr:
        raise ConversionError("PDF 中没有可编辑文字；扫描页面需要可用的 Tesseract OCR 引擎。")

    ocr_confidences: list[float] = []
    if needs_ocr:
        engine = tesseract_engine()
        if engine is None:
            raise ConversionError("PDF 含扫描页面；PDF→Word 扫描页重建需要可用的 Tesseract OCR 引擎和对应语言模型。")
        if len(needs_ocr) > MAX_OCR_PAGES:
            raise ConversionError("需要 OCR 的扫描页数量超过本机 100 页上限。")
        pdf_document = QPdfDocument()
        try:
            pdf_document.load(str(source))
            if pdf_document.status() != QPdfDocument.Status.Ready or pdf_document.pageCount() != len(pages):
                raise ConversionError("Qt PDF 引擎无法读取扫描页面。")
            dimensions: dict[int, QSize] = {}
            total_pixels = 0
            for page_index in needs_ocr:
                point_size = pdf_document.pagePointSize(page_index)
                width = max(1, math.ceil(point_size.width() * OCR_DPI / 72))
                height = max(1, math.ceil(point_size.height() * OCR_DPI / 72))
                pixels = width * height
                if pixels > MAX_OCR_PAGE_PIXELS or total_pixels + pixels > MAX_PDF_LAYOUT_PIXELS:
                    raise ConversionError("扫描页面的 OCR 总像素数超过本机安全上限。")
                dimensions[page_index] = QSize(width, height)
                total_pixels += pixels
            deadline = time.monotonic() + MAX_OCR_DOCUMENT_SECONDS
            with tempfile.TemporaryDirectory(prefix="fluke-pdf-word-ocr-") as temp_name:
                temp_dir = Path(temp_name)
                for scan_index, page_index in enumerate(needs_ocr, start=1):
                    image = pdf_document.render(page_index, dimensions[page_index])
                    if image.isNull():
                        raise ConversionError("无法渲染 PDF 页面供 OCR。")
                    image_path = temp_dir / f"scan-{scan_index:04d}.png"
                    if not image.save(str(image_path), "PNG"):
                        raise ConversionError("无法暂存 PDF 扫描页图片。")
                    del image
                    timeout = min(MAX_OCR_PAGE_SECONDS, deadline - time.monotonic())
                    if timeout <= 0:
                        raise ConversionError("PDF→Word OCR 总处理时间超过本机上限。")
                    try:
                        raw_lines = recognize_tesseract_tsv(engine, image_path, timeout)
                    except RuntimeError as exc:
                        raise ConversionError(str(exc)) from None
                    page_width = float(pages[page_index]["width"])
                    page_height = float(pages[page_index]["height"])
                    scale_x = page_width / dimensions[page_index].width()
                    scale_y = page_height / dimensions[page_index].height()
                    lines = []
                    for raw in raw_lines:
                        lines.append({
                            "text": str(raw["text"]),
                            "x0": float(raw["x0"]) * scale_x,
                            "top": float(raw["top"]) * scale_y,
                            "x1": float(raw["x1"]) * scale_x,
                            "bottom": float(raw["bottom"]) * scale_y,
                            "font_size": 10.0,
                            "font_name": "Arial",
                            "bold": False,
                            "italic": False,
                            "confidence": float(raw["confidence"]),
                        })
                    if not lines:
                        raise ConversionError("PDF 扫描页面 OCR 未识别出文字。")
                    page_text = "\n".join(line["text"] for line in lines)
                    total_text_bytes += len(page_text.encode("utf-8"))
                    if total_text_bytes > MAX_EXTRACTED_TEXT_BYTES:
                        raise ConversionError("OCR 提取出的文字超过本机 32 MiB 安全上限。")
                    pages[page_index]["lines"] = lines
                    ocr_confidences.extend(float(line["confidence"]) for line in lines)
                    if progress_callback is not None:
                        progress_callback(30 + round(scan_index * 50 / len(needs_ocr)))
        finally:
            pdf_document.close()

    document = Document()
    document.styles["Normal"].font.name = "Arial"
    section_count = len(pages)
    total_lines = sum(len(page["lines"]) for page in pages)
    if total_lines == 0:
        raise ConversionError("PDF 中没有可写入 Word 的文字。")
    for page_index, page in enumerate(pages):
        section = document.sections[0] if page_index == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
        width = float(page["width"])
        height = float(page["height"])
        if not 1 <= width <= 1_584 or not 1 <= height <= 1_584:
            raise ConversionError("PDF 页面尺寸超出 Word 文档支持范围。")
        section.page_width = Pt(width)
        section.page_height = Pt(height)
        section.top_margin = Pt(0)
        section.bottom_margin = Pt(0)
        section.left_margin = Pt(0)
        section.right_margin = Pt(0)
        section.header_distance = Pt(0)
        section.footer_distance = Pt(0)

        cursor = 0.0
        for line in sorted(page["lines"], key=lambda value: (float(value["top"]), float(value["x0"]))):
            text = "".join(
                char for char in str(line["text"])
                if char in "\t\n\r" or ord(char) >= 32 and ord(char) not in (0xFFFE, 0xFFFF)
            )
            if not text:
                continue
            top = min(height, max(0.0, float(line["top"])))
            bottom = min(height, max(top + 1.0, float(line["bottom"])))
            line_height = min(72.0, max(float(line["font_size"]) * 1.15, bottom - top, 6.0))
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(min(width - 1, max(0.0, float(line["x0"]))))
            paragraph.paragraph_format.space_before = Pt(max(0.0, top - cursor))
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(line_height)
            paragraph.paragraph_format.widow_control = False
            run = paragraph.add_run(text)
            run.font.name = str(line["font_name"])
            run.font.size = Pt(float(line["font_size"]))
            run.bold = bool(line["bold"])
            run.italic = bool(line["italic"])
            run_properties = run._element.get_or_add_rPr()
            run_fonts = run_properties.rFonts
            if run_fonts is None:
                run_fonts = OxmlElement("w:rFonts")
                run_properties.insert(0, run_fonts)
            run_fonts.set(qn("w:ascii"), str(line["font_name"]))
            run_fonts.set(qn("w:hAnsi"), str(line["font_name"]))
            run_fonts.set(qn("w:eastAsia"), str(line["font_name"]))
            cursor = max(cursor, top) + line_height
        if progress_callback is not None:
            progress_callback(80 + round((page_index + 1) * 20 / section_count))

    document.core_properties.comments = (
        "Text was reconstructed as editable Word paragraphs from PDF page coordinates. "
        "PDF vector artwork and embedded images are not included."
    )
    if ocr_confidences:
        average = sum(ocr_confidences) / len(ocr_confidences)
        document.core_properties.subject = f"Scanned pages reconstructed with local OCR; average confidence {average:.0f}%."
    document.save(destination)
    if destination.stat().st_size > MAX_PDF_LAYOUT_DOCX_BYTES:
        raise ConversionError("DOCX 输出超过本机 128 MiB 安全上限。")


def _normalize_pdf_table(
    rows: list[list[str | None]],
    *,
    preserve_grid: bool = False,
) -> list[list[str]]:
    normalized = [
        ["" if value is None else str(value).strip() for value in row]
        for row in rows
    ]
    if not preserve_grid:
        normalized = [row for row in normalized if any(value for value in row)]
    if not normalized:
        return []
    width = max(len(row) for row in normalized)
    normalized = [row + [""] * (width - len(row)) for row in normalized]
    if preserve_grid:
        return normalized if len(normalized) >= 2 and width >= 2 else []
    active_columns = [
        index for index in range(width)
        if any(row[index] for row in normalized)
    ]
    normalized = [[row[index] for index in active_columns] for row in normalized]
    if len(normalized) < 2 or len(active_columns) < 2:
        return []
    return normalized


def _cluster_pdf_coordinates(values: list[float], tolerance: float = 1.0) -> list[float]:
    clusters: list[list[float]] = []
    for value in sorted(values):
        if clusters and value - clusters[-1][-1] <= tolerance:
            clusters[-1].append(value)
        else:
            clusters.append([value])
    return [sum(cluster) / len(cluster) for cluster in clusters]


def _pdf_coordinate_index(value: float, coordinates: list[float], tolerance: float = 1.0) -> int | None:
    if not coordinates:
        return None
    index = min(range(len(coordinates)), key=lambda item: abs(coordinates[item] - value))
    return index if abs(coordinates[index] - value) <= tolerance else None


def _pdf_edge_covers_segment(
    edges: list[dict[str, object]],
    orientation: str,
    coordinate: float,
    start: float,
    end: float,
    tolerance: float = 1.0,
) -> bool:
    intervals: list[tuple[float, float]] = []
    for edge in edges:
        if edge.get("orientation") != orientation:
            continue
        try:
            if orientation == "v":
                position = (float(edge["x0"]) + float(edge["x1"])) / 2
                edge_start = float(edge["top"])
                edge_end = float(edge["bottom"])
            else:
                position = (float(edge["top"]) + float(edge["bottom"])) / 2
                edge_start = float(edge["x0"])
                edge_end = float(edge["x1"])
        except (KeyError, TypeError, ValueError):
            continue
        if abs(position - coordinate) <= tolerance:
            left = max(start, edge_start)
            right = min(end, edge_end)
            if right > left:
                intervals.append((left, right))
    intervals.sort()
    covered_until = start
    for left, right in intervals:
        if left > covered_until + tolerance:
            return False
        covered_until = max(covered_until, right)
        if covered_until >= end - tolerance:
            return True
    return covered_until >= end - tolerance


def _pdf_edge_overlaps_segment(
    edges: list[dict[str, object]],
    orientation: str,
    coordinate: float,
    start: float,
    end: float,
    tolerance: float = 1.0,
) -> bool:
    for edge in edges:
        if edge.get("orientation") != orientation:
            continue
        try:
            if orientation == "v":
                position = (float(edge["x0"]) + float(edge["x1"])) / 2
                edge_start = float(edge["top"])
                edge_end = float(edge["bottom"])
            else:
                position = (float(edge["top"]) + float(edge["bottom"])) / 2
                edge_start = float(edge["x0"])
                edge_end = float(edge["x1"])
        except (KeyError, TypeError, ValueError):
            continue
        overlap = min(end, edge_end) - max(start, edge_start)
        if abs(position - coordinate) <= tolerance and overlap > tolerance:
            return True
    return False


def _pdf_table_merge_ranges(
    table: Any,
    rows: list[list[str]],
    page_edges: list[dict[str, object]],
) -> list[tuple[int, int, int, int]]:
    """Return only text-backed one-dimensional merges supported by complete grid geometry."""
    raw_cells = getattr(table, "cells", None)
    if not isinstance(raw_cells, list) or not raw_cells:
        return []
    cells: list[tuple[float, float, float, float]] = []
    try:
        for raw_cell in raw_cells:
            if not isinstance(raw_cell, (tuple, list)) or len(raw_cell) != 4:
                return []
            x0, top, x1, bottom = (float(value) for value in raw_cell)
            if x1 <= x0 or bottom <= top:
                return []
            cells.append((x0, top, x1, bottom))
    except (TypeError, ValueError):
        return []

    x_grid = _cluster_pdf_coordinates([value for cell in cells for value in (cell[0], cell[2])])
    y_grid = _cluster_pdf_coordinates([value for cell in cells for value in (cell[1], cell[3])])
    column_starts = _cluster_pdf_coordinates([cell[0] for cell in cells])
    row_starts = _cluster_pdf_coordinates([cell[1] for cell in cells])
    if len(x_grid) < 3 or len(y_grid) < 3:
        return []
    if len(x_grid) - 1 != len(column_starts) or len(y_grid) - 1 != len(row_starts):
        return []
    if any(abs(left - right) > 1.0 for left, right in zip(x_grid[:-1], column_starts)):
        return []
    if any(abs(left - right) > 1.0 for left, right in zip(y_grid[:-1], row_starts)):
        return []
    if len(rows) != len(row_starts) or any(len(row) != len(column_starts) for row in rows):
        return []
    grid_slot_count = len(row_starts) * len(column_starts)
    if grid_slot_count > MAX_PDF_TABLE_CELLS:
        return []

    occupied: set[tuple[int, int]] = set()
    mapped_cells: list[tuple[int, int, int, int]] = []
    for x0, top, x1, bottom in cells:
        column_start = _pdf_coordinate_index(x0, x_grid)
        column_end = _pdf_coordinate_index(x1, x_grid)
        row_start = _pdf_coordinate_index(top, y_grid)
        row_end = _pdf_coordinate_index(bottom, y_grid)
        if None in (column_start, column_end, row_start, row_end):
            return []
        assert column_start is not None and column_end is not None
        assert row_start is not None and row_end is not None
        if column_end <= column_start or row_end <= row_start:
            return []
        span_slot_count = (row_end - row_start) * (column_end - column_start)
        if len(occupied) + span_slot_count > grid_slot_count:
            return []
        mapped_cells.append((row_start, column_start, row_end, column_end))
        for row_index in range(row_start, row_end):
            for column_index in range(column_start, column_end):
                slot = (row_index, column_index)
                if slot in occupied:
                    return []
                occupied.add(slot)

    expected_slots = {
        (row_index, column_index)
        for row_index in range(len(row_starts))
        for column_index in range(len(column_starts))
    }
    if occupied != expected_slots:
        return []

    merges: list[tuple[int, int, int, int]] = []
    for row_start, column_start, row_end, column_end in mapped_cells:
        row_span = row_end - row_start
        column_span = column_end - column_start
        if row_span == 1 and column_span == 1:
            continue
        # Two-dimensional spans and empty anchors are ambiguous; leave the grid intact.
        if (row_span > 1) == (column_span > 1):
            continue
        anchor = rows[row_start][column_start].strip()
        if not anchor:
            continue
        if any(
            rows[row_index][column_index].strip()
            for row_index in range(row_start, row_end)
            for column_index in range(column_start, column_end)
            if (row_index, column_index) != (row_start, column_start)
        ):
            continue

        left, top = x_grid[column_start], y_grid[row_start]
        right, bottom = x_grid[column_end], y_grid[row_end]
        if not all((
            _pdf_edge_covers_segment(page_edges, "h", top, left, right),
            _pdf_edge_covers_segment(page_edges, "h", bottom, left, right),
            _pdf_edge_covers_segment(page_edges, "v", left, top, bottom),
            _pdf_edge_covers_segment(page_edges, "v", right, top, bottom),
        )):
            continue
        if row_span == 1:
            has_internal_separator = any(
                _pdf_edge_overlaps_segment(page_edges, "v", x_grid[index], top, bottom)
                for index in range(column_start + 1, column_end)
            )
        else:
            has_internal_separator = any(
                _pdf_edge_overlaps_segment(page_edges, "h", y_grid[index], left, right)
                for index in range(row_start + 1, row_end)
            )
        if not has_internal_separator:
            merges.append((row_start + 1, column_start + 1, row_end, column_end))
    return merges


def _table_bbox_is_duplicate(candidate: tuple[float, float, float, float], existing: tuple[float, float, float, float]) -> bool:
    left = max(candidate[0], existing[0])
    top = max(candidate[1], existing[1])
    right = min(candidate[2], existing[2])
    bottom = min(candidate[3], existing[3])
    intersection = max(0.0, right - left) * max(0.0, bottom - top)
    candidate_area = max(0.0, candidate[2] - candidate[0]) * max(0.0, candidate[3] - candidate[1])
    existing_area = max(0.0, existing[2] - existing[0]) * max(0.0, existing[3] - existing[1])
    smaller_area = min(candidate_area, existing_area)
    return smaller_area > 0 and intersection / smaller_area >= 0.75


def _pdf_bboxes_intersect(first: tuple[float, float, float, float], second: tuple[float, float, float, float]) -> bool:
    return max(first[0], second[0]) < min(first[2], second[2]) and max(first[1], second[1]) < min(first[3], second[3])


def _subtract_pdf_bbox(
    source_bbox: tuple[float, float, float, float],
    cut_bbox: tuple[float, float, float, float],
) -> list[tuple[float, float, float, float]]:
    x0, top, x1, bottom = source_bbox
    cut_x0, cut_top, cut_x1, cut_bottom = cut_bbox
    left = max(x0, cut_x0)
    upper = max(top, cut_top)
    right = min(x1, cut_x1)
    lower = min(bottom, cut_bottom)
    if left >= right or upper >= lower:
        return [source_bbox]
    regions = [
        (x0, top, left, bottom),
        (right, top, x1, bottom),
        (left, top, right, upper),
        (left, lower, right, bottom),
    ]
    return [region for region in regions if region[2] - region[0] >= 10 and region[3] - region[1] >= 10]


def _ocr_page_tables(
    raw_lines: list[dict[str, object]],
    page_width: int,
) -> list[dict[str, object]]:
    """Find repeated OCR column layouts and return table rows with word confidence."""
    words: list[dict[str, object]] = []
    for line in raw_lines:
        line_words = line.get("words")
        if not isinstance(line_words, list):
            continue
        for word in line_words:
            if not isinstance(word, dict):
                continue
            text = str(word.get("text", "")).strip()
            try:
                x0 = float(word["x0"])
                top = float(word["top"])
                x1 = float(word["x1"])
                bottom = float(word["bottom"])
                confidence = max(0.0, min(100.0, float(word.get("confidence", 0))))
            except (KeyError, TypeError, ValueError):
                continue
            if text and x1 > x0 and bottom > top:
                words.append({
                    "text": text,
                    "x0": x0,
                    "top": top,
                    "x1": x1,
                    "bottom": bottom,
                    "confidence": confidence,
                })
    if len(words) < 4 or page_width < 1:
        return []

    words.sort(key=lambda word: ((float(word["top"]) + float(word["bottom"])) / 2, float(word["x0"])))
    median_height = sorted(float(word["bottom"]) - float(word["top"]) for word in words)[len(words) // 2]
    row_tolerance = max(3.0, median_height * 0.55)
    row_groups: list[dict[str, object]] = []
    for word in words:
        center = (float(word["top"]) + float(word["bottom"])) / 2
        matches = [
            group for group in row_groups
            if abs(float(group["center"]) - center) <= row_tolerance
        ]
        if matches:
            group = min(matches, key=lambda value: abs(float(value["center"]) - center))
            group_words = group["words"]
            assert isinstance(group_words, list)
            group_words.append(word)
            group["center"] = sum(
                (float(item["top"]) + float(item["bottom"])) / 2 for item in group_words
            ) / len(group_words)
        else:
            row_groups.append({"center": center, "words": [word]})

    row_groups.sort(key=lambda group: float(group["center"]))
    segmented_rows: list[dict[str, object]] = []
    gap_threshold = max(14.0, median_height * 0.85)
    for group in row_groups:
        row_words = group["words"]
        assert isinstance(row_words, list)
        row_words.sort(key=lambda word: float(word["x0"]))
        cells: list[list[dict[str, object]]] = []
        for word in row_words:
            if not cells:
                cells.append([word])
                continue
            previous = cells[-1][-1]
            gap = float(word["x0"]) - float(previous["x1"])
            if gap > gap_threshold:
                cells.append([word])
            else:
                cells[-1].append(word)
        normalized_cells = []
        for cell_words in cells:
            normalized_cells.append({
                "text": " ".join(str(word["text"]) for word in cell_words),
                "confidence": sum(float(word["confidence"]) for word in cell_words) / len(cell_words),
                "x0": min(float(word["x0"]) for word in cell_words),
                "x1": max(float(word["x1"]) for word in cell_words),
            })
        if normalized_cells:
            segmented_rows.append({"center": float(group["center"]), "cells": normalized_cells})

    if not segmented_rows:
        return []
    max_row_gap = max(80.0, median_height * 6)
    runs: list[list[dict[str, object]]] = []
    for row in segmented_rows:
        if not runs or float(row["center"]) - float(runs[-1][-1]["center"]) > max_row_gap:
            runs.append([row])
        else:
            runs[-1].append(row)

    tables: list[dict[str, object]] = []
    alignment_tolerance = max(18.0, page_width * 0.025)
    for run in runs:
        aligned_groups: list[dict[str, object]] = []
        for row in run:
            cells = row["cells"]
            assert isinstance(cells, list)
            starts = [float(cell["x0"]) for cell in cells]
            matches = [
                group for group in aligned_groups
                if len(group["starts"]) == len(starts)
                and all(abs(float(expected) - actual) <= alignment_tolerance
                        for expected, actual in zip(group["starts"], starts))
            ]
            if matches:
                group = min(
                    matches,
                    key=lambda value: sum(
                        abs(float(expected) - actual)
                        for expected, actual in zip(value["starts"], starts)
                    ),
                )
                group_rows = group["rows"]
                assert isinstance(group_rows, list)
                group_rows.append(row)
                group["starts"] = [
                    (float(expected) * (len(group_rows) - 1) + actual) / len(group_rows)
                    for expected, actual in zip(group["starts"], starts)
                ]
            else:
                aligned_groups.append({"starts": starts, "rows": [row]})

        best = max(
            (group for group in aligned_groups if len(group["rows"]) >= 2),
            key=lambda group: len(group["rows"]),
            default=None,
        )
        if best is None:
            continue
        column_starts = [float(value) for value in best["starts"]]
        core_rows = best["rows"]
        core_ids = {id(row) for row in core_rows}
        core_centers = [float(row["center"]) for row in core_rows]
        minimum_center = min(core_centers) - median_height
        maximum_center = max(core_centers) + median_height
        table_rows: list[list[dict[str, object]]] = []
        aligned_rows: list[tuple[float, list[dict[str, object]]]] = []
        for row in run:
            row_center = float(row["center"])
            row_cells = row["cells"]
            assert isinstance(row_cells, list)
            if not row_cells or len(row_cells) > len(column_starts):
                continue
            if id(row) not in core_ids and not minimum_center <= row_center <= maximum_center:
                continue
            if id(row) not in core_ids and len(row_cells) >= len(column_starts):
                continue

            placements: dict[int, dict[str, object]] = {}
            for cell in row_cells:
                cell_start = float(cell["x0"])
                choices = sorted(
                    (abs(start - cell_start), column_index)
                    for column_index, start in enumerate(column_starts)
                    if column_index not in placements
                )
                if not choices or choices[0][0] > alignment_tolerance:
                    placements = {}
                    break
                placements[choices[0][1]] = cell
            if not placements:
                continue
            padded_row = [
                {"text": "", "confidence": 0.0, "x0": 0.0, "x1": 0.0}
                for _ in column_starts
            ]
            for column_index, cell in placements.items():
                padded_row[column_index] = cell
            aligned_rows.append((row_center, padded_row))

        table_rows = [row for _, row in sorted(aligned_rows, key=lambda item: item[0])]
        tables.append({
            "starts": [value / page_width for value in column_starts],
            "rows": table_rows,
        })
    return tables


class _LimitedSeekableStream:
    """Enforce an output ceiling while allowing PDF writers to seek and patch xrefs."""

    def __init__(self, stream: Any, max_bytes: int) -> None:
        self._stream = stream
        self._max_bytes = max_bytes
        self._high_water = 0

    def write(self, data: bytes) -> int:
        position = self._stream.tell()
        end = position + len(data)
        if max(self._high_water, end) > self._max_bytes:
            raise ConversionError("拆分后文件总量超过本机安全处理上限。")
        written = self._stream.write(data)
        self._high_water = max(self._high_water, position + (written or 0))
        return written

    def truncate(self, size: int | None = None) -> int:
        target = self._stream.tell() if size is None else size
        if target > self._max_bytes:
            raise ConversionError("拆分后文件总量超过本机安全处理上限。")
        result = self._stream.truncate(size)
        self._high_water = min(self._high_water, target)
        return result

    def __getattr__(self, name: str) -> Any:
        return getattr(self._stream, name)


def _write_pdf_tables_xlsx(
    source: Path,
    destination: Path,
    progress_callback: Callable[[int], None] | None = None,
) -> None:
    try:
        import pdfplumber
    except ImportError as exc:
        raise ConversionError("当前环境没有安装 PDF 表格提取组件 pdfplumber。") from exc
    try:
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import PatternFill
    except ImportError as exc:
        raise ConversionError("当前环境没有安装 XLSX 支持组件 openpyxl。") from exc

    workbook = None
    try:
        from pypdf import PdfReader

        source_reader = PdfReader(str(source), strict=False)
        try:
            if source_reader.is_encrypted:
                raise ConversionError("PDF 已加密；请先使用原密码解密。")
            page_count = len(source_reader.pages)
        finally:
            source_reader.close()
        del source_reader
        if page_count == 0:
            raise ConversionError("PDF 中没有可提取的页面。")
        if page_count > MAX_PDF_TABLE_PAGES:
            raise ConversionError("PDF 页数超过本机表格提取上限。")

        workbook = Workbook()
        summary = workbook.active
        summary.title = "提取说明"
        summary.append(["项目", "内容"])
        summary.append(["来源文件", source.name])
        summary.cell(row=2, column=2).data_type = "s"
        summary.append(["PDF 页数", page_count])
        summary.append(["表格数量", 0])
        summary.append(["提取方式", "识别页面线条网格与文字对齐表格；有本机 Tesseract 时，对含图片页面补充 OCR 表格识别。"])
        summary.append([
            "适用范围",
            "数字有框表仅在格线和单元格矩形完整匹配时还原单向合并；扫描表格按稳定列位置识别并记录 OCR 置信度，相邻页同列结构会尝试续接。复杂跨页结构及未满足几何条件的合并仍需人工检查。",
        ])

        extracted_tables: list[tuple[int, list[list[str]], list[tuple[int, int, int, int]]]] = []
        ocr_tables: list[dict[str, object]] = []
        page_texts: list[tuple[int, str]] = []
        pages_without_text: list[int] = []
        ocr_page_candidates: list[tuple[int, float, float]] = []
        total_text_bytes = 0
        total_table_text_bytes = 0
        total_ocr_text_bytes = 0
        total_cells = 0
        total_page_chars = 0
        with pdfplumber.open(str(source)) as document:
            if len(document.pages) != page_count:
                raise ConversionError("PDF 页面数量在读取过程中发生变化，请重新添加文件。")
            for page_index, page in enumerate(document.pages, start=1):
                page_chars = page.chars
                page_char_count = len(page_chars)
                total_page_chars += page_char_count
                if page_char_count > MAX_PDF_TABLE_PAGE_CHARS:
                    raise ConversionError("PDF 单页字符数量过大，已停止表格布局分析。")
                if total_page_chars > MAX_PDF_TABLE_TOTAL_CHARS:
                    raise ConversionError("PDF 字符总量过大，已停止表格布局分析。")
                if len(page.edges) > MAX_PDF_TABLE_PAGE_EDGES:
                    raise ConversionError("PDF 单页线条数量过大，已停止表格布局分析。")
                text = page.extract_text() or ""
                text = text.strip()
                if text:
                    total_text_bytes += len(text.encode("utf-8"))
                    if total_text_bytes > MAX_EXTRACTED_TEXT_BYTES:
                        raise ConversionError("PDF 提取出的文字超过本机 32 MiB 安全上限。")
                    page_texts.append((page_index, text))
                else:
                    pages_without_text.append(page_index)
                has_page_images = bool(page.images)

                candidates: list[
                    tuple[
                        tuple[float, float, float, float],
                        list[list[str]],
                        list[tuple[int, int, int, int]],
                    ]
                ] = []
                line_tables = page.find_tables()
                line_boxes: list[tuple[float, float, float, float]] = []
                for table in line_tables:
                    rows = _normalize_pdf_table(table.extract(), preserve_grid=True)
                    if not rows or not any(value for row in rows for value in row):
                        continue
                    bbox = tuple(float(value) for value in table.bbox)
                    line_boxes.append(bbox)
                    merge_ranges = _pdf_table_merge_ranges(table, rows, page.edges)
                    candidates.append((bbox, rows, merge_ranges))

                text_settings = {
                    "vertical_strategy": "text",
                    "horizontal_strategy": "text",
                    "min_words_vertical": 2,
                    "min_words_horizontal": 1,
                }
                text_tables = page.find_tables(table_settings=text_settings)
                for text_table in text_tables:
                    text_bbox = tuple(float(value) for value in text_table.bbox)
                    overlapping_lines = [
                        bbox for bbox in line_boxes
                        if _pdf_bboxes_intersect(text_bbox, bbox)
                    ]
                    regions = [text_bbox]
                    for line_bbox in overlapping_lines:
                        regions = [
                            remainder
                            for region in regions
                            for remainder in _subtract_pdf_bbox(region, line_bbox)
                        ]
                    if not overlapping_lines:
                        regions = [text_bbox]

                    for region in regions:
                        if region == text_bbox:
                            tables_to_check = [text_table]
                        else:
                            cropped_page = page.crop(region, strict=False)
                            tables_to_check = cropped_page.find_tables(table_settings=text_settings)
                        for table in tables_to_check:
                            rows = _normalize_pdf_table(table.extract())
                            if not rows:
                                continue
                            bbox = tuple(float(value) for value in table.bbox)
                            if any(_table_bbox_is_duplicate(bbox, existing) for existing, _, _ in candidates):
                                continue
                            candidates.append((bbox, rows, []))

                for _, table_rows, merge_ranges in candidates:
                    if any(len(row) > 16_384 for row in table_rows):
                        raise ConversionError("PDF 表格列数超过 Excel 的工作表上限。")
                    total_cells += sum(len(row) for row in table_rows)
                    if total_cells > MAX_PDF_TABLE_CELLS:
                        raise ConversionError("PDF 表格单元格总数超过本机处理上限。")
                    if len(extracted_tables) >= MAX_PDF_TABLES:
                        raise ConversionError("PDF 表格数量超过本机处理上限。")
                    if any(len(value) > 32_767 for row in table_rows for value in row):
                        raise ConversionError("PDF 中有单元格文字超过 Excel 的长度上限。")
                    total_table_text_bytes += sum(
                        len(value.encode("utf-8")) for row in table_rows for value in row
                    )
                    if total_table_text_bytes > MAX_PDF_TABLE_TEXT_BYTES:
                        raise ConversionError("PDF 表格文字总量超过本机 64 MiB 安全上限。")
                    extracted_tables.append((page_index, table_rows, merge_ranges))
                if has_page_images and (not text or not candidates):
                    ocr_page_candidates.append((page_index - 1, float(page.width), float(page.height)))
                page.close()
                if progress_callback is not None:
                    progress_callback(round(page_index * 25 / page_count))

        ocr_engine = tesseract_engine() if ocr_page_candidates else None
        if ocr_page_candidates and ocr_engine is not None:
            if len(ocr_page_candidates) > MAX_PDF_TABLE_OCR_PAGES:
                raise ConversionError("需要 OCR 的扫描页数量超过本机 100 页上限。")
            pdf_document = QPdfDocument()
            pdf_document.load(str(source))
            if pdf_document.status() != QPdfDocument.Status.Ready or pdf_document.pageCount() != page_count:
                pdf_document.close()
                raise ConversionError("Qt PDF 引擎无法读取扫描页面。")
            dimensions: dict[int, QSize] = {}
            total_pixels = 0
            for page_index, _page_width, _page_height in ocr_page_candidates:
                point_size = pdf_document.pagePointSize(page_index)
                width = max(1, math.ceil(point_size.width() * OCR_DPI / 72))
                height = max(1, math.ceil(point_size.height() * OCR_DPI / 72))
                pixels = width * height
                if pixels > MAX_OCR_PAGE_PIXELS or total_pixels + pixels > MAX_OCR_TOTAL_PIXELS:
                    pdf_document.close()
                    raise ConversionError("扫描页面的 OCR 总像素数超过本机安全上限。")
                dimensions[page_index] = QSize(width, height)
                total_pixels += pixels
            deadline = time.monotonic() + MAX_OCR_DOCUMENT_SECONDS
            try:
                with tempfile.TemporaryDirectory(prefix="fluke-pdf-table-ocr-") as temp_name:
                    temp_dir = Path(temp_name)
                    for scan_index, (page_index, _page_width, _page_height) in enumerate(ocr_page_candidates, start=1):
                        image = pdf_document.render(page_index, dimensions[page_index])
                        if image.isNull():
                            raise ConversionError("无法渲染 PDF 页面供 OCR。")
                        image_path = temp_dir / f"scan-{scan_index:04d}.png"
                        if not image.save(str(image_path), "PNG"):
                            raise ConversionError("无法暂存 PDF 扫描页图片。")
                        del image
                        timeout = min(MAX_OCR_PAGE_SECONDS, deadline - time.monotonic())
                        if timeout <= 0:
                            raise ConversionError("PDF 表格 OCR 总处理时间超过本机上限。")
                        try:
                            raw_lines = recognize_tesseract_tsv(ocr_engine, image_path, timeout)
                        except RuntimeError as exc:
                            raise ConversionError(str(exc)) from None
                        page_text = "\n".join(
                            str(line.get("text", "")).strip()
                            for line in raw_lines
                            if str(line.get("text", "")).strip()
                        ).strip()
                        if page_text:
                            page_text_bytes = len(page_text.encode("utf-8"))
                            total_ocr_text_bytes += page_text_bytes
                            if total_ocr_text_bytes > MAX_PDF_TABLE_OCR_TEXT_BYTES:
                                raise ConversionError("OCR 提取出的文字超过本机 32 MiB 安全上限。")
                            if (page_index + 1) not in {item[0] for item in page_texts}:
                                page_texts.append((page_index + 1, page_text))
                        page_tables = _ocr_page_tables(raw_lines, dimensions[page_index].width())
                        for table in page_tables:
                            table["page"] = page_index + 1
                            ocr_tables.append(table)
                        if progress_callback is not None:
                            progress_callback(25 + round(scan_index * 50 / len(ocr_page_candidates)))
            finally:
                pdf_document.close()

        # Join only adjacent-page OCR tables whose normalized column starts match.
        merged_ocr_tables: list[dict[str, object]] = []
        for table in sorted(ocr_tables, key=lambda item: int(item["page"])):
            current_starts = table["starts"]
            assert isinstance(current_starts, list)
            continuations = []
            for previous in merged_ocr_tables:
                previous_starts = previous["starts"]
                assert isinstance(previous_starts, list)
                same_columns = len(previous_starts) == len(current_starts) and all(
                    abs(float(left) - float(right)) <= 0.025
                    for left, right in zip(previous_starts, current_starts)
                )
                if same_columns and int(table["page"]) == int(previous["last_page"]) + 1:
                    continuations.append(previous)
            if len(continuations) == 1:
                previous = continuations[0]
                previous_rows = previous["rows"]
                current_rows = table["rows"]
                assert isinstance(previous_rows, list) and isinstance(current_rows, list)
                current_first = current_rows[0] if current_rows else []
                previous_first = previous_rows[0] if previous_rows else []
                repeated_header = bool(current_first and previous_first) and [
                    str(cell["text"]).casefold() for cell in current_first
                ] == [str(cell["text"]).casefold() for cell in previous_first]
                previous_rows.extend(current_rows[1:] if repeated_header else current_rows)
                previous["last_page"] = int(table["page"])
                previous_pages = previous["pages"]
                assert isinstance(previous_pages, list)
                previous_pages.append(int(table["page"]))
                continue
            table["last_page"] = int(table["page"])
            table["pages"] = [int(table["page"])]
            merged_ocr_tables.append(table)

        if len(extracted_tables) + len(merged_ocr_tables) > MAX_PDF_TABLES:
            raise ConversionError("PDF 表格数量超过本机处理上限。")

        summary.cell(row=4, column=2, value=len(extracted_tables) + len(merged_ocr_tables))
        summary.append(["无内嵌文字的页面", "、".join(map(str, pages_without_text)) or "无"])
        summary.append(["OCR 表格数量", len(merged_ocr_tables)])
        summary.append([
            "OCR 扫描页",
            "、".join(str(page + 1) for page, _, _ in ocr_page_candidates) or "无",
        ])
        if ocr_page_candidates and ocr_engine is None:
            summary.append(["扫描页说明", "未检测到可用的 Tesseract 与语言模型，扫描页未执行 OCR。"])
        elif ocr_page_candidates and not merged_ocr_tables:
            summary.append(["扫描页说明", "OCR 已完成，但未发现列位置稳定的表格；识别正文已保存在页面正文工作表。"])
        if not extracted_tables and not merged_ocr_tables and not page_texts:
            raise ConversionError("未检测到 PDF 内嵌文字或可识别表格；扫描页需要可用的 Tesseract OCR 引擎。")

        for table_index, (page_index, rows, merge_ranges) in enumerate(extracted_tables, start=1):
            sheet = workbook.create_sheet(f"表{table_index:04d}_P{page_index:04d}")
            for row_index, values in enumerate(rows, start=1):
                for column_index, value in enumerate(values, start=1):
                    cell = sheet.cell(row=row_index, column=column_index, value=value)
                    # PDF content is untrusted; keep formula-looking text inert in Excel.
                    cell.data_type = "s"
            for row_start, column_start, row_end, column_end in merge_ranges:
                sheet.merge_cells(
                    start_row=row_start,
                    start_column=column_start,
                    end_row=row_end,
                    end_column=column_end,
                )

        ocr_comment_cells = 0
        low_confidence_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
        for table_index, table in enumerate(merged_ocr_tables, start=1):
            pages = table["pages"]
            rows = table["rows"]
            assert isinstance(pages, list) and isinstance(rows, list)
            page_label = "、".join(str(page) for page in pages)
            sheet = workbook.create_sheet(f"OCR表{table_index:04d}_P{pages[0]:04d}")
            for row_index, row_values in enumerate(rows, start=1):
                if not isinstance(row_values, list):
                    continue
                for column_index, item in enumerate(row_values, start=1):
                    if not isinstance(item, dict):
                        continue
                    text_value = str(item["text"])
                    if not text_value:
                        continue
                    if len(text_value) > 32_767:
                        raise ConversionError("PDF 中有单元格文字超过 Excel 的长度上限。")
                    total_table_text_bytes += len(text_value.encode("utf-8"))
                    if total_table_text_bytes > MAX_PDF_TABLE_TEXT_BYTES:
                        raise ConversionError("PDF 表格文字总量超过本机 64 MiB 安全上限。")
                    cell = sheet.cell(row=row_index, column=column_index, value=text_value)
                    cell.data_type = "s"
                    confidence = max(0.0, min(100.0, float(item["confidence"])))
                    if ocr_comment_cells >= MAX_PDF_TABLE_OCR_CELLS:
                        raise ConversionError("扫描 PDF 表格单元格超过本机 OCR 注释上限。")
                    cell.comment = Comment(
                        f"本机 OCR 置信度：{confidence:.0f}%\n扫描页：{page_label}",
                        "FLUKE",
                    )
                    if confidence < 60:
                        cell.fill = low_confidence_fill
                    ocr_comment_cells += 1
                    total_cells += 1
                    if total_cells > MAX_PDF_TABLE_CELLS:
                        raise ConversionError("PDF 表格单元格总数超过本机处理上限。")

        if page_texts:
            body = workbook.create_sheet("页面正文")
            body.append(["页码", "分段", "正文"])
            for page_index, text in page_texts:
                for segment_index, start in enumerate(range(0, len(text), 30_000), start=1):
                    cell_text = text[start:start + 30_000]
                    total_cells += 1
                    if total_cells > MAX_PDF_TABLE_CELLS:
                        raise ConversionError("PDF 表格与正文总量超过本机处理上限。")
                    row_index = body.max_row + 1
                    body.cell(row=row_index, column=1, value=page_index)
                    body.cell(row=row_index, column=2, value=segment_index)
                    body_cell = body.cell(row=row_index, column=3, value=cell_text)
                    body_cell.data_type = "s"
        if progress_callback is not None:
            progress_callback(95)

        workbook.save(destination)
        if destination.stat().st_size > MAX_PDF_TABLE_XLSX_BYTES:
            raise ConversionError("XLSX 输出超过本机 128 MiB 安全上限。")
        if progress_callback is not None:
            progress_callback(100)
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"PDF 表格提取失败：{exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _write_pdf_pages_zip(source: Path, destination: Path, pages_per_file: int = 1) -> None:
    try:
        from pypdf import PdfReader, PdfWriter
    except ImportError as exc:
        raise ConversionError("当前环境没有安装 PDF 文本支持组件 pypdf。") from exc
    try:
        if isinstance(pages_per_file, bool) or not isinstance(pages_per_file, int):
            raise ConversionError("每个拆分 PDF 的页数必须是整数。")
        if not 1 <= pages_per_file <= MAX_PDF_PAGES:
            raise ConversionError("每个拆分 PDF 的页数必须在 1 到本机上限之间。")
        members: list[tuple[str, int]] = []
        total_bytes = 0
        with PdfReader(str(source), strict=False) as reader:
            if reader.is_encrypted:
                raise ConversionError("PDF 已加密，当前版本不会尝试密码破解或解锁。")
            page_count = len(reader.pages)
            if page_count == 0:
                raise ConversionError("PDF 中没有可拆分的页面。")
            if page_count > MAX_PDF_PAGES:
                raise ConversionError("PDF 页数超过本机拆分上限。")
            with tempfile.TemporaryDirectory(prefix="fluke-pdf-split-") as temp_dir:
                with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_STORED) as archive:
                    for start in range(0, page_count, pages_per_file):
                        end = min(page_count, start + pages_per_file)
                        writer = PdfWriter()
                        for index in range(start, end):
                            writer.add_page(reader.pages[index])
                        temp_pdf = Path(temp_dir) / "group.pdf"
                        remaining_bytes = MAX_ARCHIVE_EXPANDED_BYTES - total_bytes
                        with temp_pdf.open("w+b") as stream:
                            writer.write(_LimitedSeekableStream(stream, remaining_bytes))
                            stream.flush()
                            group_size = stream.seek(0, os.SEEK_END)
                            if total_bytes + group_size > MAX_ARCHIVE_EXPANDED_BYTES:
                                raise ConversionError("拆分后文件总量超过本机安全处理上限。")
                            stream.seek(0)
                            if pages_per_file == 1:
                                name = f"page_{start + 1:04d}.pdf"
                            else:
                                name = f"pages_{start + 1:04d}-{end:04d}.pdf"
                            with archive.open(name, "w") as member:
                                shutil.copyfileobj(stream, member, length=1024 * 1024)
                        del writer
                        total_bytes += group_size
                        members.append((name, end - start))
        with zipfile.ZipFile(destination) as archive:
            if archive.testzip() is not None:
                raise ConversionError("拆分后的 ZIP 完整性检查未通过。")
            with tempfile.TemporaryDirectory(prefix="fluke-pdf-verify-") as temp_dir:
                for index, (name, expected_pages) in enumerate(members):
                    temp_pdf = Path(temp_dir) / f"verify-{index:04d}.pdf"
                    with archive.open(name) as member, temp_pdf.open("wb") as output:
                        shutil.copyfileobj(member, output, length=1024 * 1024)
                    with PdfReader(str(temp_pdf), strict=False) as page_pdf:
                        if page_pdf.is_encrypted or len(page_pdf.pages) != expected_pages:
                            raise ConversionError("拆分结果中的 PDF 页数校验失败。")
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError("PDF 文件损坏或页面结构不受支持，无法按页拆分。") from exc


def _write_pdf_images_zip(source: Path, destination: Path) -> None:
    document = QPdfDocument()
    try:
        document.load(str(source))
        if document.status() != QPdfDocument.Status.Ready:
            raise ConversionError("Qt PDF 引擎无法读取这个文件。")
        page_count = document.pageCount()
        if page_count < 1:
            raise ConversionError("PDF 中没有可导出的页面。")
        if page_count > MAX_PDF_PAGES:
            raise ConversionError("PDF 页数超过本机页面图片导出上限。")

        dimensions: list[QSize] = []
        total_pixels = 0
        for page_index in range(page_count):
            point_size = document.pagePointSize(page_index)
            if point_size.width() <= 0 or point_size.height() <= 0:
                raise ConversionError("PDF 页面尺寸无效，无法导出页面图片。")
            width = max(1, math.ceil(point_size.width() * PDF_PAGE_IMAGE_DPI / 72))
            height = max(1, math.ceil(point_size.height() * PDF_PAGE_IMAGE_DPI / 72))
            pixels = width * height
            if pixels > MAX_PDF_IMAGE_PAGE_PIXELS:
                raise ConversionError("PDF 页面尺寸超过本机页面图片像素上限。")
            if total_pixels + pixels > MAX_PDF_IMAGE_TOTAL_PIXELS:
                raise ConversionError("PDF 页面图片总像素数超过本机处理上限。")
            dimensions.append(QSize(width, height))
            total_pixels += pixels

        names: list[str] = []
        total_bytes = 0
        with tempfile.TemporaryDirectory(prefix="fluke-pdf-images-") as temp_dir:
            image_path = Path(temp_dir) / "page.png"
            with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_STORED) as archive:
                for page_index, size in enumerate(dimensions):
                    image = document.render(page_index, size)
                    if image.isNull():
                        raise ConversionError("无法渲染 PDF 页面。")
                    if not image.save(str(image_path), "PNG"):
                        raise ConversionError("无法写入 PDF 页面图片。")
                    del image
                    page_size = image_path.stat().st_size
                    if total_bytes + page_size > MAX_ARCHIVE_EXPANDED_BYTES:
                        raise ConversionError("页面图片总量超过本机安全处理上限。")
                    name = f"page_{page_index + 1:04d}.png"
                    archive.write(image_path, arcname=name)
                    total_bytes += page_size
                    names.append(name)

        with zipfile.ZipFile(destination) as archive:
            if archive.testzip() is not None or archive.namelist() != names:
                raise ConversionError("页面图片 ZIP 完整性检查未通过。")
            for name in names:
                with archive.open(name) as image_file:
                    if image_file.read(8) != b"\x89PNG\r\n\x1a\n":
                        raise ConversionError("页面图片输出校验失败。")
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError("PDF 文件损坏或页面结构不受支持，无法导出页面图片。") from exc
    finally:
        document.close()


def _write_pdf_security(
    source: Path, destination: Path, operation: str, password: str
) -> None:
    """Encrypt or decrypt a PDF without exposing the supplied password in errors."""
    try:
        from pypdf import PdfReader, PdfWriter
    except ImportError as exc:
        raise ConversionError("当前环境没有安装 PDF 文本支持组件 pypdf。") from exc
    if not password:
        raise ConversionError("请输入 PDF 密码。")
    if len(password) > MAX_PDF_PASSWORD_LENGTH:
        raise ConversionError("PDF 密码长度超过本机输入上限。")
    if operation == "pdf-encrypt" and len(password) < MIN_PDF_PASSWORD_LENGTH:
        raise ConversionError("PDF 加密密码至少需要 8 个字符。")
    try:
        with PdfReader(str(source), strict=False) as reader:
            if operation == "pdf-encrypt" and reader.is_encrypted:
                raise ConversionError("这个 PDF 已加密，请先使用原密码解密。")
            if operation == "pdf-decrypt":
                if not reader.is_encrypted:
                    raise ConversionError("这个 PDF 没有加密，无需解密。")
                if not reader.decrypt(password):
                    raise ConversionError("PDF 密码不正确，无法解密。")
            page_count = len(reader.pages)
            if not page_count:
                raise ConversionError("PDF 中没有可处理的页面。")
            if page_count > MAX_PDF_PAGES:
                raise ConversionError("PDF 页数超过本机安全处理上限。")
            writer = PdfWriter()
            writer.clone_document_from_reader(reader)
            if operation == "pdf-encrypt":
                writer.encrypt(user_password=password, algorithm="AES-256")
            with destination.open("wb") as stream:
                writer.write(stream)
        if destination.stat().st_size > MAX_ARCHIVE_EXPANDED_BYTES:
            raise ConversionError("PDF 输出超过本机安全处理上限。")

        with PdfReader(str(destination), strict=False) as verified:
            if operation == "pdf-encrypt":
                if not verified.is_encrypted or not verified.decrypt(password):
                    raise ConversionError("加密后的 PDF 完整性检查未通过。")
            elif verified.is_encrypted:
                raise ConversionError("解密后的 PDF 完整性检查未通过。")
            if len(verified.pages) != page_count:
                raise ConversionError("PDF 页面完整性检查未通过。")
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError("PDF 文件损坏或安全处理失败，未能完成加密/解密。") from None


def _write_pdf_merge(sources: list[Path], destination: Path) -> None:
    try:
        from pypdf import PdfReader, PdfWriter
    except ImportError as exc:
        raise ConversionError("当前环境没有安装 PDF 文本支持组件 pypdf。") from exc
    writer = PdfWriter()
    expected_pages = 0
    try:
        with ExitStack() as stack:
            for source in sources:
                stream = stack.enter_context(source.open("rb"))
                reader = PdfReader(stream, strict=False)
                if reader.is_encrypted:
                    raise ConversionError("合并列表中含有加密 PDF，请先使用原密码解密。")
                expected_pages += len(reader.pages)
                if expected_pages > MAX_PDF_PAGES:
                    raise ConversionError("合并后的 PDF 页数超过本机处理上限。")
                writer.append(reader)
            if not expected_pages:
                raise ConversionError("合并列表中的 PDF 没有可用页面。")
            with destination.open("wb") as stream:
                writer.write(stream)
        if destination.stat().st_size > MAX_MERGED_PDF_BYTES:
            raise ConversionError("合并后的 PDF 超过本机 512 MiB 输出上限。")
        with PdfReader(str(destination), strict=False) as check_reader:
            if check_reader.is_encrypted or len(check_reader.pages) != expected_pages:
                raise ConversionError("合并后的 PDF 页面完整性检查未通过。")
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError("PDF 文件损坏或页面结构不受支持，无法合并。") from None


def merge_pdf_files(
    sources: list[str | Path], output_directory: str | Path | None = None
) -> Path:
    source_paths = [Path(source).expanduser().resolve() for source in sources]
    if len(source_paths) < 2:
        raise ConversionError("至少选择两个未加密 PDF 才能合并。")
    if len({str(path).casefold() for path in source_paths}) != len(source_paths):
        raise ConversionError("合并列表中存在重复 PDF 文件。")
    for path in source_paths:
        _check_source(path, MAX_DOCUMENT_BYTES)
        if path.suffix.lower() != ".pdf" or not _looks_like_pdf(path):
            raise ConversionError("合并列表中有无效 PDF 文件。")
        if _pdf_encryption_state_for_path(path) is not False:
            raise ConversionError("加密或无法读取的 PDF 不能直接合并。")
    destination = _unique_destination(source_paths[0], "pdf", output_directory)
    temporary: Path | None = None
    try:
        temporary = _temporary_path(destination)
        _write_pdf_merge(source_paths, temporary)
        os.replace(temporary, destination)
    except ConversionError:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        destination.unlink(missing_ok=True)
        raise
    except Exception as exc:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        destination.unlink(missing_ok=True)
        raise ConversionError("合并 PDF 失败，未保留部分输出。") from None
    return destination


def _document_content(path: Path, extension: str) -> tuple[str, list[tuple[str, Any]] | None]:
    if extension in TEXT_EXTENSIONS:
        text = _read_text_file(path)
        if extension in ("html", "htm"):
            parser = _HTMLTextExtractor()
            parser.feed(text)
            text = parser.text()
        return text, None
    if extension == "docx":
        blocks = _docx_blocks(path)
        lines: list[str] = []
        for kind, content in blocks:
            if kind == "table":
                lines.extend("\t".join(row) for row in content)
            elif kind == "heading":
                _level, value = content
                lines.append(value)
            else:
                lines.append(content)
        return "\n\n".join(lines), blocks
    if extension == "pdf":
        return _pdf_text(path), None
    if extension == "epub":
        _title, blocks = _epub_blocks(path)
        lines: list[str] = []
        for kind, content in blocks:
            if kind == "heading":
                _level, value = content
                lines.append(value)
            else:
                lines.append(content)
        return "\n\n".join(lines), blocks
    raise ConversionError("不适用于文档转换的输入格式。")


def _markdown_html(text: str) -> str:
    try:
        from markdown_it import MarkdownIt
    except ImportError as exc:
        raise ConversionError("当前环境没有安装 Markdown 结构解析组件 markdown-it-py。") from exc

    parser = MarkdownIt("commonmark", {"html": False})
    parser.enable(("table", "strikethrough"))

    def render_image_alt(_renderer: Any, tokens: list[Any], index: int, _options: Any, _env: Any) -> str:
        # Conversion stays local. A remote Markdown image is represented by its
        # alt text instead of being fetched or embedded as an external resource.
        token = tokens[index]
        return escape(str(token.content or token.attrGet("alt") or ""))

    parser.add_render_rule("image", render_image_alt)
    return parser.render(text)


def _add_docx_run(
    paragraph: Any,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
    code: bool = False,
    strike: bool = False,
    href: str = "",
) -> None:
    if not text:
        return
    if href and urlsplit(href).scheme.lower() in {"http", "https", "mailto"}:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.opc.constants import RELATIONSHIP_TYPE as RT

        relationship_id = paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship_id)
        run_element = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        if bold:
            properties.append(OxmlElement("w:b"))
        if italic:
            properties.append(OxmlElement("w:i"))
        if strike:
            properties.append(OxmlElement("w:strike"))
        if code:
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), "Consolas")
            fonts.set(qn("w:hAnsi"), "Consolas")
            properties.append(fonts)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        properties.append(color)
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        properties.append(underline)
        run_element.append(properties)
        text_element = OxmlElement("w:t")
        text_element.text = text
        if text[:1].isspace() or text[-1:].isspace():
            text_element.set(qn("xml:space"), "preserve")
        run_element.append(text_element)
        hyperlink.append(run_element)
        paragraph._p.append(hyperlink)
        return

    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.strike = strike
    if code:
        run.font.name = "Consolas"


def _append_markdown_inline(
    node: Any,
    paragraph: Any,
    *,
    bold: bool = False,
    italic: bool = False,
    code: bool = False,
    strike: bool = False,
    href: str = "",
) -> None:
    if node.text:
        _add_docx_run(paragraph, node.text, bold=bold, italic=italic, code=code, strike=strike, href=href)
    for child in node:
        tag = str(child.tag).lower() if isinstance(child.tag, str) else ""
        if tag == "br":
            paragraph.add_run().add_break()
        else:
            _append_markdown_inline(
                child,
                paragraph,
                bold=bold or tag in {"strong", "b"},
                italic=italic or tag in {"em", "i"},
                code=code or tag == "code",
                strike=strike or tag in {"del", "s"},
                href=str(child.get("href", "")) if tag == "a" else href,
            )
        if child.tail:
            _add_docx_run(paragraph, child.tail, bold=bold, italic=italic, code=code, strike=strike, href=href)


def _add_docx_paragraph(
    document: Any,
    node: Any,
    *,
    style: str | None = None,
    indent_level: int = 0,
) -> Any:
    from docx.shared import Pt

    paragraph = document.add_paragraph(style=style)
    if indent_level:
        paragraph.paragraph_format.left_indent = Pt(18 * indent_level)
    _append_markdown_inline(node, paragraph)
    return paragraph


def _append_lxml_text(node: Any, text: str) -> None:
    if not text:
        return
    if len(node):
        node[-1].tail = (node[-1].tail or "") + text
    else:
        node.text = (node.text or "") + text


def _add_docx_markdown_list(node: Any, document: Any, depth: int = 0) -> None:
    from copy import deepcopy
    from lxml import html as lxml_html

    tag = str(node.tag).lower()
    paragraph_style = "List Number" if tag == "ol" else "List Bullet"
    for item in node:
        if str(item.tag).lower() != "li":
            continue
        pending = lxml_html.Element("span")
        pending.text = item.text

        def flush_inline() -> None:
            nonlocal pending
            if (pending.text and pending.text.strip()) or len(pending):
                _add_docx_paragraph(document, pending, style=paragraph_style, indent_level=depth)
            pending = lxml_html.Element("span")

        for child in item:
            child_tag = str(child.tag).lower() if isinstance(child.tag, str) else ""
            child_tail = child.tail or ""
            if child_tag in {"ul", "ol"}:
                flush_inline()
                _add_docx_markdown_list(child, document, depth + 1)
            elif child_tag in {"p", "blockquote", "pre", "table", "hr"}:
                flush_inline()
                if child_tag == "p":
                    _add_docx_paragraph(document, child, style=paragraph_style, indent_level=depth)
                else:
                    _add_docx_markdown_blocks(child, document)
            else:
                cloned = deepcopy(child)
                cloned.tail = None
                pending.append(cloned)
            _append_lxml_text(pending, child_tail)
        flush_inline()


def _add_docx_markdown_table(node: Any, document: Any) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    rows = [row for row in node.xpath(".//tr")]
    cells_by_row = [row.xpath("./th | ./td") for row in rows]
    columns = max((len(row) for row in cells_by_row), default=0)
    if not columns:
        return
    table = document.add_table(rows=len(cells_by_row), cols=columns)
    table.style = "Table Grid"
    for row_index, row_cells in enumerate(cells_by_row):
        for column_index, source_cell in enumerate(row_cells):
            target_cell = table.cell(row_index, column_index)
            paragraph = target_cell.paragraphs[0]
            paragraph.clear()
            _append_markdown_inline(source_cell, paragraph, bold=str(source_cell.tag).lower() == "th")
            align = source_cell.get("style", "").lower()
            if "text-align:right" in align:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif "text-align:center" in align:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_docx_markdown_blocks(node: Any, document: Any, quote_depth: int = 0) -> None:
    from docx.shared import Pt

    for child in node:
        tag = str(child.tag).lower() if isinstance(child.tag, str) else ""
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            level = int(tag[1])
            paragraph = document.add_heading(level=level)
            _append_markdown_inline(child, paragraph)
        elif tag == "p":
            _add_docx_paragraph(document, child, style="Quote" if quote_depth else None)
        elif tag in {"ul", "ol"}:
            _add_docx_markdown_list(child, document)
        elif tag == "blockquote":
            _add_docx_markdown_blocks(child, document, quote_depth + 1)
        elif tag == "table":
            _add_docx_markdown_table(child, document)
        elif tag == "pre":
            code_text = "".join(child.itertext())
            if code_text.endswith("\n"):
                code_text = code_text[:-1]
            for line in code_text.split("\n") or [""]:
                paragraph = document.add_paragraph(style="No Spacing")
                run = paragraph.add_run(line)
                run.font.name = "Consolas"
                run.font.size = Pt(9)
        elif tag == "hr":
            document.add_paragraph("────────────────")
        elif tag in {"div", "section", "article"}:
            _add_docx_markdown_blocks(child, document, quote_depth)


def _write_markdown_docx(path: Path, text: str) -> None:
    try:
        from docx import Document
        from lxml import html as lxml_html
    except ImportError as exc:
        raise ConversionError("当前环境缺少 Markdown 或 DOCX 文档支持组件。") from exc
    document = Document()
    rendered = _markdown_html(text)
    root = lxml_html.fragment_fromstring(rendered, create_parent="div")
    _add_docx_markdown_blocks(root, document)
    document.save(path)


def _docx_html(blocks: list[tuple[str, Any]]) -> str:
    parts: list[str] = []
    for kind, content in blocks:
        if kind == "table":
            rows = ["<tr>" + "".join(f"<td>{escape(str(cell))}</td>" for cell in row) + "</tr>" for row in content]
            parts.append("<table>" + "".join(rows) + "</table>")
        elif kind == "heading":
            level, value = content
            parts.append(f"<h{min(int(level), 6)}>{escape(value)}</h{min(int(level), 6)}>")
        else:
            parts.append("<p>" + escape(content) + "</p>")
    return "\n".join(parts)


def _docx_markdown(blocks: list[tuple[str, Any]]) -> str:
    parts: list[str] = []
    for kind, content in blocks:
        if kind == "table":
            rows = [[str(cell).replace("|", "\\|").replace("\n", " ") for cell in row] for row in content]
            if rows:
                parts.append("| " + " | ".join(rows[0]) + " |")
                parts.append("| " + " | ".join("---" for _ in rows[0]) + " |")
                parts.extend("| " + " | ".join(row) + " |" for row in rows[1:])
        elif kind == "heading":
            level, value = content
            parts.append("#" * min(int(level), 6) + " " + value)
        else:
            parts.append(content)
    return "\n\n".join(parts)


def _table_html(headers: list[str], rows: list[list[Any]]) -> str:
    heading = "<tr>" + "".join(f"<th>{escape(str(value))}</th>" for value in headers) + "</tr>"
    body = "".join(
        "<tr>" + "".join(f"<td>{escape(str(value))}</td>" for value in row) + "</tr>"
        for row in rows
    )
    return "<table><thead>" + heading + "</thead><tbody>" + body + "</tbody></table>"


def _write_table_docx(path: Path, headers: list[str], rows: list[list[Any]]) -> None:
    try:
        from docx import Document
    except ImportError as exc:
        raise ConversionError("当前环境没有安装 DOCX 支持组件 python-docx。") from exc
    document = Document()
    if not headers:
        document.add_paragraph("没有可导出的表格内容。")
    else:
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        for index, value in enumerate(headers):
            table.rows[0].cells[index].text = str(value)
        for row in rows:
            cells = table.add_row().cells
            for index, value in enumerate(row[:len(headers)]):
                cells[index].text = str(value)
    document.save(path)


def _wrap_html(body: str) -> str:
    return (
        "<!doctype html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\">"
        "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">"
        "<title>Converted document</title></head><body>" + body + "</body></html>"
    )


def _epub_output_chapters(
    source_extension: str,
    text: str,
    blocks: list[tuple[str, Any]] | None,
    book_title: str,
) -> list[tuple[str, str]]:
    if blocks is not None:
        grouped: list[tuple[str, list[tuple[str, Any]]]] = []
        current_title = book_title
        current_blocks: list[tuple[str, Any]] = []
        for block in blocks:
            kind, content = block
            if kind == "heading" and str(content[0]) == "1":
                if current_blocks:
                    grouped.append((current_title, current_blocks))
                current_title = str(content[1])
                current_blocks = [block]
            else:
                current_blocks.append(block)
        if current_blocks:
            grouped.append((current_title, current_blocks))
        if grouped:
            return [(title, _docx_html(items)) for title, items in grouped]

    if source_extension in ("md", "markdown"):
        grouped_markdown: list[tuple[str, list[str]]] = []
        current_title = book_title
        current_lines: list[str] = []
        for line in text.splitlines():
            heading = re.match(r"^#\s+(.+)$", line)
            if heading:
                if current_lines:
                    grouped_markdown.append((current_title, current_lines))
                current_title = heading.group(1).strip()
                current_lines = [line]
            else:
                current_lines.append(line)
        if current_lines:
            grouped_markdown.append((current_title, current_lines))
        if grouped_markdown:
            return [
                (title, _markdown_html("\n".join(lines)))
                for title, lines in grouped_markdown
                if any(line.strip() for line in lines)
            ]

    if source_extension in ("html", "htm"):
        text_parser = _HTMLTextExtractor()
        text_parser.feed(text)
        text = text_parser.text()
    return [(book_title, "<pre>" + escape(text) + "</pre>")]


def _write_epub(path: Path, title: str, chapters: list[tuple[str, str]]) -> None:
    clean_title = (title.strip() or "Untitled")[:255]
    nonempty_chapters: list[tuple[str, str]] = []
    for index, (name, body) in enumerate(chapters, start=1):
        text_parser = _HTMLTextExtractor()
        text_parser.feed(body)
        if text_parser.text():
            nonempty_chapters.append((name.strip() or f"第 {index} 章", body))
    chapters = nonempty_chapters
    if not chapters:
        raise ConversionError("源文件没有可写入 EPUB 的正文内容。")
    if len(chapters) > MAX_PDF_PAGES:
        raise ConversionError("电子书章节数量超过本机输出上限。")
    if sum(len(body.encode("utf-8")) for _name, body in chapters) > MAX_EXTRACTED_TEXT_BYTES:
        raise ConversionError("电子书正文超过本机 EPUB 输出上限。")
    safe_title = escape(clean_title, quote=True)
    book_id = f"urn:uuid:{uuid.uuid4()}"
    modified = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    chapter_items: list[tuple[str, str, str]] = []
    nav_items: list[str] = []
    spine_items: list[str] = []
    for index, (chapter_title, body_html) in enumerate(chapters, start=1):
        filename = f"chapter-{index:04d}.xhtml"
        item_id = f"chapter-{index:04d}"
        safe_chapter_title = escape(chapter_title[:255], quote=True)
        chapter = (
            '<?xml version="1.0" encoding="utf-8"?>'
            '<html xmlns="http://www.w3.org/1999/xhtml" xml:lang="zh-CN" lang="zh-CN">'
            "<head><title>" + safe_chapter_title + "</title><meta charset=\"utf-8\"/></head>"
            "<body>" + body_html + "</body></html>"
        )
        chapter_items.append((item_id, filename, chapter))
        nav_items.append(f'<li><a href="{filename}">{safe_chapter_title}</a></li>')
        spine_items.append(f'<itemref idref="{item_id}"/>')
    navigation = (
        '<?xml version="1.0" encoding="utf-8"?>'
        '<html xmlns="http://www.w3.org/1999/xhtml" '
        'xmlns:epub="http://www.idpf.org/2007/ops" xml:lang="zh-CN" lang="zh-CN">'
        "<head><title>" + safe_title + "</title></head><body>"
        '<nav epub:type="toc" id="toc"><h1>目录</h1><ol>'
        + "".join(nav_items)
        + "</ol></nav></body></html>"
    )
    package = (
        '<?xml version="1.0" encoding="utf-8"?>'
        '<package xmlns="http://www.idpf.org/2007/opf" '
        'xmlns:dcterms="http://purl.org/dc/terms/" version="3.0" '
        'unique-identifier="book-id" xml:lang="zh-CN">'
        "<metadata xmlns:dc=\"http://purl.org/dc/elements/1.1/\">"
        '<dc:identifier id="book-id">' + book_id + "</dc:identifier>"
        "<dc:title>" + safe_title + "</dc:title>"
        "<dc:language>zh-CN</dc:language>"
        '<meta property="dcterms:modified">' + modified + "</meta></metadata>"
        '<manifest><item id="nav" href="nav.xhtml" '
        'media-type="application/xhtml+xml" properties="nav"/>'
        + "".join(
            f'<item id="{item_id}" href="{filename}" media-type="application/xhtml+xml"/>'
            for item_id, filename, _chapter in chapter_items
        )
        + "</manifest><spine>"
        + "".join(spine_items)
        + "</spine></package>"
    )
    container = (
        '<?xml version="1.0" encoding="utf-8"?>'
        '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">'
        '<rootfiles><rootfile full-path="OEBPS/content.opf" '
        'media-type="application/oebps-package+xml"/></rootfiles></container>'
    )
    try:
        with zipfile.ZipFile(path, "w") as archive:
            mimetype = zipfile.ZipInfo("mimetype", date_time=(2020, 1, 1, 0, 0, 0))
            mimetype.compress_type = zipfile.ZIP_STORED
            archive.writestr(mimetype, b"application/epub+zip")
            resources = [
                ("META-INF/container.xml", container),
                ("OEBPS/content.opf", package),
                ("OEBPS/nav.xhtml", navigation),
            ] + [(f"OEBPS/{filename}", chapter) for _item_id, filename, chapter in chapter_items]
            for name, content in resources:
                info = zipfile.ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                archive.writestr(info, content.encode("utf-8"))
        with zipfile.ZipFile(path) as archive:
            if archive.testzip() is not None:
                raise ConversionError("生成的 EPUB 完整性检查未通过。")
            if archive.namelist()[0] != "mimetype" or archive.getinfo("mimetype").compress_type != zipfile.ZIP_STORED:
                raise ConversionError("生成的 EPUB 容器结构校验失败。")
    except ConversionError:
        raise
    except (OSError, zipfile.BadZipFile, ValueError) as exc:
        raise ConversionError("无法创建 EPUB 文件，请检查磁盘空间和目标文件夹。") from exc


def _write_docx(path: Path, text: str, source_extension: str) -> None:
    try:
        from docx import Document
    except ImportError as exc:
        raise ConversionError("当前环境没有安装 DOCX 支持组件 python-docx。") from exc
    document = Document()
    if source_extension in ("md", "markdown"):
        for line in text.splitlines():
            heading = re.match(r"^(#{1,6})\s+(.+)$", line)
            if heading:
                document.add_heading(heading.group(2), level=len(heading.group(1)))
            elif line.strip():
                document.add_paragraph(re.sub(r"[*`]+", "", line))
    elif source_extension in ("html", "htm"):
        parser = _HTMLTextExtractor()
        parser.feed(text)
        for line in parser.text().splitlines():
            if line:
                document.add_paragraph(line)
    else:
        for paragraph in text.split("\n\n"):
            if paragraph.strip():
                document.add_paragraph(paragraph)
    document.save(path)


def _write_pdf(path: Path, html: str) -> None:
    from PySide6.QtGui import QFont, QFontDatabase, QTextDocument

    writer = QPdfWriter(str(path))
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setResolution(96)
    writer.setPageMargins(QMarginsF(12, 12, 12, 12), QPageLayout.Unit.Millimeter)
    painter = QPainter()
    if not painter.begin(writer):
        raise ConversionError("无法创建 PDF 输出。")
    document = QTextDocument()
    # The Qt offscreen test platform does not enumerate Windows fonts. Load the
    # built-in CJK font explicitly there and use it for stable Chinese output.
    if os.name == "nt":
        windows_dir = Path(os.environ.get("WINDIR", "C:/Windows"))
        cjk_font = windows_dir / "Fonts" / "msyh.ttc"
        if cjk_font.is_file():
            QFontDatabase.addApplicationFont(str(cjk_font))
        document.setDefaultFont(QFont("Microsoft YaHei UI", 10))
    document.setHtml(html)
    rect = writer.pageLayout().paintRectPixels(writer.resolution())
    document.setPageSize(rect.size())
    page_count = max(1, document.pageCount())
    for page in range(page_count):
        if page:
            writer.newPage()
        painter.save()
        painter.translate(rect.left(), rect.top())
        document.drawContents(painter, QRectF(0, page * rect.height(), rect.width(), rect.height()))
        painter.restore()
    painter.end()


def _write_image_pdf(path: Path, image: Any) -> None:
    writer = QPdfWriter(str(path))
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setResolution(150)
    writer.setPageMargins(QMarginsF(12, 12, 12, 12), QPageLayout.Unit.Millimeter)
    rect = writer.pageLayout().paintRectPixels(writer.resolution())
    image_rect = QRectF(0, 0, image.width(), image.height())
    scale = min(rect.width() / image.width(), rect.height() / image.height())
    draw_rect = QRectF(
        rect.left() + (rect.width() - image.width() * scale) / 2,
        rect.top() + (rect.height() - image.height() * scale) / 2,
        image.width() * scale,
        image.height() * scale,
    )
    painter = QPainter()
    if not painter.begin(writer):
        raise ConversionError("无法创建 PDF 输出。")
    painter.drawImage(draw_rect, image, image_rect)
    painter.end()


def _draw_image_pdf_page(painter: QPainter, rect: QRectF, image: Any) -> None:
    image_rect = QRectF(0, 0, image.width(), image.height())
    scale = min(rect.width() / image.width(), rect.height() / image.height())
    draw_rect = QRectF(
        rect.left() + (rect.width() - image.width() * scale) / 2,
        rect.top() + (rect.height() - image.height() * scale) / 2,
        image.width() * scale,
        image.height() * scale,
    )
    painter.drawImage(draw_rect, image, image_rect)


def _write_images_pdf(
    source_paths: list[Path],
    destination: Path,
    progress_callback: Callable[[int], None] | None = None,
) -> None:
    if not 1 <= len(source_paths) <= MAX_IMAGE_MERGE_FILES:
        raise ConversionError(f"PDF 最多可包含 {MAX_IMAGE_MERGE_FILES} 张图片。")
    canonical = [path.expanduser().resolve() for path in source_paths]
    if len(set(canonical)) != len(canonical):
        raise ConversionError("合并列表中不能重复选择同一张图片。")

    writer = QPdfWriter(str(destination))
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setResolution(150)
    writer.setPageMargins(QMarginsF(12, 12, 12, 12), QPageLayout.Unit.Millimeter)
    rect = writer.pageLayout().paintRectPixels(writer.resolution())
    painter = QPainter()
    if not painter.begin(writer):
        raise ConversionError("无法创建 PDF 输出。")
    total_pixels = 0
    try:
        for index, source in enumerate(canonical):
            _check_source(source, MAX_IMAGE_BYTES)
            if source.suffix.lower().lstrip(".") not in supported_image_extensions():
                raise ConversionError(f"不支持的图片格式：{source.name}")
            reader = QImageReader(str(source))
            reader.setAutoTransform(True)
            size = reader.size()
            if size.isValid() and size.width() * size.height() > MAX_IMAGE_PIXELS:
                raise ConversionError("图片尺寸过大，超过本机安全处理上限。")
            image = reader.read()
            if image.isNull():
                raise ConversionError(f"无法解码图片：{source.name}（{reader.errorString()}）")
            pixels = image.width() * image.height()
            if pixels > MAX_IMAGE_PIXELS or total_pixels + pixels > MAX_PDF_IMAGE_TOTAL_PIXELS:
                raise ConversionError("合并图片的总像素数超过本机安全处理上限。")
            total_pixels += pixels
            if index:
                if not writer.newPage():
                    raise ConversionError("无法继续写入多页 PDF。")
            _draw_image_pdf_page(painter, rect, image)
            if destination.stat().st_size > MAX_MERGED_PDF_BYTES:
                raise ConversionError("生成的 PDF 超过本机 512 MiB 输出上限。")
            if progress_callback:
                progress_callback(int((index + 1) * 100 / len(canonical)))
    finally:
        painter.end()
    if not destination.is_file() or destination.stat().st_size == 0:
        raise ConversionError("PDF 写入失败，没有生成有效文件。")
    if destination.stat().st_size > MAX_MERGED_PDF_BYTES:
        raise ConversionError("生成的 PDF 超过本机 512 MiB 输出上限。")
    document = QPdfDocument()
    try:
        document.load(str(destination))
        if document.status() != QPdfDocument.Status.Ready or document.pageCount() != len(canonical):
            raise ConversionError("生成的 PDF 页面完整性检查未通过。")
    finally:
        document.close()


def images_to_pdf(
    source_paths: list[str | Path],
    output_directory: str | Path | None = None,
    progress_callback: Callable[[int], None] | None = None,
    output_stem: str = "",
) -> Path:
    paths = [Path(path).expanduser().resolve() for path in source_paths]
    if not paths:
        raise ConversionError("没有可合并的图片。")
    destination = _unique_destination(
        paths[0], "pdf", output_directory,
        stem_override=Path(output_stem).name if output_stem else "",
    )
    temporary = _temporary_path(destination, preserve_extension=True)
    try:
        _write_images_pdf(paths, temporary, progress_callback)
        os.replace(temporary, destination)
        return destination
    except ConversionError:
        temporary.unlink(missing_ok=True)
        destination.unlink(missing_ok=True)
        raise
    except Exception as exc:
        temporary.unlink(missing_ok=True)
        destination.unlink(missing_ok=True)
        raise ConversionError(f"图片合并失败：{exc}") from exc


def _write_zip_images_pdf(
    source: Path,
    destination: Path,
    progress_callback: Callable[[int], None] | None = None,
) -> None:
    _check_source(source, MAX_IMAGE_BYTES)
    readable = supported_image_extensions()
    extracted_paths: list[Path] = []
    total_expanded = 0
    try:
        with zipfile.ZipFile(source, "r") as archive, tempfile.TemporaryDirectory(
            prefix="fluke-zip-images-"
        ) as temp_name:
            infos = archive.infolist()
            if len(infos) > MAX_ARCHIVE_MEMBERS:
                raise ConversionError("ZIP 内部文件数量过多，已停止处理。")
            if sum(info.file_size for info in infos) > MAX_ARCHIVE_EXPANDED_BYTES:
                raise ConversionError("ZIP 解压后的内容超过本机安全处理上限。")
            image_infos = sorted(
                (
                    info for info in infos
                    if not info.is_dir()
                    and PurePosixPath(info.filename.replace("\\", "/")).suffix.lower().lstrip(".") in readable
                ),
                key=lambda info: (
                    tuple(
                        (1, int(part)) if part.isdigit() else (0, part.casefold())
                        for part in re.split(r"(\d+)", info.filename.replace("\\", "/"))
                    ),
                    info.header_offset,
                ),
            )
            if not image_infos:
                raise ConversionError("ZIP 中没有可读取的图片文件。")
            if len(image_infos) > MAX_IMAGE_MERGE_FILES:
                raise ConversionError(f"ZIP 中图片超过 {MAX_IMAGE_MERGE_FILES} 张上限。")
            temp_root = Path(temp_name)
            for index, info in enumerate(image_infos, start=1):
                if info.flag_bits & 0x1:
                    raise ConversionError("ZIP 中包含加密图片，无法处理。")
                if info.file_size <= 0 or info.file_size > MAX_IMAGE_BYTES:
                    raise ConversionError("ZIP 内图片为空或超过本机单文件大小上限。")
                if info.file_size / max(1, info.compress_size) > 250:
                    raise ConversionError("ZIP 图片压缩比过高，已停止处理。")
                total_expanded += info.file_size
                if total_expanded > MAX_ARCHIVE_EXPANDED_BYTES:
                    raise ConversionError("ZIP 图片总展开大小超过本机安全处理上限。")
                suffix = PurePosixPath(info.filename.replace("\\", "/")).suffix.lower()
                target = temp_root / f"image-{index:04d}{suffix}"
                copied = 0
                with archive.open(info, "r") as source_stream, target.open("xb") as output_stream:
                    while True:
                        chunk = source_stream.read(1024 * 1024)
                        if not chunk:
                            break
                        copied += len(chunk)
                        if copied > info.file_size or copied > MAX_IMAGE_BYTES:
                            raise ConversionError("ZIP 图片实际展开大小超过声明值或安全上限。")
                        output_stream.write(chunk)
                if copied != info.file_size:
                    raise ConversionError("ZIP 图片长度与目录记录不符，压缩包可能损坏。")
                extracted_paths.append(target)
            _write_images_pdf(extracted_paths, destination, progress_callback)
    except ConversionError:
        raise
    except (OSError, RuntimeError, EOFError, zlib.error, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise ConversionError("ZIP 文件损坏或内部图片无法安全读取。") from exc


def convert_file(
    source: str | Path,
    target_format: str,
    output_directory: str | Path | None = None,
    password: str = "",
    progress_callback: Callable[[int], None] | None = None,
    ocr_language: str = "",
    split_pages_per_file: int = 1,
    video_codec: str = "",
) -> Path:
    source_path = Path(source).expanduser().resolve()
    extension = source_path.suffix.lower().lstrip(".")
    target = str(target_format).lower().lstrip(".")
    available = {item["value"] for item in formats_for(source_path)}
    if target not in available:
        raise ConversionError("所选格式不适用于这个文件，或当前设备没有对应的转换组件。")

    image_sources = supported_image_extensions()
    image_targets = writable_image_extensions()
    is_image = extension in image_sources
    is_raw_image = (
        extension in RAW_CAMERA_IMAGE_EXTENSIONS
        and extension not in image_sources
        and _rawpy_module() is not None
    )
    is_image_ocr = target == "image-ocr-txt" and (is_image or is_raw_image)
    is_pdf_page = extension == "pdf" and target in image_targets
    is_pdf_split = extension == "pdf" and target == "zip"
    is_pdf_images_zip = extension == "pdf" and target == "pdf-images-zip"
    is_pdf_table = extension == "pdf" and target == "xlsx"
    is_pdf_docx = extension == "pdf" and target == "docx"
    is_pdf_security = extension == "pdf" and target in {"pdf-encrypt", "pdf-decrypt"}
    is_pdf_ocr = extension == "pdf" and target == "pdf-ocr"
    is_zip_images_pdf = extension == "zip" and target == "zip-images-pdf"
    is_ofd = extension == "ofd"
    is_mobi = extension == "mobi"
    is_media = extension in MEDIA_INPUT_EXTENSIONS
    is_tabular = extension in TABULAR_FORMATS
    office_operation = _office_operation(extension, target)
    _check_source(
        source_path,
        MAX_IMAGE_BYTES if is_image or is_raw_image else MAX_TABULAR_BYTES if is_tabular
        else MAX_MEDIA_BYTES if is_media else MAX_IMAGE_BYTES if is_zip_images_pdf
        else MAX_EBOOK_SOURCE_BYTES if is_mobi else MAX_DOCUMENT_BYTES,
    )
    if target == "epub" and source_path.stat().st_size > MAX_EBOOK_SOURCE_BYTES:
        raise ConversionError("源文件超过本机 EPUB 输入上限。")
    destination_extension = (
        "txt" if is_image_ocr
        else "pdf" if is_pdf_security or is_pdf_ocr or is_zip_images_pdf
        else "gif" if target == "gif-animation"
        else "zip" if is_pdf_split or is_pdf_images_zip
        or (is_ofd and target == "ofd-png-zip")
        else target
    )
    destination = _unique_destination(source_path, destination_extension, output_directory)
    temporary: Path | None = None
    try:
        temporary = _temporary_path(
            destination,
            preserve_extension=bool(office_operation) or is_media or is_mobi,
        )
        if is_pdf_split:
            _write_pdf_pages_zip(source_path, temporary, split_pages_per_file)
        elif is_pdf_images_zip:
            _write_pdf_images_zip(source_path, temporary)
        elif is_zip_images_pdf:
            _write_zip_images_pdf(source_path, temporary, progress_callback)
        elif is_ofd:
            _convert_ofd(source_path, target, temporary)
        elif is_pdf_table:
            _write_pdf_tables_xlsx(source_path, temporary, progress_callback)
        elif is_pdf_docx:
            _write_pdf_layout_docx(source_path, temporary, progress_callback)
        elif is_pdf_security:
            _write_pdf_security(source_path, temporary, target, password)
        elif is_pdf_ocr:
            try:
                create_searchable_pdf(source_path, temporary, progress_callback, ocr_language)
            except RuntimeError as exc:
                raise ConversionError(str(exc)) from None
        elif is_media:
            try:
                convert_media(source_path, temporary, target, progress_callback, video_codec)
            except ValueError as exc:
                raise ConversionError(str(exc)) from None
        elif is_image_ocr:
            _write_image_ocr_text(
                source_path,
                temporary,
                ocr_language,
                progress_callback,
                is_raw_image,
            )
        elif target == "gif-animation":
            _write_pillow_animation(source_path, temporary, "GIF", progress_callback)
        elif (
            target == "webp"
            and _pillow_can_save("WEBP")
            and _pillow_frame_count(source_path) > 1
        ):
            _write_pillow_animation(source_path, temporary, "WEBP", progress_callback)
        elif is_image or is_raw_image or is_pdf_page:
            if is_raw_image:
                if progress_callback:
                    progress_callback(10)
                image = _read_raw_camera_image(source_path)
                if progress_callback:
                    progress_callback(90)
            else:
                reader = QImageReader(str(source_path))
                reader.setAutoTransform(True)
                size = reader.size()
                if size.isValid() and size.width() * size.height() > MAX_IMAGE_PIXELS:
                    raise ConversionError("图片或 PDF 页面尺寸过大，超过当前转换器的安全处理上限。")
                image = reader.read()
                if image.isNull():
                    raise ConversionError(f"无法解码图片或 PDF 首页面：{reader.errorString()}")
                if image.width() * image.height() > MAX_IMAGE_PIXELS:
                    raise ConversionError("图片或 PDF 页面尺寸过大，超过本机安全处理上限。")
            if target == "pdf":
                _write_image_pdf(temporary, image)
            elif target == "ico":
                _write_multi_size_ico(image, temporary)
            else:
                writer = QImageWriter(str(temporary), IMAGE_FORMATS[target]["qt"].encode("ascii"))
                if target in ("jpg", "jpeg"):
                    writer.setQuality(92)
                try:
                    written = writer.write(image)
                    writer_error = writer.errorString() if not written else ""
                finally:
                    # Release the codec handle before atomic replacement on Windows.
                    del writer
                if not written:
                    raise ConversionError(f"无法写入目标图片：{writer_error}")
        elif office_operation:
            # Office SaveAs/Export APIs require a non-existent destination and
            # select their output format explicitly through the automation API.
            temporary.unlink(missing_ok=True)
            _run_office_conversion(office_operation, source_path, temporary)
        elif is_mobi:
            temporary.unlink(missing_ok=True)
            _run_ebook_convert(source_path, temporary, target)
        elif is_tabular:
            headers, rows = _read_table(source_path, extension)
            if target in TABULAR_FORMATS:
                _write_table(temporary, target, headers, rows)
            elif target == "html":
                temporary.write_text(_wrap_html(_table_html(headers, rows)), encoding="utf-8", newline="")
            elif target == "docx":
                _write_table_docx(temporary, headers, rows)
            elif target == "pdf":
                _write_pdf(temporary, _wrap_html(_table_html(headers, rows)))
            else:
                raise ConversionError("尚不支持所选表格输出格式。")
        else:
            text, blocks = _document_content(source_path, extension)
            if target == "epub":
                ebook_title = source_path.stem
                if blocks is not None:
                    first_heading = next(
                        (content[1] for kind, content in blocks if kind == "heading" and str(content[0]) == "1"),
                        None,
                    )
                    if first_heading:
                        ebook_title = str(first_heading)
                elif extension in ("md", "markdown"):
                    first_heading = re.search(r"^#\s+(.+)$", text, re.MULTILINE)
                    if first_heading:
                        ebook_title = first_heading.group(1).strip()
                chapters = _epub_output_chapters(extension, text, blocks, ebook_title)
                _write_epub(temporary, ebook_title, chapters)
            elif target == "txt":
                temporary.write_text(text, encoding="utf-8", newline="")
            elif target == "md":
                temporary.write_text(
                    _docx_markdown(blocks) if blocks is not None else text,
                    encoding="utf-8",
                    newline="",
                )
            elif target == "html":
                if blocks is not None:
                    body = _docx_html(blocks)
                elif extension == "md":
                    body = _markdown_html(text)
                else:
                    body = "<pre>" + escape(text) + "</pre>"
                temporary.write_text(_wrap_html(body), encoding="utf-8", newline="")
            elif target == "docx":
                if extension in {"md", "markdown"}:
                    _write_markdown_docx(temporary, text)
                else:
                    _write_docx(temporary, text, extension)
            elif target == "pdf":
                if blocks is not None:
                    body = _docx_html(blocks)
                elif extension == "md":
                    body = _markdown_html(text)
                else:
                    body = "<pre>" + escape(text) + "</pre>"
                _write_pdf(temporary, _wrap_html(body))
            else:
                raise ConversionError("尚不支持所选文档输出格式。")
        os.replace(temporary, destination)
    except ConversionError:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        destination.unlink(missing_ok=True)
        raise
    except Exception as exc:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        destination.unlink(missing_ok=True)
        raise ConversionError(f"转换失败：{exc}") from exc
    return destination


class _WorkerSignals(QObject):
    started = Signal(str)
    finished = Signal(str, bool, str, str)
    progress = Signal(str, int)


class _ConversionWorker(QRunnable):
    def __init__(
        self,
        job_id: str,
        source: str,
        target: str,
        output_directory: str,
        password: str = "",
        source_paths: tuple[str, ...] = (),
        ocr_language: str = "",
        split_pages_per_file: int = 1,
        video_codec: str = "",
        output_stem: str = "",
    ):
        super().__init__()
        self.job_id = job_id
        self.source = source
        self.target = target
        self.output_directory = output_directory
        self.password = password
        self.source_paths = source_paths
        self.ocr_language = ocr_language
        self.split_pages_per_file = split_pages_per_file
        self.video_codec = video_codec
        self.output_stem = output_stem
        self.signals = _WorkerSignals()

    @Slot()
    def run(self) -> None:
        self.signals.started.emit(self.job_id)
        try:
            if self.target == "pdf-merge":
                output = merge_pdf_files(self.source_paths, self.output_directory or None)
            elif self.target == "images-pdf":
                output = images_to_pdf(
                    self.source_paths,
                    self.output_directory or None,
                    lambda progress: self.signals.progress.emit(self.job_id, progress),
                    self.output_stem,
                )
            else:
                output = convert_file(
                    self.source,
                    self.target,
                    self.output_directory or None,
                    self.password,
                    lambda progress: self.signals.progress.emit(self.job_id, progress),
                    self.ocr_language,
                    self.split_pages_per_file,
                    self.video_codec,
                )
            self.signals.finished.emit(self.job_id, True, str(output), "")
        except (ConversionError, OSError) as exc:
            self.signals.finished.emit(self.job_id, False, "", str(exc))
        except Exception as exc:
            message = str(exc) or "转换过程中遇到意外问题，请检查文件后重试。"
            self.signals.finished.emit(self.job_id, False, "", message)
        finally:
            self.password = ""
            self.source_paths = ()
            self.ocr_language = ""
            self.split_pages_per_file = 1
            self.video_codec = ""


class _SaveWorkerSignals(QObject):
    finished = Signal(str, bool, str, str)


class _SaveWorker(QRunnable):
    def __init__(
        self,
        job_id: str,
        source_path: str,
        staged_path: str,
        staging_root: str,
        output_directory: str,
    ):
        super().__init__()
        self.job_id = job_id
        self.source_path = source_path
        self.staged_path = staged_path
        self.staging_root = staging_root
        self.output_directory = output_directory
        self.signals = _SaveWorkerSignals()

    @Slot()
    def run(self) -> None:
        try:
            staged = Path(self.staged_path).resolve(strict=True)
            staging_root = Path(self.staging_root).resolve(strict=True)
            staged.relative_to(staging_root)
            output_directory = Path(self.output_directory).expanduser()
            try:
                output_directory.resolve().relative_to(staging_root)
            except ValueError:
                pass
            else:
                raise ConversionError("输出位置不能指向应用的临时转换目录。")
            destination = _save_staged_result(
                Path(self.source_path), staged, output_directory,
            )
            self.signals.finished.emit(self.job_id, True, str(destination), "")
        except Exception as exc:
            message = str(exc) or "无法保存转换结果，请检查输出位置后重试。"
            self.signals.finished.emit(self.job_id, False, "", message)


class ConverterBridge(QObject):
    jobsChanged = Signal()

    def __init__(self, settings: QSettings | None = None) -> None:
        super().__init__()
        self._settings = settings
        self._staging_root = Path(tempfile.mkdtemp(prefix="fluke-converter-"))
        owner_file = self._staging_root / _STAGING_OWNER_FILE
        try:
            owner_file.write_text(json.dumps({"pid": os.getpid()}), encoding="utf-8")
        except OSError:
            shutil.rmtree(self._staging_root, ignore_errors=True)
            raise
        _cleanup_abandoned_staging_roots()
        self._jobs: list[dict[str, Any]] = []
        remembered_directory = str(settings.value("converter/outputDirectory", "") or "") if settings else ""
        self._output_directory = (
            str(Path(remembered_directory).resolve())
            if remembered_directory and Path(remembered_directory).is_dir()
            else ""
        )
        self._task_output_directory = ""
        if settings and remembered_directory and not self._output_directory:
            settings.setValue("converter/outputDirectory", "")
            settings.sync()
        raw_targets = settings.value("converter/targetFormats", {}) if settings else {}
        self._remembered_targets = (
            {str(key).lower().lstrip("."): str(value).lower().lstrip(".") for key, value in raw_targets.items()}
            if isinstance(raw_targets, dict)
            else {}
        )
        raw_remember_targets = (
            settings.value("converter/rememberTargetFormats", None)
            if settings is not None
            else None
        )
        if raw_remember_targets is None:
            # Preserve an existing pre-4.1 preference, while making the
            # previously implicit behavior explicit in the UI.
            self._remember_target_formats = bool(self._remembered_targets)
        elif isinstance(raw_remember_targets, bool):
            self._remember_target_formats = raw_remember_targets
        else:
            self._remember_target_formats = str(raw_remember_targets).strip().lower() in {
                "1", "true", "yes", "on"
            }
        if not self._remember_target_formats:
            self._remembered_targets = {}
        raw_video_codecs = settings.value("converter/videoCodecs", {}) if settings else {}
        self._remembered_video_codecs = (
            {str(key).lower(): str(value).lower() for key, value in raw_video_codecs.items()}
            if isinstance(raw_video_codecs, dict)
            else {}
        )
        self._pool = QThreadPool(self)
        self._pool.setMaxThreadCount(2)
        self._workers: dict[str, _ConversionWorker] = {}
        self._save_workers: dict[str, _SaveWorker] = {}
        self._job_started_at: dict[str, float] = {}
        self._elapsed_timer = QTimer(self)
        self._elapsed_timer.setInterval(1000)
        self._elapsed_timer.timeout.connect(self._update_elapsed_times)
        self._job_passwords: dict[str, str] = {}
        self._job_ocr_languages: dict[str, str] = {}
        self._job_pdf_split_group_sizes: dict[str, int] = {}
        self._job_video_codecs: dict[str, str] = {}
        self._merge_sources: dict[str, tuple[str, ...]] = {}
        self._image_merge_sources: dict[str, tuple[str, ...]] = {}
        self._image_merge_output_stems: dict[str, str] = {}
        self._job_staging_dirs: dict[str, Path] = {}
        application = QCoreApplication.instance()
        if application is not None:
            application.aboutToQuit.connect(self.cleanupStagedResults)

    @Property("QVariant", notify=jobsChanged)
    def jobs(self) -> list[dict[str, Any]]:
        result = []
        for job in self._jobs:
            snapshot = dict(job)
            snapshot["availableFormats"] = [
                dict(option) for option in job.get("availableFormats", [])
            ]
            source_value = str(job.get("sourcePath", ""))
            if source_value:
                if "saveDirectory" in job:
                    planned_directory = str(job.get("saveDirectory") or Path(source_value).parent)
                else:
                    planned_directory = (
                        self._task_output_directory
                        or self._output_directory
                        or str(Path(source_value).parent)
                    )
                snapshot["saveDirectory"] = str(Path(planned_directory).expanduser().resolve())
            result.append(snapshot)
        return result

    @Property(bool, notify=jobsChanged)
    def rememberTargetFormats(self) -> bool:
        return self._remember_target_formats

    @Property("QVariant", notify=jobsChanged)
    def commonTargetFormats(self) -> list[dict[str, str]]:
        candidates = [
            job for job in self._jobs
            if job.get("taskType") == "convert"
            and job.get("canConvert")
            and job.get("status") == "ready"
        ]
        if len(candidates) < 2:
            return []
        common = {
            item["value"] for item in candidates[0].get("availableFormats", [])
        }
        for job in candidates[1:]:
            common.intersection_update(
                item["value"] for item in job.get("availableFormats", [])
            )
        return [
            {"label": item["label"], "value": item["value"]}
            for item in candidates[0].get("availableFormats", [])
            if item.get("value") in common
        ]

    @Property(str, notify=jobsChanged)
    def outputDirectory(self) -> str:
        return self._output_directory

    @Property(str, notify=jobsChanged)
    def taskOutputDirectory(self) -> str:
        return self._task_output_directory

    @Property(str, constant=True)
    def inputFilePatterns(self) -> str:
        """Current local input formats for the native file picker."""
        return " ".join(f"*.{extension}" for extension in sorted(input_extensions()))

    @Property(QUrl, notify=jobsChanged)
    def outputDirectoryUrl(self) -> QUrl:
        return QUrl.fromLocalFile(self._output_directory) if self._output_directory else QUrl()

    @Property(QUrl, notify=jobsChanged)
    def taskOutputDirectoryUrl(self) -> QUrl:
        return QUrl.fromLocalFile(self._task_output_directory) if self._task_output_directory else QUrl()

    @Property("QVariant", notify=jobsChanged)
    def acceptedExtensions(self) -> list[str]:
        return sorted(input_extensions())

    @Property("QVariant", notify=jobsChanged)
    def ocrLanguages(self) -> list[dict[str, str]]:
        return available_ocr_languages()

    @Property(int, constant=True)
    def maxPdfPages(self) -> int:
        return MAX_PDF_PAGES

    @Slot(str, result="QVariant")
    def formatsForPath(self, source: str) -> list[dict[str, str]]:
        return formats_for(source)

    @Slot(str, result="QVariant")
    def videoCodecsForJob(self, job_id: str) -> list[dict[str, str]]:
        job = self._find_job(job_id)
        if not job or job.get("taskType") != "convert":
            return []
        source_extension = Path(str(job.get("sourcePath", ""))).suffix.lower().lstrip(".")
        if source_extension not in VIDEO_INPUT_EXTENSIONS:
            return []
        return available_video_codecs(str(job.get("targetFormat", "")))

    @Slot(str, result=str)
    def videoCodecForJob(self, job_id: str) -> str:
        job = self._find_job(job_id)
        if not job:
            return ""
        options = self.videoCodecsForJob(job_id)
        valid = {item["value"] for item in options}
        selected = self._job_video_codecs.get(job_id)
        if selected in valid:
            return selected
        remembered = self._remembered_video_codecs.get(str(job.get("targetFormat", "")))
        if remembered in valid:
            return str(remembered)
        return str(options[0]["value"]) if options else ""

    @Slot(str, str)
    def setJobVideoCodec(self, job_id: str, codec: str) -> None:
        job = self._find_job(job_id)
        if not job or job.get("status") not in ("ready", "failed"):
            return
        valid = {item["value"] for item in self.videoCodecsForJob(job_id)}
        selected = str(codec).lower()
        if selected not in valid:
            return
        self._job_video_codecs[job_id] = selected
        self._remembered_video_codecs[str(job["targetFormat"])] = selected
        if self._settings:
            self._settings.setValue("converter/videoCodecs", self._remembered_video_codecs)
            self._settings.sync()
        self.jobsChanged.emit()

    @Slot(str, result="QVariant")
    def previewInfo(self, job_id: str) -> dict[str, str]:
        job = self._find_job(job_id)
        if not job or job.get("status") != "done":
            return {"kind": "", "url": ""}
        output = Path(str(job.get("outputPath") or job.get("stagedPath", "")))
        if not output.is_file():
            return {"kind": "", "url": ""}
        extension = output.suffix.lower().lstrip(".")
        if extension in IMAGE_FORMATS:
            kind = "image"
        elif extension == "pdf":
            kind = "pdf"
        elif extension in TEXT_EXTENSIONS | {"csv", "tsv", "json", "html", "htm"}:
            kind = "text"
        elif extension in AUDIO_INPUT_EXTENSIONS:
            kind = "audio"
        elif extension in VIDEO_INPUT_EXTENSIONS:
            kind = "video"
        else:
            kind = ""
        return {
            "kind": kind,
            "url": QUrl.fromLocalFile(str(output.resolve())).toString() if kind else "",
        }

    @Slot(str, result="QVariant")
    def previewText(self, job_id: str) -> dict[str, Any]:
        job = self._find_job(job_id)
        if not job or job.get("status") != "done":
            return {"text": "", "truncated": False}
        output = Path(str(job.get("outputPath") or job.get("stagedPath", "")))
        text_extensions = TEXT_EXTENSIONS | {"csv", "tsv", "json", "html", "htm"}
        if not output.is_file() or output.suffix.lower().lstrip(".") not in text_extensions:
            return {"text": "", "truncated": False}
        try:
            with output.open("rb") as stream:
                payload = stream.read(MAX_TEXT_PREVIEW_BYTES + 1)
        except OSError:
            return {"text": "无法读取转换结果预览。", "truncated": False}
        truncated = len(payload) > MAX_TEXT_PREVIEW_BYTES
        payload = payload[:MAX_TEXT_PREVIEW_BYTES]
        if payload.startswith((b"\xff\xfe", b"\xfe\xff")):
            content = payload.decode("utf-16", errors="replace")
        else:
            content = payload.decode("utf-8-sig", errors="replace")
        if output.suffix.lower() in (".html", ".htm"):
            parser = _HTMLTextExtractor()
            parser.feed(content)
            content = parser.text()
        return {"text": content, "truncated": truncated}

    @Slot()
    def mergeQueuedPdfs(self) -> None:
        self.mergePdfJobs([item["jobId"] for item in self._mergeable_pdf_items()])

    @Slot(result="QVariant")
    def pdfMergeCandidates(self) -> dict[str, Any]:
        items = self._mergeable_pdf_items()
        public_items = [
            {key: item[key] for key in ("jobId", "fileName", "folder")}
            for item in items
        ]
        excluded = [
            exclusion
            for job in self._jobs
            if (exclusion := self._pdf_merge_exclusion(job)) is not None
        ]
        return {"items": public_items, "skipped": len(excluded), "excluded": excluded}

    @Slot("QVariant", result="QVariant")
    def mergePdfJobs(self, requested_job_ids: Any) -> dict[str, Any]:
        job_ids = [str(value) for value in self._variant_items(requested_job_ids)]
        if len(job_ids) < 2:
            return {"ok": False, "error": "合并至少需要两份有效 PDF。"}
        if len(set(job_ids)) != len(job_ids):
            return {"ok": False, "error": "合并列表中不能重复选择同一文件。"}

        eligible = {item["jobId"]: item["sourcePath"] for item in self._mergeable_pdf_items()}
        if any(job_id not in eligible for job_id in job_ids):
            return {"ok": False, "error": "队列状态已变化，请重新检查合并顺序。"}
        sources = [eligible[job_id] for job_id in job_ids]
        self._start_pdf_merge_job(sources)
        return {"ok": True}

    @Slot(result="QVariant")
    def imageMergeCandidates(self) -> dict[str, Any]:
        items = self._mergeable_image_items()
        public_items = [
            {key: item[key] for key in ("jobId", "fileName", "folder")}
            for item in items
        ]
        excluded = [
            exclusion
            for job in self._jobs
            if (exclusion := self._image_merge_exclusion(job)) is not None
        ]
        return {"items": public_items, "skipped": len(excluded), "excluded": excluded}

    @Slot("QVariant", result="QVariant")
    def mergeImagesToPdf(self, requested_job_ids: Any) -> dict[str, Any]:
        job_ids = [str(value) for value in self._variant_items(requested_job_ids)]
        if not 2 <= len(job_ids) <= MAX_IMAGE_MERGE_FILES:
            return {"ok": False, "error": f"合并需要选择 2 至 {MAX_IMAGE_MERGE_FILES} 张图片。"}
        if len(set(job_ids)) != len(job_ids):
            return {"ok": False, "error": "合并列表中不能重复选择同一张图片。"}
        eligible = {item["jobId"]: item["sourcePath"] for item in self._mergeable_image_items()}
        if any(job_id not in eligible for job_id in job_ids):
            return {"ok": False, "error": "队列状态已变化，请重新检查图片合并顺序。"}
        self._start_image_merge_job([eligible[job_id] for job_id in job_ids])
        return {"ok": True}

    @Slot("QVariant", result="QVariant")
    def addImageFolderToPdf(self, value: Any) -> dict[str, Any]:
        folder_value = self._local_path(value)
        if not folder_value:
            return {"ok": False, "error": "请选择有效的本机文件夹。"}
        folder = Path(folder_value).expanduser()
        if not folder.is_dir():
            return {"ok": False, "error": "请选择有效的本机文件夹。"}

        readable = supported_image_extensions()
        candidates: list[Path] = []
        try:
            with os.scandir(folder) as entries:
                for entry in entries:
                    if not entry.is_file(follow_symlinks=False):
                        continue
                    source = Path(entry.path)
                    if source.suffix.lower().lstrip(".") not in readable:
                        continue
                    candidates.append(source)
                    if len(candidates) > MAX_IMAGE_MERGE_FILES:
                        return {
                            "ok": False,
                            "error": "文件夹图片数量超过本机处理上限。",
                        }
        except OSError:
            return {"ok": False, "error": "无法读取所选文件夹。"}

        candidates.sort(
            key=lambda path: tuple(
                (1, int(part)) if part.isdigit() else (0, part.casefold())
                for part in re.split(r"(\d+)", path.name)
            )
        )
        sources: list[str] = []
        skipped = 0
        for source in candidates:
            try:
                if any(option["value"] == "pdf" for option in formats_for(source)):
                    sources.append(str(source.resolve()))
                else:
                    skipped += 1
            except (OSError, RuntimeError, ConversionError):
                skipped += 1
        if not sources:
            return {"ok": False, "error": "所选文件夹中没有可读取的图片。"}

        resolved_folder = folder.resolve()
        self._start_image_merge_job(
            sources,
            folder_label=resolved_folder.name,
            output_stem=resolved_folder.name,
        )
        return {"ok": True, "count": len(sources), "skipped": skipped}

    def _mergeable_pdf_items(self) -> list[dict[str, str]]:
        items: list[dict[str, str]] = []
        for job in self._jobs:
            if job.get("taskType") == "pdf-merge":
                continue
            if self._pdf_merge_exclusion(job) is not None:
                continue
            source = Path(job["sourcePath"])
            items.append({
                "jobId": str(job["id"]),
                "fileName": str(job["fileName"]),
                "folder": str(source.parent),
                "sourcePath": str(source),
            })
        return items

    def _mergeable_image_items(self) -> list[dict[str, str]]:
        items: list[dict[str, str]] = []
        for job in self._jobs:
            if job.get("taskType") in ("pdf-merge", "images-pdf"):
                continue
            if self._image_merge_exclusion(job) is not None:
                continue
            source = Path(job["sourcePath"])
            items.append({
                "jobId": str(job["id"]),
                "fileName": str(job["fileName"]),
                "folder": str(source.parent),
                "sourcePath": str(source),
            })
        return items

    @staticmethod
    def _pdf_merge_exclusion(job: dict[str, Any]) -> dict[str, str] | None:
        if job.get("taskType") == "pdf-merge":
            return None
        source = Path(str(job.get("sourcePath", "")))
        if source.suffix.lower() != ".pdf":
            return None
        reason_code = ""
        if not source.is_file():
            reason_code = "missing"
        elif not _looks_like_pdf(source):
            reason_code = "invalid"
        else:
            encryption_state = _pdf_encryption_state_for_path(source)
            if job.get("targetFormat") == "pdf-decrypt":
                reason_code = "decrypt_first"
            elif encryption_state is True:
                reason_code = "encrypted"
            elif encryption_state is None:
                reason_code = "unreadable"
            elif job.get("status") == "running":
                reason_code = "running"
            elif job.get("status") == "unsupported":
                reason_code = "unsupported"
            elif not job.get("canConvert"):
                reason_code = "not_convertible"
        if not reason_code:
            return None
        return {
            "fileName": str(job.get("fileName") or source.name),
            "folder": str(source.parent),
            "reasonCode": reason_code,
        }

    @staticmethod
    def _image_merge_exclusion(job: dict[str, Any]) -> dict[str, str] | None:
        if job.get("taskType") in ("pdf-merge", "images-pdf"):
            return None
        source = Path(str(job.get("sourcePath", "")))
        if source.suffix.lower().lstrip(".") not in supported_image_extensions():
            return None
        reason_code = ""
        if not source.is_file():
            reason_code = "missing"
        elif job.get("taskType") != "convert":
            reason_code = "not_convertible"
        else:
            try:
                _check_source(source, MAX_IMAGE_BYTES)
                reader = QImageReader(str(source))
                reader.setAutoTransform(True)
                size = reader.size()
                if size.isValid() and size.width() * size.height() > MAX_IMAGE_PIXELS:
                    reason_code = "too_large"
                elif not reader.canRead():
                    reason_code = "unreadable"
            except (OSError, RuntimeError, ConversionError):
                reason_code = "unreadable"
            if not reason_code and job.get("status") == "running":
                reason_code = "running"
            elif not reason_code and job.get("status") == "unsupported":
                reason_code = "unsupported"
        if not reason_code:
            return None
        return {
            "fileName": str(job.get("fileName") or source.name),
            "folder": str(source.parent),
            "reasonCode": reason_code,
        }

    def _start_pdf_merge_job(self, sources: list[str]) -> None:
        job_id = uuid.uuid4().hex
        job = {
            "id": job_id,
            "sourcePath": sources[0],
            "fileName": f"合并 {len(sources)} 个 PDF",
            "taskType": "pdf-merge",
            "targetFormat": "pdf-merge",
            "status": "ready",
            "statusLabel": "等待转换",
            "progress": 0,
            "progressMode": "indeterminate",
            "elapsedSeconds": 0,
            "outputPath": "",
            "stagedPath": "",
            "saveError": "",
            "savePending": False,
            "error": "",
            "canConvert": True,
        }
        self._jobs.append(job)
        self._merge_sources[job_id] = tuple(sources)
        self._start_job(job)
        self.jobsChanged.emit()

    def _start_image_merge_job(
        self,
        sources: list[str],
        *,
        folder_label: str = "",
        output_stem: str = "",
    ) -> None:
        job_id = uuid.uuid4().hex
        job = {
            "id": job_id,
            "sourcePath": sources[0],
            "fileName": (
                f"{folder_label} · {len(sources)} 张图片"
                if folder_label else f"合并 {len(sources)} 张图片"
            ),
            "taskType": "images-pdf",
            "targetFormat": "images-pdf",
            "status": "ready",
            "statusLabel": "等待转换",
            "progress": 0,
            "progressMode": "indeterminate",
            "elapsedSeconds": 0,
            "outputPath": "",
            "stagedPath": "",
            "saveError": "",
            "savePending": False,
            "error": "",
            "canConvert": True,
        }
        self._jobs.append(job)
        self._image_merge_sources[job_id] = tuple(sources)
        if output_stem:
            self._image_merge_output_stems[job_id] = output_stem
        self._start_job(job)
        self.jobsChanged.emit()

    @Slot(str, str)
    def setJobPassword(self, job_id: str, password: str) -> None:
        job = self._find_job(job_id)
        if not job or job["status"] not in ("ready", "failed"):
            return
        if job["targetFormat"] not in ("pdf-encrypt", "pdf-decrypt"):
            self._job_passwords.pop(job_id, None)
            return
        self._job_passwords[job_id] = str(password)[:MAX_PDF_PASSWORD_LENGTH]

    @Slot(str, result=str)
    def passwordForJob(self, job_id: str) -> str:
        return self._job_passwords.get(job_id, "")

    @Slot(str, result=str)
    def ocrLanguageForJob(self, job_id: str) -> str:
        selected = self._job_ocr_languages.get(job_id, "")
        if selected:
            return selected
        engine = tesseract_engine()
        return engine.language if engine else ""

    @Slot(str, str)
    def setJobOcrLanguage(self, job_id: str, language: str) -> None:
        job = self._find_job(job_id)
        if (
            not job
            or job["status"] not in ("ready", "failed")
            or job["targetFormat"] not in {"pdf-ocr", "image-ocr-txt"}
        ):
            return
        valid_languages = {item["value"] for item in available_ocr_languages()}
        if language in valid_languages:
            self._job_ocr_languages[job_id] = language
            self.jobsChanged.emit()

    @Slot(str, result=int)
    def pdfSplitGroupSizeForJob(self, job_id: str) -> int:
        return self._job_pdf_split_group_sizes.get(job_id, 1)

    @Slot(str, int)
    def setPdfSplitGroupSize(self, job_id: str, pages_per_file: int) -> None:
        job = self._find_job(job_id)
        if not job or job["status"] not in ("ready", "failed") or job["targetFormat"] != "zip":
            return
        if not 1 <= pages_per_file <= MAX_PDF_PAGES:
            return
        if self._job_pdf_split_group_sizes.get(job_id, 1) != pages_per_file:
            self._job_pdf_split_group_sizes[job_id] = pages_per_file
            self.jobsChanged.emit()

    @Slot("QVariant")
    def addFiles(self, values: Any) -> None:
        for value in self._variant_items(values):
            path = self._local_path(value)
            if not path:
                continue
            source = Path(path).expanduser()
            if not source.is_file():
                continue
            resolved = str(source.resolve())
            if any(job["sourcePath"] == resolved for job in self._jobs):
                continue
            formats = formats_for(resolved)
            if not formats:
                self._jobs.append({
                    "id": uuid.uuid4().hex,
                    "sourcePath": resolved,
                    "fileName": source.name,
                    "availableFormats": [],
                    "targetFormat": "",
                    "status": "unsupported",
                    "statusLabel": "暂不支持此格式",
                    "progress": 0,
                    "elapsedSeconds": 0,
                    "outputPath": "",
                    "stagedPath": "",
                    "saveError": "",
                    "savePending": False,
                    "error": "此文件类型在当前设备没有可用的本机转换组件。",
                    "canConvert": False,
                })
                continue
            extension = source.suffix.lower().lstrip(".")
            preferred = (
                self._remembered_targets.get(extension)
                if self._remember_target_formats
                else ""
            ) or {
                "csv": "json", "tsv": "csv", "json": "csv", "xlsx": "csv"
            }.get(extension, "png")
            default_target = formats[0]["value"]
            if extension in VIDEO_INPUT_EXTENSIONS:
                default_target = next(
                    (
                        item["value"]
                        for video_target in VIDEO_TARGETS
                        for item in formats
                        if item["value"] == video_target
                    ),
                    default_target,
                )
            target = next(
                (item["value"] for item in formats if item["value"] == preferred),
                default_target,
            )
            self._jobs.append({
                "id": uuid.uuid4().hex,
                "sourcePath": resolved,
                "fileName": source.name,
                "availableFormats": [dict(item) for item in formats],
                "taskType": "convert",
                "targetFormat": target,
                "status": "ready",
                "statusLabel": "等待转换",
                "progress": 0,
                "progressMode": "indeterminate",
                "elapsedSeconds": 0,
                "outputPath": "",
                "stagedPath": "",
                "saveError": "",
                "savePending": False,
                "error": "",
                "canConvert": True,
            })
        self.jobsChanged.emit()

    @Slot()
    def refreshEngineAvailability(self) -> None:
        """Refresh queued file targets after a bundled engine install/update."""
        changed = False
        for job in self._jobs:
            if job.get("taskType") not in (None, "convert") or job.get("status") not in ("ready", "unsupported"):
                continue
            source = str(job.get("sourcePath", ""))
            if not Path(source).is_file():
                continue
            formats = formats_for(source)
            if not formats:
                continue
            available = {item["value"] for item in formats}
            selected = str(job.get("targetFormat", ""))
            if selected not in available:
                extension = Path(source).suffix.lower().lstrip(".")
                preferred = (
                    self._remembered_targets.get(extension, "")
                    if self._remember_target_formats
                    else ""
                )
                selected = preferred if preferred in available else formats[0]["value"]
            job.update({
                "availableFormats": [dict(item) for item in formats],
                "taskType": "convert",
                "targetFormat": selected,
                "status": "ready",
                "statusLabel": "等待转换",
                "progress": 0,
                "elapsedSeconds": 0,
                "error": "",
                "canConvert": True,
            })
            changed = True
        if changed:
            self.jobsChanged.emit()

    @Slot(str, str)
    def setTargetFormat(self, job_id: str, target_format: str) -> None:
        job = self._find_job(job_id)
        if not job or job["status"] != "ready" or job.get("taskType") in ("pdf-merge", "images-pdf"):
            return
        target = str(target_format).lower().lstrip(".")
        options = job.get("availableFormats") or formats_for(job["sourcePath"])
        if target in {item["value"] for item in options}:
            extension = Path(job["sourcePath"]).suffix.lower().lstrip(".")
            if self._remember_target_formats:
                self._remembered_targets[extension] = target
                self._persist_target_preferences()
            if target != job["targetFormat"]:
                self._job_passwords.pop(job_id, None)
                self._job_ocr_languages.pop(job_id, None)
                self._job_pdf_split_group_sizes.pop(job_id, None)
            job["targetFormat"] = target
            self.jobsChanged.emit()

    @Slot(str, result="QVariant")
    def setCommonTargetFormat(self, target_format: str) -> dict[str, Any]:
        target = str(target_format).lower().lstrip(".")
        if target not in {item["value"] for item in self.commonTargetFormats}:
            return {"ok": False, "applied": 0, "error": "所选格式不是当前队列文件的共同目标。"}
        candidates = [
            job for job in self._jobs
            if job.get("taskType") == "convert"
            and job.get("canConvert")
            and job.get("status") == "ready"
            and target in {
                item["value"] for item in job.get("availableFormats", [])
            }
        ]
        if len(candidates) < 2:
            return {"ok": False, "applied": 0, "error": "至少需要两个等待转换的文件。"}
        changed = False
        for job in candidates:
            extension = Path(job["sourcePath"]).suffix.lower().lstrip(".")
            if self._remember_target_formats:
                self._remembered_targets[extension] = target
            if job["targetFormat"] != target:
                self._job_passwords.pop(str(job["id"]), None)
                self._job_ocr_languages.pop(str(job["id"]), None)
                self._job_pdf_split_group_sizes.pop(str(job["id"]), None)
                self._job_video_codecs.pop(str(job["id"]), None)
                job["targetFormat"] = target
                changed = True
        self._persist_target_preferences()
        if changed:
            self.jobsChanged.emit()
        return {"ok": True, "applied": len(candidates), "changed": changed}

    @Slot(bool)
    def setRememberTargetFormats(self, enabled: bool) -> None:
        enabled = bool(enabled)
        if enabled == self._remember_target_formats:
            return
        self._remember_target_formats = enabled
        if not enabled:
            self._remembered_targets = {}
        self._persist_target_preferences()
        self.jobsChanged.emit()

    def _persist_target_preferences(self) -> None:
        if self._settings is None:
            return
        self._settings.setValue(
            "converter/rememberTargetFormats", self._remember_target_formats
        )
        self._settings.setValue(
            "converter/targetFormats",
            self._remembered_targets if self._remember_target_formats else {},
        )
        self._settings.sync()

    @Slot("QVariant")
    def setOutputDirectory(self, value: Any) -> None:
        path = self._local_path(value)
        if path and Path(path).is_dir():
            self._output_directory = str(Path(path).resolve())
        else:
            self._output_directory = ""
        if self._settings:
            self._settings.setValue("converter/outputDirectory", self._output_directory)
            self._settings.sync()
        self.jobsChanged.emit()

    @Slot("QVariant")
    def setTaskOutputDirectory(self, value: Any) -> None:
        path = self._local_path(value)
        self._task_output_directory = str(Path(path).resolve()) if path and Path(path).is_dir() else ""
        self.jobsChanged.emit()

    @Slot()
    def clearTaskOutputDirectory(self) -> None:
        if not self._task_output_directory:
            return
        self._task_output_directory = ""
        self.jobsChanged.emit()

    @Slot()
    def convertReady(self) -> None:
        for job in self._jobs:
            if job["status"] != "ready" or not job["canConvert"]:
                continue
            self._start_job(job)
        self.jobsChanged.emit()

    def _start_job(self, job: dict[str, Any]) -> None:
        job["saveDirectory"] = self._task_output_directory or self._output_directory
        job["status"] = "running"
        job["statusLabel"] = "正在转换"
        job["progress"] = 0
        job["elapsedSeconds"] = 0
        job["outputPath"] = ""
        job["stagedPath"] = ""
        job["saveError"] = ""
        job["savePending"] = False
        staging_directory = Path(tempfile.mkdtemp(prefix="job-", dir=self._staging_root))
        self._job_staging_dirs[job["id"]] = staging_directory
        worker = _ConversionWorker(
            job["id"], job["sourcePath"], job["targetFormat"], str(staging_directory),
            self._job_passwords.get(job["id"], ""),
            self._image_merge_sources.get(job["id"], self._merge_sources.get(job["id"], ())),
            self._job_ocr_languages.get(job["id"], ""),
            self._job_pdf_split_group_sizes.get(job["id"], 1),
            self.videoCodecForJob(job["id"]),
            self._image_merge_output_stems.get(job["id"], ""),
        )
        worker.signals.started.connect(self._on_started)
        worker.signals.finished.connect(self._on_finished)
        worker.signals.progress.connect(self._on_progress)
        self._workers[job["id"]] = worker
        self._pool.start(worker)

    @Slot(str)
    def _on_started(self, job_id: str) -> None:
        job = self._find_job(job_id)
        if not job or job.get("status") != "running":
            return
        self._job_started_at[job_id] = time.monotonic()
        if not self._elapsed_timer.isActive():
            self._elapsed_timer.start()
        job["elapsedSeconds"] = 0
        self.jobsChanged.emit()

    @Slot(str)
    def removeJob(self, job_id: str) -> None:
        if any(
            job["id"] == job_id and (job["status"] == "running" or job.get("savePending"))
            for job in self._jobs
        ):
            return
        self._remove_job_staging(job_id)
        self._job_passwords.pop(job_id, None)
        self._job_ocr_languages.pop(job_id, None)
        self._job_pdf_split_group_sizes.pop(job_id, None)
        self._job_video_codecs.pop(job_id, None)
        self._merge_sources.pop(job_id, None)
        self._image_merge_sources.pop(job_id, None)
        self._image_merge_output_stems.pop(job_id, None)
        self._job_started_at.pop(job_id, None)
        self._jobs = [job for job in self._jobs if job["id"] != job_id]
        if not self._jobs:
            self._task_output_directory = ""
        self.jobsChanged.emit()

    @Slot(str)
    def retryJob(self, job_id: str) -> None:
        job = self._find_job(job_id)
        if not job or job["status"] != "failed":
            return
        job.update(
            status="ready", statusLabel="等待转换", progress=0,
            outputPath="", stagedPath="", saveError="", savePending=False,
            error="", elapsedSeconds=0,
        )
        if job["targetFormat"] not in ("pdf-encrypt", "pdf-decrypt"):
            self._job_passwords.pop(job_id, None)
        self.jobsChanged.emit()

    @Slot(str, result="QVariant")
    def saveJob(self, job_id: str) -> dict[str, Any]:
        job = self._find_job(job_id)
        if not job or job.get("status") != "done" or not job.get("stagedPath"):
            return {"ok": False, "error": "这个任务没有待保存的转换结果。"}
        if job.get("savePending"):
            return {"ok": False, "error": "这个结果正在保存。"}
        queued = self._queue_save_job(job)
        if queued:
            self.jobsChanged.emit()
            return {"ok": True, "pending": True}
        self.jobsChanged.emit()
        return {"ok": False, "error": str(job.get("saveError", "无法开始保存。"))}

    @Slot(result="QVariant")
    def saveAll(self) -> dict[str, Any]:
        pending = [
            job for job in self._jobs
            if job.get("status") == "done" and job.get("stagedPath") and not job.get("savePending")
        ]
        scheduled = 0
        for job in pending:
            if self._queue_save_job(job):
                scheduled += 1
        if pending:
            self.jobsChanged.emit()
        return {"ok": scheduled == len(pending), "scheduled": scheduled, "failed": len(pending) - scheduled}

    def _queue_save_job(self, job: dict[str, Any]) -> bool:
        job_id = str(job["id"])
        source_path = Path(str(job["sourcePath"]))
        if "saveDirectory" in job:
            saved_directory = str(job.get("saveDirectory") or "")
        else:
            saved_directory = self._task_output_directory or self._output_directory
        destination_directory = (
            Path(saved_directory).expanduser()
            if saved_directory else source_path.parent
        )
        worker = _SaveWorker(
            job_id, str(source_path), str(job.get("stagedPath", "")),
            str(self._staging_root), str(destination_directory),
        )
        worker.signals.finished.connect(self._on_save_finished)
        self._save_workers[job_id] = worker
        job["savePending"] = True
        job["saveError"] = ""
        self._pool.start(worker)
        return True

    @Slot(str, bool, str, str)
    def _on_save_finished(self, job_id: str, succeeded: bool, output: str, error: str) -> None:
        self._save_workers.pop(job_id, None)
        job = self._find_job(job_id)
        if not job:
            self._remove_job_staging(job_id)
            return
        job["savePending"] = False
        if succeeded:
            self._remove_job_staging(job_id)
            job.update(outputPath=output, stagedPath="", saveError="")
        else:
            job["saveError"] = error
        self.jobsChanged.emit()

    def _remove_job_staging(self, job_id: str) -> None:
        directory = self._job_staging_dirs.pop(job_id, None)
        if directory is not None:
            shutil.rmtree(directory, ignore_errors=True)

    @Slot()
    def cleanupStagedResults(self) -> None:
        active_ids = {
            job["id"] for job in self._jobs if job.get("status") == "running"
        }
        active_ids |= {
            job["id"] for job in self._jobs if job.get("savePending")
        }
        for job_id in tuple(self._job_staging_dirs):
            if job_id not in active_ids:
                self._remove_job_staging(job_id)
        if not active_ids:
            shutil.rmtree(self._staging_root, ignore_errors=True)

    @Slot()
    def clearFinished(self) -> None:
        kept_ids = {
            job["id"] for job in self._jobs
            if job["status"] in ("ready", "running") or job.get("savePending")
        }
        for job in self._jobs:
            if job["id"] not in kept_ids:
                self._remove_job_staging(job["id"])
        self._job_passwords = {
            job_id: password for job_id, password in self._job_passwords.items()
            if job_id in kept_ids
        }
        self._job_ocr_languages = {
            job_id: language for job_id, language in self._job_ocr_languages.items()
            if job_id in kept_ids
        }
        self._job_pdf_split_group_sizes = {
            job_id: page_count for job_id, page_count in self._job_pdf_split_group_sizes.items()
            if job_id in kept_ids
        }
        self._job_video_codecs = {
            job_id: codec for job_id, codec in self._job_video_codecs.items()
            if job_id in kept_ids
        }
        self._merge_sources = {
            job_id: sources for job_id, sources in self._merge_sources.items()
            if job_id in kept_ids
        }
        self._image_merge_sources = {
            job_id: sources for job_id, sources in self._image_merge_sources.items()
            if job_id in kept_ids
        }
        self._image_merge_output_stems = {
            job_id: stem for job_id, stem in self._image_merge_output_stems.items()
            if job_id in kept_ids
        }
        self._jobs = [
            job for job in self._jobs
            if job["status"] in ("ready", "running") or job.get("savePending")
        ]
        if not self._jobs:
            self._task_output_directory = ""
        self.jobsChanged.emit()

    @Slot(str, bool, str, str)
    def _on_finished(self, job_id: str, succeeded: bool, output: str, error: str) -> None:
        job = self._find_job(job_id)
        self._workers.pop(job_id, None)
        started_at = self._job_started_at.pop(job_id, None)
        self._job_passwords.pop(job_id, None)
        self._job_ocr_languages.pop(job_id, None)
        if not job:
            if not self._job_started_at:
                self._elapsed_timer.stop()
            return
        if started_at is not None:
            job["elapsedSeconds"] = max(0, int(time.monotonic() - started_at))
        if succeeded:
            self._job_pdf_split_group_sizes.pop(job_id, None)
            self._merge_sources.pop(job_id, None)
            self._image_merge_sources.pop(job_id, None)
            self._image_merge_output_stems.pop(job_id, None)
            job.update(
                status="done", statusLabel="转换完成", progress=100,
                outputPath="", stagedPath=output, saveError="", savePending=False, error="",
            )
        else:
            self._remove_job_staging(job_id)
            job.update(
                status="failed", statusLabel="转换失败", progress=0,
                outputPath="", stagedPath="", saveError="", savePending=False, error=error,
            )
        self.jobsChanged.emit()
        if not self._job_started_at:
            self._elapsed_timer.stop()

    @Slot()
    def _update_elapsed_times(self) -> None:
        now = time.monotonic()
        changed = False
        for job_id, started_at in self._job_started_at.items():
            job = self._find_job(job_id)
            if not job or job.get("status") != "running":
                continue
            elapsed = max(0, int(now - started_at))
            if job.get("elapsedSeconds") != elapsed:
                job["elapsedSeconds"] = elapsed
                changed = True
        if changed:
            self.jobsChanged.emit()

    @Slot(str, int)
    def _on_progress(self, job_id: str, progress: int) -> None:
        job = self._find_job(job_id)
        if not job or job["status"] != "running":
            return
        job["progress"] = max(0, min(100, int(progress)))
        job["progressMode"] = "determinate"
        self.jobsChanged.emit()

    def _find_job(self, job_id: str) -> dict[str, Any] | None:
        return next((job for job in self._jobs if job["id"] == job_id), None)

    @staticmethod
    def _variant_items(value: Any) -> list[Any]:
        if isinstance(value, QJSValue):
            value = value.toVariant()
        if isinstance(value, (list, tuple)):
            return list(value)
        if value is None:
            return []
        return [value]

    @staticmethod
    def _local_path(value: Any) -> str:
        if isinstance(value, QJSValue):
            value = value.toVariant()
        if isinstance(value, QUrl):
            return value.toLocalFile()
        text = str(value or "")
        if text.lower().startswith("file:"):
            return QUrl(text).toLocalFile()
        return text
