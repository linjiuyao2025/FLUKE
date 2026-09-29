from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from wanxiang.spreadsheet_export import SpreadsheetExportError, export_xlsx


class SpreadsheetExportTests(unittest.TestCase):
    def test_exports_header_only_workbook_and_returns_resolved_path(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "empty.xlsx"

            result = export_xlsx(target, "记录", ["日期", "金额"], [])

            self.assertEqual(result, target.resolve())
            self.assertTrue(result.is_file())
            workbook = load_workbook(result, data_only=False)
            self.assertEqual(workbook.sheetnames, ["记录"])
            worksheet = workbook["记录"]
            self.assertEqual(worksheet.max_row, 1)
            self.assertEqual(worksheet.max_column, 2)
            self.assertEqual([cell.value for cell in worksheet[1]], ["日期", "金额"])
            workbook.close()

    def test_preserves_row_column_and_scalar_types_after_reopen(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "typed.xlsx"
            records = [
                [date(2026, 9, 27), 12, Decimal("2.50"), True, None],
                [datetime(2026, 9, 27, 8, 30), -3.25, 0, False, "说明"],
            ]

            result = export_xlsx(
                target,
                "明细",
                ["日期", "数量", "金额", "已完成", "备注"],
                records,
            )

            workbook = load_workbook(result, data_only=False)
            worksheet = workbook["明细"]
            self.assertEqual((worksheet.max_row, worksheet.max_column), (3, 5))
            # Excel stores a calendar date as a date-formatted serial. openpyxl
            # exposes that serial as midnight datetime after reopening.
            self.assertEqual(worksheet["A2"].value.date(), date(2026, 9, 27))
            self.assertEqual(worksheet["A2"].data_type, "d")
            self.assertEqual(worksheet["B2"].value, 12)
            self.assertEqual(worksheet["B2"].data_type, "n")
            self.assertEqual(worksheet["C2"].value, 2.5)
            self.assertEqual(worksheet["C2"].data_type, "n")
            self.assertIs(worksheet["D2"].value, True)
            self.assertEqual(worksheet["D2"].data_type, "b")
            self.assertIsNone(worksheet["E2"].value)
            self.assertEqual(worksheet["A3"].value, datetime(2026, 9, 27, 8, 30))
            self.assertEqual(worksheet["A3"].data_type, "d")
            self.assertEqual(worksheet["B3"].value, -3.25)
            self.assertIs(worksheet["D3"].value, False)
            self.assertEqual(worksheet["E3"].value, "说明")
            self.assertEqual(worksheet["E3"].data_type, "s")
            workbook.close()

    def test_formula_like_user_strings_and_headers_are_stored_as_text(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "text.xlsx"
            formula_text = '=HYPERLINK("https://example.invalid","打开")'

            export_xlsx(
                target,
                "文本",
                ["=标题"],
                [[formula_text], ["=1+1"], ["+cmd"]],
            )

            workbook = load_workbook(target, data_only=False)
            worksheet = workbook["文本"]
            self.assertEqual(worksheet["A1"].value, "=标题")
            self.assertEqual(worksheet["A1"].data_type, "s")
            self.assertEqual(worksheet["A2"].value, formula_text)
            self.assertEqual(worksheet["A2"].data_type, "s")
            self.assertEqual(worksheet["A3"].value, "=1+1")
            self.assertEqual(worksheet["A3"].data_type, "s")
            self.assertEqual(worksheet["A4"].value, "+cmd")
            self.assertEqual(worksheet["A4"].data_type, "s")
            workbook.close()

    def test_invalid_rows_and_values_raise_readable_errors_without_creating_target(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "invalid.xlsx"

            with self.assertRaisesRegex(SpreadsheetExportError, "第 2 行有 1 列"):
                export_xlsx(target, "记录", ["日期", "金额"], [["今天"]])
            self.assertFalse(target.exists())

            with self.assertRaisesRegex(SpreadsheetExportError, "数据类型不受支持"):
                export_xlsx(target, "记录", ["值"], [[object()]])
            self.assertFalse(target.exists())

    def test_failed_save_and_failed_atomic_replace_preserve_existing_file(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            directory = Path(folder)
            target = directory / "existing.xlsx"
            original = b"keep original bytes"
            target.write_bytes(original)

            with patch("wanxiang.spreadsheet_export.Workbook.save", side_effect=OSError("合成保存故障")):
                with self.assertRaisesRegex(SpreadsheetExportError, "目标文件未被替换"):
                    export_xlsx(target, "记录", ["值"], [[1]])
            self.assertEqual(target.read_bytes(), original)
            self.assertEqual(list(directory.glob(".existing.*.tmp.xlsx")), [])

            with patch("wanxiang.spreadsheet_export.os.replace", side_effect=PermissionError("合成替换故障")):
                with self.assertRaisesRegex(SpreadsheetExportError, "目标文件未被替换"):
                    export_xlsx(target, "记录", ["值"], [[2]])
            self.assertEqual(target.read_bytes(), original)
            self.assertEqual(list(directory.glob(".existing.*.tmp.xlsx")), [])

    def test_bad_sheet_name_and_target_suffix_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            directory = Path(folder)
            with self.assertRaisesRegex(SpreadsheetExportError, "工作表名称不能包含"):
                export_xlsx(directory / "bad.xlsx", "财务/健身", ["值"], [])
            with self.assertRaisesRegex(SpreadsheetExportError, "必须以 .xlsx 结尾"):
                export_xlsx(directory / "records.csv", "记录", ["值"], [])


if __name__ == "__main__":
    unittest.main()
