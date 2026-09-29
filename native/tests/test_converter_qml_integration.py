from __future__ import annotations

import csv
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import struct
import tempfile
import time
import unittest
import zipfile
from types import SimpleNamespace
from unittest.mock import patch

import numpy as np

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QCoreApplication, QEvent, QObject, QSettings, QUrl, Qt, QMetaObject
from PySide6.QtGui import QAccessible, QImage, QImageReader, QImageWriter
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication
from PySide6.QtQuickControls2 import QQuickStyle
from PySide6.QtQml import QQmlApplicationEngine

from main import (
    ArchiveBridge,
    BackupBridge,
    DailyBridge,
    FinanceBridge,
    FitnessBridge,
    HabitBridge,
    IssuePreferencesBridge,
    MediaBridge,
    MigrationBridge,
    NewsBridge,
    PlannerBridge,
    ReadingBridge,
    ShoppingBridge,
    WeatherBridge,
)
from wanxiang.brand import BrandBridge, BrandRepository
from wanxiang.converter_engines import ConverterEngineUpdateBridge
from wanxiang.converter import (
    ConversionError,
    ConverterBridge,
    IMAGE_FORMATS,
    READ_ONLY_IMAGE_EXTENSIONS,
    AUDIO_INPUT_EXTENSIONS,
    VIDEO_INPUT_EXTENSIONS,
    VIDEO_TARGETS,
    convert_file,
    formats_for,
    _cleanup_abandoned_staging_roots,
    _STAGING_OWNER_FILE,
    _office_engines,
    _office_operation,
    supported_image_extensions,
    writable_image_extensions,
)
from wanxiang.localization import LocalizationBridge, LocalizationRepository
from wanxiang.pdf_ocr import TesseractEngine


class ConverterServiceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])
        cls.app.setQuitOnLastWindowClosed(False)

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory(prefix="wanxiang-converter-service-")
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        self.output = self.root / "converted"
        self.output.mkdir()

    def _wait_for_saves(self, bridge: ConverterBridge, timeout_seconds: float = 8.0) -> list[dict[str, object]]:
        deadline = time.monotonic() + timeout_seconds
        while time.monotonic() < deadline:
            self.app.processEvents()
            jobs = bridge.jobs
            if not any(job.get("savePending") for job in jobs):
                return jobs
            QTest.qWait(20)
        self.fail(f"save jobs did not finish: {bridge.jobs!r}")

    def test_abandoned_staging_sessions_are_removed_only_for_dead_marked_processes(self) -> None:
        temp_root = self.root / "system-temp"
        temp_root.mkdir()
        dead_session = temp_root / "fluke-converter-dead"
        live_session = temp_root / "fluke-converter-live"
        unmarked_session = temp_root / "fluke-converter-unmarked"
        for session in (dead_session, live_session, unmarked_session):
            session.mkdir()
            (session / "result.pdf").write_bytes(b"staged")
        (dead_session / _STAGING_OWNER_FILE).write_text(json.dumps({"pid": 111}), encoding="utf-8")
        (live_session / _STAGING_OWNER_FILE).write_text(json.dumps({"pid": 222}), encoding="utf-8")

        with patch("wanxiang.converter._process_is_running", side_effect=lambda pid: pid == 222):
            _cleanup_abandoned_staging_roots(temp_root)

        self.assertFalse(dead_session.exists())
        self.assertTrue(live_session.joinpath("result.pdf").is_file())
        self.assertTrue(unmarked_session.joinpath("result.pdf").is_file())

    def test_preview_info_classifies_audio_and_video_outputs(self) -> None:
        bridge = ConverterBridge()
        for extension, expected_kind in (("mp3", "audio"), ("mp4", "video")):
            with self.subTest(extension=extension):
                output = self.output / f"preview.{extension}"
                output.write_bytes(b"synthetic media")
                job_id = f"preview-{extension}"
                bridge._jobs.append({
                    "id": job_id,
                    "status": "done",
                    "outputPath": str(output),
                })
                result = bridge.previewInfo(job_id)
                self.assertEqual(result["kind"], expected_kind)
                self.assertTrue(result["url"].startswith("file:///"))

        self.assertIn("mp3", AUDIO_INPUT_EXTENSIONS)
        self.assertIn("mp4", VIDEO_INPUT_EXTENSIONS)

    @staticmethod
    def _csv_rows(path: Path, delimiter: str = ",") -> list[list[str]]:
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            return list(csv.reader(stream, delimiter=delimiter))

    @staticmethod
    def _xlsx_rows(path: Path) -> list[list[object]]:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            return [list(row) for row in workbook.active.iter_rows(values_only=True)]
        finally:
            workbook.close()

    def _read_table(self, path: Path) -> list[list[object]]:
        extension = path.suffix.lower().lstrip(".")
        if extension == "csv":
            return self._csv_rows(path)
        if extension == "tsv":
            return self._csv_rows(path, "\t")
        if extension == "json":
            records = json.loads(path.read_text(encoding="utf-8"))
            headers = list(records[0]) if records else []
            return [headers] + [[record.get(key, "") for key in headers] for record in records]
        return self._xlsx_rows(path)

    def _table_source(self, extension: str) -> Path:
        path = self.root / f"fixture.{extension}"
        rows = [
            ["名称", "备注", "数值"],
            ["合成书目", "含逗号,和制表符\t", "17"],
            ["第二行", "保留原文", "0"],
        ]
        if extension in ("csv", "tsv"):
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                csv.writer(stream, delimiter="\t" if extension == "tsv" else ",").writerows(rows)
        elif extension == "json":
            path.write_text(
                json.dumps(
                    [dict(zip(rows[0], row)) for row in rows[1:]], ensure_ascii=False, indent=2
                ),
                encoding="utf-8",
            )
        else:
            from openpyxl import Workbook

            workbook = Workbook()
            sheet = workbook.active
            for row in rows:
                sheet.append(row)
            workbook.save(path)
        return path

    def test_each_supported_table_format_converts_to_every_other_table_format(self) -> None:
        expected = [
            ["名称", "备注", "数值"],
            ["合成书目", "含逗号,和制表符\t", "17"],
            ["第二行", "保留原文", "0"],
        ]
        extensions = ("csv", "tsv", "json", "xlsx")
        for source_extension in extensions:
            source = self._table_source(source_extension)
            source_before = source.read_bytes()
            for target_extension in extensions:
                if target_extension == source_extension:
                    continue
                with self.subTest(source=source_extension, target=target_extension):
                    self.assertIn(
                        target_extension,
                        {entry["value"] for entry in formats_for(source)},
                    )
                    destination = convert_file(source, target_extension, self.output)
                    self.assertEqual(destination.suffix, f".{target_extension}")
                    self.assertNotEqual(destination.resolve(), source.resolve())
                    self.assertEqual(self._read_table(destination), expected)
                    self.assertEqual(source.read_bytes(), source_before)

    def test_csv_formula_like_text_is_neutralized_but_numbers_are_preserved(self) -> None:
        source = self.root / "formula-fixture.csv"
        source.write_text(
            "value\n=1+1\n  @SUM(A1:A2)\n+SUM(A1:A2)\n-4\nordinary\n",
            encoding="utf-8",
        )
        destination = convert_file(source, "tsv", self.output)
        self.assertEqual(
            self._csv_rows(destination, "\t"),
            [["value"], ["'=1+1"], ["'  @SUM(A1:A2)"], ["'+SUM(A1:A2)"], ["-4"], ["ordinary"]],
        )
        self.assertEqual(source.read_text(encoding="utf-8"), "value\n=1+1\n  @SUM(A1:A2)\n+SUM(A1:A2)\n-4\nordinary\n")

    def test_image_conversion_uses_only_local_reader_writer_intersection(self) -> None:
        readable = {bytes(value).decode("ascii").lower() for value in QImageReader.supportedImageFormats()}
        writable = {bytes(value).decode("ascii").lower() for value in QImageWriter.supportedImageFormats()}
        self.assertEqual(
            supported_image_extensions(),
            {extension for extension, info in IMAGE_FORMATS.items() if info["qt"] in readable},
        )
        self.assertEqual(
            writable_image_extensions(),
            {
                extension for extension, info in IMAGE_FORMATS.items()
                if info["qt"] in writable and extension not in READ_ONLY_IMAGE_EXTENSIONS
            },
        )
        source = self.root / "synthetic-image.png"
        image = QImage(18, 12, QImage.Format.Format_RGB32)
        image.fill(0xFF426A7D)
        self.assertTrue(image.save(str(source), "PNG"))
        source_before = hashlib.sha256(source.read_bytes()).digest()
        available = {item["value"] for item in formats_for(source)} & writable_image_extensions()
        self.assertTrue(available, "the Qt runtime should expose at least one alternate image writer")
        for target in sorted(available):
            with self.subTest(target=target):
                destination = convert_file(source, target, self.output)
                reader = QImageReader(str(destination))
                self.assertTrue(reader.canRead(), f"output {target} could not be decoded")
                image = reader.read()
                self.assertFalse(image.isNull(), f"output {target} did not contain a readable image")
                if target in {"ico", "icns"}:
                    self.assertGreater(image.width() * image.height(), 0)
                elif target != "pdf":
                    self.assertEqual((image.width(), image.height()), (18, 12))
                self.assertEqual(hashlib.sha256(source.read_bytes()).digest(), source_before)

    def test_image_and_pdf_convert_in_both_directions_with_first_page_limit(self) -> None:
        source = self.root / "page.png"
        image = QImage(30, 18, QImage.Format.Format_RGB32)
        image.fill(0xFF315761)
        self.assertTrue(image.save(str(source), "PNG"))
        self.assertIn("pdf", {item["value"] for item in formats_for(source)})
        pdf = convert_file(source, "pdf", self.output)
        self.assertEqual(pdf.suffix, ".pdf")
        if "pdf" not in {bytes(value).decode("ascii").lower() for value in QImageReader.supportedImageFormats()}:
            self.skipTest("this Qt installation does not include the PDF image reader")
        self.assertIn("png", {item["value"] for item in formats_for(pdf)})
        png = convert_file(pdf, "png", self.output)
        reader = QImageReader(str(png))
        self.assertTrue(reader.canRead())
        rendered = reader.read()
        self.assertFalse(rendered.isNull())
        self.assertGreater(rendered.width(), 30, "PDF-to-image exports the rendered first page")

    def test_epub_reads_spine_order_and_round_trips_basic_text_formats(self) -> None:
        from docx import Document
        from pypdf import PdfReader

        source = self.root / "synthetic-book.epub"
        container = (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">'
            '<rootfiles><rootfile full-path="OEBPS/content.opf" '
            'media-type="application/oebps-package+xml"/></rootfiles></container>'
        )
        package = (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" unique-identifier="book-id">'
            '<metadata><dc:identifier id="book-id">urn:synthetic:book</dc:identifier>'
            '<dc:title>合成电子书</dc:title><dc:language>zh-CN</dc:language></metadata>'
            '<manifest><item id="second" href="text/second.xhtml" media-type="application/xhtml+xml"/>'
            '<item id="first" href="text/first.xhtml" media-type="application/xhtml+xml"/></manifest>'
            '<spine><itemref idref="first"/><itemref idref="second"/></spine></package>'
        )
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("mimetype", b"application/epub+zip", compress_type=zipfile.ZIP_STORED)
            archive.writestr("META-INF/container.xml", container)
            archive.writestr("OEBPS/content.opf", package)
            archive.writestr(
                "OEBPS/text/second.xhtml",
                '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>后章</title></head>'
                "<body><h1>第二章</h1><p>按目录顺序读取。</p></body></html>",
            )
            archive.writestr(
                "OEBPS/text/first.xhtml",
                '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>前章</title></head>'
                "<body><h1>第一章</h1><p>电子书正文。</p><script>不应执行</script></body></html>",
            )

        targets = {item["value"] for item in formats_for(source)}
        self.assertTrue({"txt", "md", "html", "docx", "pdf"}.issubset(targets))
        markdown = convert_file(source, "md", self.output).read_text(encoding="utf-8")
        self.assertLess(markdown.index("第一章"), markdown.index("第二章"))
        self.assertIn("电子书正文。", markdown)
        self.assertIn("按目录顺序读取。", markdown)
        self.assertNotIn("不应执行", markdown)
        self.assertGreater(source.stat().st_size, 0)

        pdf = convert_file(source, "pdf", self.output)
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(str(pdf)).pages)
        self.assertIn("第一章", pdf_text)
        self.assertIn("第二章", pdf_text)
        docx = Document(convert_file(source, "docx", self.output))
        self.assertIn("第一章", "\n".join(paragraph.text for paragraph in docx.paragraphs))

        markdown_source = self.root / "guide.md"
        markdown_source.write_text(
            "# 电子书标题\n\n正文段落。\n\n## 小节\n\n更多内容。\n\n# 第二章\n\n章节内容。",
            encoding="utf-8",
        )
        epub_output = convert_file(markdown_source, "epub", self.output)
        from lxml import etree

        with zipfile.ZipFile(epub_output) as archive:
            self.assertEqual(archive.namelist()[0], "mimetype")
            self.assertEqual(archive.getinfo("mimetype").compress_type, zipfile.ZIP_STORED)
            self.assertIsNone(archive.testzip())
            self.assertIn("OEBPS/content.opf", archive.namelist())
            package = etree.fromstring(archive.read("OEBPS/content.opf"))
            chapters = package.xpath("//*[local-name()='spine']/*[local-name()='itemref']")
            self.assertEqual(len(chapters), 2)
            self.assertTrue(package.xpath("//*[local-name()='meta' and @property='dcterms:modified']"))
            for name in archive.namelist():
                if name.endswith((".xml", ".opf", ".xhtml")):
                    etree.fromstring(archive.read(name))
        epub_markdown = convert_file(epub_output, "md", self.output).read_text(encoding="utf-8")
        self.assertIn("电子书标题", epub_markdown)
        self.assertIn("更多内容。", epub_markdown)
        self.assertIn("第二章", epub_markdown)

    def test_epub_rejects_path_traversal_and_invalid_container(self) -> None:
        source = self.root / "unsafe.epub"
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("mimetype", b"application/epub+zip", compress_type=zipfile.ZIP_STORED)
            archive.writestr(
                "META-INF/container.xml",
                '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                '<rootfiles><rootfile full-path="../outside.opf"/></rootfiles></container>',
            )
        self.assertIn("txt", {item["value"] for item in formats_for(source)})
        with self.assertRaisesRegex(ConversionError, "超出压缩包范围"):
            convert_file(source, "txt", self.output)
        self.assertFalse(any(self.output.iterdir()))

    def test_pdf_splits_into_valid_single_page_pdfs_in_a_zip(self) -> None:
        from pypdf import PdfReader, PdfWriter

        source = self.root / "two-pages.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.add_blank_page(width=72, height=72)
        with source.open("wb") as stream:
            writer.write(stream)
        self.assertIn("zip", {item["value"] for item in formats_for(source)})
        archive_path = convert_file(source, "zip", self.output)
        self.assertEqual(archive_path.suffix, ".zip")
        with zipfile.ZipFile(archive_path) as archive:
            self.assertEqual(archive.namelist(), ["page_0001.pdf", "page_0002.pdf"])
            self.assertIsNone(archive.testzip())
            for name in archive.namelist():
                self.assertEqual(len(PdfReader(io.BytesIO(archive.read(name))).pages), 1)

    def test_pdf_splits_into_configurable_consecutive_page_groups(self) -> None:
        from pypdf import PdfReader, PdfWriter

        source = self.root / "five-pages.pdf"
        writer = PdfWriter()
        page_widths = [72, 96, 120, 144, 168]
        for width in page_widths:
            writer.add_blank_page(width=width, height=72)
        with source.open("wb") as stream:
            writer.write(stream)

        grouped = convert_file(source, "zip", self.output, split_pages_per_file=2)
        with zipfile.ZipFile(grouped) as archive:
            names = ["pages_0001-0002.pdf", "pages_0003-0004.pdf", "pages_0005-0005.pdf"]
            self.assertEqual(archive.namelist(), names)
            self.assertIsNone(archive.testzip())
            observed_widths = []
            for name, expected_count in zip(names, (2, 2, 1)):
                pages = PdfReader(io.BytesIO(archive.read(name))).pages
                self.assertEqual(len(pages), expected_count)
                observed_widths.extend(float(page.mediabox.width) for page in pages)
            self.assertEqual(observed_widths, page_widths)

        one_group = convert_file(source, "zip", self.output, split_pages_per_file=10)
        with zipfile.ZipFile(one_group) as archive:
            self.assertEqual(archive.namelist(), ["pages_0001-0005.pdf"])
            self.assertEqual(len(PdfReader(io.BytesIO(archive.read(archive.namelist()[0]))).pages), 5)

        for invalid_size in (0, 2_001, True):
            with self.subTest(invalid_size=invalid_size):
                with self.assertRaises(ConversionError):
                    convert_file(source, "zip", self.output, split_pages_per_file=invalid_size)

    def test_pdf_exports_all_pages_to_a_valid_png_zip(self) -> None:
        from pypdf import PdfReader, PdfWriter

        source = self.root / "three-pages.pdf"
        writer = PdfWriter()
        for width in (72, 144, 216):
            writer.add_blank_page(width=width, height=72)
        with source.open("wb") as stream:
            writer.write(stream)

        self.assertIn("pdf-images-zip", {item["value"] for item in formats_for(source)})
        archive_path = convert_file(source, "pdf-images-zip", self.output)
        self.assertEqual(archive_path.suffix, ".zip")
        self.assertEqual(len(PdfReader(str(source)).pages), 3)
        with zipfile.ZipFile(archive_path) as archive:
            self.assertEqual(archive.namelist(), [
                "page_0001.png", "page_0002.png", "page_0003.png",
            ])
            self.assertIsNone(archive.testzip())
            for name in archive.namelist():
                payload = archive.read(name)
                self.assertTrue(payload.startswith(b"\x89PNG\r\n\x1a\n"))
                image = QImage.fromData(payload)
                self.assertFalse(image.isNull())
                self.assertGreater(image.width(), 0)
                self.assertGreater(image.height(), 0)

    def test_pdf_digital_tables_export_to_separate_safe_xlsx_sheets(self) -> None:
        from openpyxl import load_workbook
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        source = self.root / "digital-tables.pdf"
        writer = PdfWriter()
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        font_ref = writer._add_object(font)
        page_labels = ("alpha", "beta")
        xs = (50, 190, 330)
        ys = (60, 90, 120, 150)
        for label in page_labels:
            page = writer.add_blank_page(width=600, height=300)
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref}),
            })
            operations = ["q", "0.5 w"]
            for x in xs:
                operations.append(f"{x} {ys[0]} m {x} {ys[-1]} l S")
            for y in ys:
                operations.append(f"{xs[0]} {y} m {xs[-1]} {y} l S")
            rows = (("Key", "Value"), (label, "17"), ("formula", "=1+1"))
            for row_index, row in enumerate(rows):
                baseline = 220 - row_index * 32
                for column_index, value in enumerate(row):
                    operations.append(
                        f"BT /F1 12 Tf 1 0 0 1 {xs[column_index] + 8} {baseline} Tm ({value}) Tj ET"
                    )
            operations.append("Q")
            content = DecodedStreamObject()
            content.set_data(("\n".join(operations) + "\n").encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(content)
        with source.open("wb") as stream:
            writer.write(stream)
        source_reader = PdfReader(source)
        self.assertEqual(len(source_reader.pages), 2)
        self.assertIn("=1+1", source_reader.pages[0].extract_text())

        self.assertIn("xlsx", {item["value"] for item in formats_for(source)})
        progress: list[int] = []
        output = convert_file(source, "xlsx", self.output, progress_callback=progress.append)
        self.assertEqual(output.suffix, ".xlsx")
        self.assertEqual(progress[-1], 100)
        workbook = load_workbook(output, data_only=False)
        try:
            self.assertEqual(workbook.sheetnames, [
                "提取说明", "表0001_P0001", "表0002_P0002", "页面正文",
            ])
            self.assertEqual(workbook["提取说明"]["B4"].value, 2)
            self.assertEqual(workbook["表0001_P0001"]["A1"].value, "Key")
            self.assertEqual(workbook["表0001_P0001"]["A2"].value, "alpha")
            formula_text = workbook["表0001_P0001"]["B3"]
            self.assertEqual(formula_text.value, "=1+1")
            self.assertEqual(formula_text.data_type, "s")
            self.assertEqual(workbook["表0002_P0002"]["A2"].value, "beta")
            self.assertTrue(workbook["页面正文"]["C2"].value)
        finally:
            workbook.close()

    def test_pdf_xlsx_restores_only_geometry_proven_merged_cells(self) -> None:
        from openpyxl import load_workbook
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        source = self.root / "merged-grid.pdf"
        writer = PdfWriter()
        page = writer.add_blank_page(width=900, height=300)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        font_ref = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref}),
        })

        operations = ["q", "0.5 w"]
        # First ruled grid: A1:B1 spans horizontally and C1:C2 spans vertically.
        for x in (40, 280):
            operations.append(f"{x} 40 m {x} 190 l S")
        operations.extend((
            "120 40 m 120 140 l S",
            "200 40 m 200 190 l S",
        ))
        operations.extend((
            "40 190 m 280 190 l S",
            "40 140 m 200 140 l S",
            "40 90 m 280 90 l S",
            "40 40 m 280 40 l S",
        ))
        for x, y, value in (
            (50, 165, "Hspan"), (210, 165, "Vspan"),
            (50, 115, "left2"), (130, 115, "right2"),
            (50, 65, "left3"), (130, 65, "mid3"), (210, 65, "right3"),
        ):
            operations.append(f"BT /F1 12 Tf 1 0 0 1 {x} {y} Tm ({value}) Tj ET")

        # A partial internal divider is not enough evidence to merge the top row.
        for x in (360, 560):
            operations.append(f"{x} 40 m {x} 140 l S")
        operations.extend(("460 40 m 460 90 l S", "460 130 m 460 135 l S"))
        for y in (40, 90, 140):
            operations.append(f"360 {y} m 560 {y} l S")
        operations.extend((
            "BT /F1 12 Tf 1 0 0 1 370 115 Tm (partial) Tj ET",
            "BT /F1 12 Tf 1 0 0 1 370 65 Tm (lower1) Tj ET",
            "BT /F1 12 Tf 1 0 0 1 470 65 Tm (lower2) Tj ET",
        ))

        # Third ruled grid has an entirely empty B column; it must stay in place and unmerged.
        for x in (640, 740, 840):
            operations.append(f"{x} 40 m {x} 140 l S")
        for y in (40, 90, 140):
            operations.append(f"640 {y} m 840 {y} l S")
        operations.extend((
            "BT /F1 12 Tf 1 0 0 1 650 115 Tm (Left) Tj ET",
            "BT /F1 12 Tf 1 0 0 1 650 65 Tm (keep) Tj ET",
            "Q",
        ))
        content = DecodedStreamObject()
        content.set_data(("\n".join(operations) + "\n").encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(content)
        with source.open("wb") as stream:
            writer.write(stream)

        output = convert_file(source, "xlsx", self.output)
        workbook = load_workbook(output, data_only=False)
        try:
            table_sheets = [sheet for sheet in workbook.worksheets if sheet.title.startswith("表")]
            merged_sheet = next(sheet for sheet in table_sheets if sheet["A1"].value == "Hspan")
            self.assertEqual(merged_sheet["C1"].value, "Vspan")
            self.assertEqual(merged_sheet["A2"].value, "left2")
            self.assertEqual(merged_sheet["C3"].value, "right3")
            self.assertEqual({str(value) for value in merged_sheet.merged_cells.ranges}, {"A1:B1", "C1:C2"})
            self.assertEqual(merged_sheet["A1"].data_type, "s")

            partial_separator_sheet = next(sheet for sheet in table_sheets if sheet["A1"].value == "partial")
            self.assertEqual(partial_separator_sheet["B2"].value, "lower2")
            self.assertEqual(list(partial_separator_sheet.merged_cells.ranges), [])

            blank_column_sheet = next(sheet for sheet in table_sheets if sheet["A1"].value == "Left")
            self.assertEqual(blank_column_sheet["A2"].value, "keep")
            self.assertIsNone(blank_column_sheet["B1"].value)
            self.assertIsNone(blank_column_sheet["B2"].value)
            self.assertEqual(list(blank_column_sheet.merged_cells.ranges), [])
        finally:
            workbook.close()

    def test_pdf_to_docx_rebuilds_editable_text_at_page_coordinates(self) -> None:
        from docx import Document
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        source = self.root / "layout-source.pdf"
        writer = PdfWriter()
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        font_ref = writer._add_object(font)
        page_specs = (
            (600, 800, ((72, 740, 16, "Layout title"), (120, 700, 10, "Indented body line"))),
            (360, 600, ((36, 540, 12, "Second page body"),)),
        )
        for width, height, text_items in page_specs:
            page = writer.add_blank_page(width=width, height=height)
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref}),
            })
            operations = [
                f"BT /F1 {font_size} Tf 1 0 0 1 {x} {y} Tm ({text}) Tj ET"
                for x, y, font_size, text in text_items
            ]
            content = DecodedStreamObject()
            content.set_data(("\n".join(operations) + "\n").encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(content)
        with source.open("wb") as stream:
            writer.write(stream)

        self.assertIn("docx", {item["value"] for item in formats_for(source)})
        progress: list[int] = []
        output = convert_file(source, "docx", self.output, progress_callback=progress.append)
        self.assertEqual(progress[-1], 100)
        self.assertEqual(len(PdfReader(source).pages), 2)
        document = Document(output)
        try:
            texts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
            self.assertEqual(texts, ["Layout title", "Indented body line", "Second page body"])
            self.assertEqual(len(document.sections), 2)
            self.assertAlmostEqual(document.sections[0].page_width.pt, 600, delta=0.1)
            self.assertAlmostEqual(document.sections[0].page_height.pt, 800, delta=0.1)
            self.assertAlmostEqual(document.sections[1].page_width.pt, 360, delta=0.1)
            self.assertGreater(document.paragraphs[0].paragraph_format.left_indent.pt, 70)
            self.assertAlmostEqual(document.paragraphs[0].runs[0].font.size.pt, 16, delta=0.1)
            self.assertAlmostEqual(document.paragraphs[1].runs[0].font.size.pt, 10, delta=0.1)
        finally:
            del document

    def test_scanned_pdf_page_uses_local_ocr_lines_in_editable_docx(self) -> None:
        import zlib

        from docx import Document
        from pypdf import PdfWriter
        from pypdf.generic import (
            DecodedStreamObject,
            DictionaryObject,
            NameObject,
            NumberObject,
        )
        from wanxiang.pdf_ocr import TesseractEngine

        source = self.root / "scanned-page.pdf"
        writer = PdfWriter()
        page = writer.add_blank_page(width=144, height=144)
        image = DecodedStreamObject()
        image.set_data(zlib.compress(bytes([255, 255, 255]) * 64 * 64))
        image.update({
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Image"),
            NameObject("/Width"): NumberObject(64),
            NameObject("/Height"): NumberObject(64),
            NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
            NameObject("/BitsPerComponent"): NumberObject(8),
            NameObject("/Filter"): NameObject("/FlateDecode"),
        })
        image_ref = writer._add_object(image)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/XObject"): DictionaryObject({NameObject("/Im0"): image_ref}),
        })
        content = DecodedStreamObject()
        content.set_data(b"q 120 0 0 120 12 12 cm /Im0 Do Q")
        page[NameObject("/Contents")] = writer._add_object(content)
        with source.open("wb") as stream:
            writer.write(stream)

        engine = TesseractEngine("local-tesseract", frozenset({"eng"}), "eng")
        with (
            patch("wanxiang.converter.tesseract_engine", return_value=engine),
            patch("wanxiang.converter.recognize_tesseract_tsv", return_value=[{
                "text": "Scanned page text",
                "x0": 30.0,
                "top": 40.0,
                "x1": 220.0,
                "bottom": 65.0,
                "confidence": 82.0,
            }]),
        ):
            output = convert_file(source, "docx", self.output)
        document = Document(output)
        try:
            self.assertEqual([paragraph.text for paragraph in document.paragraphs], ["Scanned page text"])
            self.assertIn("average confidence 82%", document.core_properties.subject)
        finally:
            del document

    def test_tesseract_tsv_parser_combines_words_and_keeps_confidence(self) -> None:
        from wanxiang.pdf_ocr import TesseractEngine, recognize_tesseract_tsv

        tsv = (
            "level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n"
            "5\t1\t1\t1\t1\t1\t10\t20\t30\t10\t80\tHello\n"
            "5\t1\t1\t1\t1\t2\t42\t19\t28\t12\t60\tworld\n"
        ).encode("utf-8")
        engine = TesseractEngine("tesseract.exe", frozenset({"eng"}), "eng")
        result = subprocess.CompletedProcess(args=[], returncode=0, stdout=tsv, stderr=b"")
        with patch("wanxiang.pdf_ocr.subprocess.run", return_value=result) as run:
            lines = recognize_tesseract_tsv(engine, self.root / "page.png", 10)
        self.assertEqual(len(lines), 1)
        self.assertEqual(lines[0]["text"], "Hello world")
        self.assertEqual(lines[0]["x0"], 10)
        self.assertEqual(lines[0]["top"], 19)
        self.assertEqual(lines[0]["x1"], 70)
        self.assertEqual(lines[0]["bottom"], 31)
        self.assertEqual(lines[0]["confidence"], 70)
        self.assertEqual([word["text"] for word in lines[0]["words"]], ["Hello", "world"])
        self.assertIn("tsv", run.call_args.args[0])

    def test_ocr_table_detection_requires_repeated_aligned_columns(self) -> None:
        from wanxiang.converter import _ocr_page_tables

        lines = []
        for row_index, (left, right) in enumerate((("Name", "Score"), ("Ada", "91"), ("Lin", "87"))):
            top = 40 + row_index * 40
            lines.append({
                "text": f"{left} {right}",
                "words": [
                    {"text": left, "x0": 40, "top": top, "x1": 90, "bottom": top + 20, "confidence": 92},
                    {"text": right, "x0": 250, "top": top, "x1": 300, "bottom": top + 20, "confidence": 74},
                ],
            })
        lines.append({
            "text": "A normal paragraph with several words",
            "words": [
                {"text": "A", "x0": 30, "top": 400, "x1": 40, "bottom": 420, "confidence": 90},
                {"text": "normal", "x0": 48, "top": 400, "x1": 100, "bottom": 420, "confidence": 90},
                {"text": "paragraph", "x0": 108, "top": 400, "x1": 180, "bottom": 420, "confidence": 90},
            ],
        })

        tables = _ocr_page_tables(lines, 600)
        self.assertEqual(len(tables), 1)
        self.assertEqual(
            [[cell["text"] for cell in row] for row in tables[0]["rows"]],
            [["Name", "Score"], ["Ada", "91"], ["Lin", "87"]],
        )
        self.assertEqual(tables[0]["rows"][1][1]["confidence"], 74)

    def test_ocr_table_detection_retains_aligned_single_cell_rows(self) -> None:
        from wanxiang.converter import _ocr_page_tables

        lines = []
        for row_index, (left, right) in enumerate((("Name", "Score"), ("Ada", "91"), ("Lin", "87"))):
            top = 40 + row_index * 40
            lines.append({
                "text": f"{left} {right}",
                "words": [
                    {"text": left, "x0": 40, "top": top, "x1": 90, "bottom": top + 20, "confidence": 92},
                    {"text": right, "x0": 250, "top": top, "x1": 300, "bottom": top + 20, "confidence": 74},
                ],
            })
        lines.append({
            "text": "Merged note",
            "words": [
                {"text": "Merged", "x0": 40, "top": 100, "x1": 105, "bottom": 120, "confidence": 68},
                {"text": "note", "x0": 112, "top": 100, "x1": 150, "bottom": 120, "confidence": 68},
            ],
        })

        tables = _ocr_page_tables(lines, 600)
        self.assertEqual(len(tables), 1)
        self.assertEqual(
            [[cell["text"] for cell in row] for row in tables[0]["rows"]],
            [["Name", "Score"], ["Ada", "91"], ["Merged note", ""], ["Lin", "87"]],
        )

    def test_scanned_pdf_xlsx_ocr_adds_confidence_and_joins_adjacent_pages(self) -> None:
        import zlib

        from openpyxl import load_workbook
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject
        from wanxiang.pdf_ocr import TesseractEngine

        source = self.root / "scanned-tables.pdf"
        writer = PdfWriter()
        for _ in range(2):
            page = writer.add_blank_page(width=144, height=144)
            image = DecodedStreamObject()
            image.set_data(zlib.compress(bytes([255, 255, 255]) * 64 * 64))
            image.update({
                NameObject("/Type"): NameObject("/XObject"),
                NameObject("/Subtype"): NameObject("/Image"),
                NameObject("/Width"): NumberObject(64),
                NameObject("/Height"): NumberObject(64),
                NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
                NameObject("/BitsPerComponent"): NumberObject(8),
                NameObject("/Filter"): NameObject("/FlateDecode"),
            })
            image_ref = writer._add_object(image)
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/XObject"): DictionaryObject({NameObject("/Im0"): image_ref}),
            })
            content = DecodedStreamObject()
            content.set_data(b"q 120 0 0 120 12 12 cm /Im0 Do Q")
            page[NameObject("/Contents")] = writer._add_object(content)
        with source.open("wb") as stream:
            writer.write(stream)

        def recognized_page(
            rows: tuple[tuple[str, str, int], ...],
            include_single_cell_row: bool = False,
        ) -> list[dict[str, object]]:
            result = []
            for row_index, (left, right, confidence) in enumerate(rows):
                top = 100 + row_index * 40
                result.append({
                    "text": f"{left} {right}",
                    "x0": 100,
                    "top": top,
                    "x1": 340,
                    "bottom": top + 20,
                    "confidence": confidence,
                    "words": [
                        {"text": left, "x0": 100, "top": top, "x1": 160, "bottom": top + 20, "confidence": confidence},
                        {"text": right, "x0": 300, "top": top, "x1": 350, "bottom": top + 20, "confidence": confidence},
                    ],
                })
            if include_single_cell_row:
                result.append({
                    "text": "Merged note",
                    "x0": 100,
                    "top": 120,
                    "x1": 188,
                    "bottom": 140,
                    "confidence": 68,
                    "words": [
                        {"text": "Merged", "x0": 100, "top": 120, "x1": 150, "bottom": 140, "confidence": 68},
                        {"text": "note", "x0": 154, "top": 120, "x1": 188, "bottom": 140, "confidence": 68},
                    ],
                })
            return result

        mocked_pages = [
            recognized_page((("Name", "Score", 95), ("Ada", "91", 85)), include_single_cell_row=True),
            recognized_page((("Name", "Score", 95), ("Lin", "87", 54))),
        ]
        engine = TesseractEngine("local-tesseract", frozenset({"eng"}), "eng")
        with (
            patch("wanxiang.converter.tesseract_engine", return_value=engine),
            patch("wanxiang.converter.recognize_tesseract_tsv", side_effect=mocked_pages) as recognize,
        ):
            output = convert_file(source, "xlsx", self.output)

        workbook = load_workbook(output, data_only=False)
        try:
            self.assertIn("页面正文", workbook.sheetnames)
            table_names = [name for name in workbook.sheetnames if name.startswith("OCR表")]
            self.assertEqual(len(table_names), 1)
            table = workbook[table_names[0]]
            self.assertEqual(
                [[table.cell(row=row, column=column).value for column in (1, 2)] for row in range(1, 5)],
                [["Name", "Score"], ["Merged note", None], ["Ada", "91"], ["Lin", "87"]],
            )
            self.assertIsNone(table["B2"].comment)
            self.assertIn("置信度：54%", table["B4"].comment.text)
            self.assertEqual(table["B4"].fill.fgColor.rgb, "00FFF2CC")
            self.assertEqual(workbook["提取说明"]["B4"].value, 1)
            self.assertEqual(recognize.call_count, 2)
        finally:
            workbook.close()

    def test_pdf_xlsx_finds_ruled_and_text_aligned_tables_on_separate_pages(self) -> None:
        from openpyxl import load_workbook
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        source = self.root / "mixed-tables.pdf"
        writer = PdfWriter()
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        font_ref = writer._add_object(font)

        def add_page(operations: list[str]) -> None:
            page = writer.add_blank_page(width=600, height=300)
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref}),
            })
            content = DecodedStreamObject()
            content.set_data(("\n".join(operations) + "\n").encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(content)

        ruled_ops = ["q", "0.5 w"]
        xs = (45, 165, 285)
        ys = (60, 90, 120, 150)
        for x in xs:
            ruled_ops.append(f"{x} {ys[0]} m {x} {ys[-1]} l S")
        for y in ys:
            ruled_ops.append(f"{xs[0]} {y} m {xs[-1]} {y} l S")
        for row_index, row in enumerate((("Key", "Value"), ("ruled", "17"), ("last", "18"))):
            for column_index, value in enumerate(row):
                ruled_ops.append(
                    f"BT /F1 12 Tf 1 0 0 1 {xs[column_index] + 8} {ys[row_index] + 8} Tm ({value}) Tj ET"
                )
        ruled_ops.append("Q")
        aligned_ops: list[str] = []
        text_rows = (("Name", "Score"), ("plain", "27"), ("second", "28"))
        for row_index, row in enumerate(text_rows):
            y = (68, 98, 128)[row_index]
            aligned_ops.append(f"BT /F1 12 Tf 1 0 0 1 350 {y} Tm ({row[0]}) Tj ET")
            aligned_ops.append(f"BT /F1 12 Tf 1 0 0 1 475 {y} Tm ({row[1]}) Tj ET")
        add_page(ruled_ops + aligned_ops)
        with source.open("wb") as stream:
            writer.write(stream)
        reader = PdfReader(source)
        self.assertEqual(len(reader.pages), 1)
        self.assertIn("ruled", reader.pages[0].extract_text())
        self.assertIn("plain", reader.pages[0].extract_text())

        output = convert_file(source, "xlsx", self.output)
        workbook = load_workbook(output, data_only=False)
        try:
            table_sheets = [name for name in workbook.sheetnames if name.startswith("表")]
            self.assertEqual(len(table_sheets), 2)
            self.assertTrue(any(workbook[name]["A2"].value == "ruled" for name in table_sheets))
            self.assertTrue(any(workbook[name]["A2"].value == "plain" for name in table_sheets))
        finally:
            workbook.close()

    def test_bridge_keeps_pdf_split_size_private_and_resets_it_on_target_change(self) -> None:
        from pypdf import PdfWriter

        source = self.root / "bridge-split.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        with source.open("wb") as stream:
            writer.write(stream)

        bridge = ConverterBridge()
        bridge.addFiles(QUrl.fromLocalFile(str(source)))
        job_id = str(bridge.jobs[0]["id"])
        bridge.setTargetFormat(job_id, "zip")
        bridge.setPdfSplitGroupSize(job_id, 4)
        self.assertEqual(bridge.pdfSplitGroupSizeForJob(job_id), 4)
        self.assertNotIn("splitPagesPerFile", bridge.jobs[0])
        bridge.setPdfSplitGroupSize(job_id, 0)
        self.assertEqual(bridge.pdfSplitGroupSizeForJob(job_id), 4)
        bridge.setTargetFormat(job_id, "pdf-encrypt")
        self.assertEqual(bridge.pdfSplitGroupSizeForJob(job_id), 1)

    def test_bridge_converts_direct_folder_images_in_natural_order(self) -> None:
        folder = self.root / "photo-set"
        folder.mkdir()
        nested = folder / "nested"
        nested.mkdir()
        for name, color in (("page10.png", 0xFFAA2211), ("page2.png", 0xFF1166AA)):
            image = QImage(24, 32, QImage.Format.Format_ARGB32)
            image.fill(color)
            self.assertTrue(image.save(str(folder / name), "PNG"))
        nested_image = QImage(24, 32, QImage.Format.Format_ARGB32)
        nested_image.fill(0xFF55AA33)
        self.assertTrue(nested_image.save(str(nested / "nested.png"), "PNG"))
        (folder / "notes.txt").write_text("ignored", encoding="utf-8")

        bridge = ConverterBridge()
        result = bridge.addImageFolderToPdf(QUrl.fromLocalFile(str(folder)))
        self.assertEqual(result, {"ok": True, "count": 2, "skipped": 0})
        job = bridge.jobs[0]
        job_id = str(job["id"])
        self.assertEqual(job["fileName"], "photo-set · 2 张图片")
        self.assertEqual(
            [Path(path).name for path in bridge._image_merge_sources[job_id]],
            ["page2.png", "page10.png"],
        )

        deadline = time.monotonic() + 8
        while time.monotonic() < deadline:
            self.app.processEvents()
            job = bridge.jobs[0]
            if job["status"] in ("done", "failed"):
                break
            QTest.qWait(20)
        self.assertEqual(job["status"], "done", str(job.get("error", "")))
        staged = Path(str(job["stagedPath"]))
        self.assertEqual(staged.name, "photo-set_converted.pdf")
        self.assertGreater(staged.stat().st_size, 0)
        from pypdf import PdfReader
        with PdfReader(staged) as reader:
            self.assertEqual(len(reader.pages), 2)
        self.assertTrue((folder / "page2.png").is_file())
        self.assertTrue((folder / "page10.png").is_file())

    def test_common_target_format_applies_to_ready_convertible_jobs(self) -> None:
        sources = []
        for name in ("first.csv", "second.tsv"):
            source = self.root / name
            source.write_text("name,value\nsynthetic,1\n", encoding="utf-8")
            sources.append(source)
        unsupported = self.root / "unsupported.bin"
        unsupported.write_bytes(b"not a supported file")

        bridge = ConverterBridge()
        bridge.addFiles([*(str(path) for path in sources), str(unsupported)])
        self.assertTrue(bridge.commonTargetFormats)
        common_values = {item["value"] for item in bridge.commonTargetFormats}
        self.assertIn("html", common_values)
        self.assertTrue(all("availableFormats" in job for job in bridge.jobs))

        before = [job["targetFormat"] for job in bridge.jobs]
        invalid = bridge.setCommonTargetFormat("not-a-target")
        self.assertFalse(invalid["ok"])
        self.assertEqual([job["targetFormat"] for job in bridge.jobs], before)

        result = bridge.setCommonTargetFormat("html")
        self.assertEqual(result, {"ok": True, "applied": 2, "changed": True})
        self.assertEqual([job["targetFormat"] for job in bridge.jobs[:2]], ["html", "html"])
        self.assertEqual(bridge.jobs[2]["status"], "unsupported")
        self.assertFalse(bridge.rememberTargetFormats)
        self.assertEqual(bridge._remembered_targets, {})

        unchanged = bridge.setCommonTargetFormat("html")
        self.assertEqual(unchanged, {"ok": True, "applied": 2, "changed": False})

        bridge.setRememberTargetFormats(True)
        bridge.setCommonTargetFormat("html")
        self.assertTrue(bridge.rememberTargetFormats)
        self.assertEqual(bridge._remembered_targets["csv"], "html")
        self.assertEqual(bridge._remembered_targets["tsv"], "html")

        bridge.setRememberTargetFormats(False)
        self.assertFalse(bridge.rememberTargetFormats)
        self.assertEqual(bridge._remembered_targets, {})

    def test_target_format_preference_is_opt_in_across_reopen(self) -> None:
        settings = QSettings(str(self.root / "converter-settings.ini"), QSettings.IniFormat)
        first = self.root / "first.csv"
        first.write_text("name,value\nfirst,1\n", encoding="utf-8")

        bridge = ConverterBridge(settings)
        bridge.addFiles([str(first)])
        bridge.setTargetFormat(bridge.jobs[0]["id"], "html")
        settings.sync()

        later_without_preference = self.root / "later-without-preference.csv"
        later_without_preference.write_text("name,value\nlater,2\n", encoding="utf-8")
        reopened_without_preference = ConverterBridge(settings)
        reopened_without_preference.addFiles([str(later_without_preference)])
        self.assertFalse(reopened_without_preference.rememberTargetFormats)
        self.assertEqual(reopened_without_preference.jobs[0]["targetFormat"], "json")

        bridge.setRememberTargetFormats(True)
        bridge.setTargetFormat(bridge.jobs[0]["id"], "html")
        settings.sync()

        later_with_preference = self.root / "later-with-preference.csv"
        later_with_preference.write_text("name,value\nlater,3\n", encoding="utf-8")
        reopened_with_preference = ConverterBridge(settings)
        reopened_with_preference.addFiles([str(later_with_preference)])
        self.assertTrue(reopened_with_preference.rememberTargetFormats)
        self.assertEqual(reopened_with_preference.jobs[0]["targetFormat"], "html")

    def test_pdf_aes256_encryption_and_original_password_decryption(self) -> None:
        from pypdf import PdfReader, PdfWriter

        source = self.root / "security.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.add_blank_page(width=144, height=144)
        with source.open("wb") as stream:
            writer.write(stream)

        available = {item["value"] for item in formats_for(source)}
        self.assertIn("pdf-encrypt", available)
        self.assertNotIn("pdf-decrypt", available)
        password = "sample-password-256"
        encrypted = convert_file(source, "pdf-encrypt", self.output, password)
        self.assertEqual(encrypted.suffix, ".pdf")
        encrypted_reader = PdfReader(str(encrypted))
        self.assertTrue(encrypted_reader.is_encrypted)
        self.assertTrue(encrypted_reader.decrypt(password))
        self.assertEqual(len(encrypted_reader.pages), 2)
        self.assertEqual(
            {item["value"] for item in formats_for(encrypted)},
            {"pdf-decrypt"},
        )

        with self.assertRaisesRegex(ConversionError, "密码不正确"):
            convert_file(encrypted, "pdf-decrypt", self.output, "wrong-password")
        with self.assertRaisesRegex(ConversionError, "至少需要 8 个字符"):
            convert_file(source, "pdf-encrypt", self.output, "short")
        with self.assertRaisesRegex(ConversionError, "请输入 PDF 密码"):
            convert_file(source, "pdf-encrypt", self.output)

        decrypted = convert_file(encrypted, "pdf-decrypt", self.output, password)
        decrypted_reader = PdfReader(str(decrypted))
        self.assertFalse(decrypted_reader.is_encrypted)
        self.assertEqual(len(decrypted_reader.pages), 2)
        self.assertFalse(any(path.name.endswith(".tmp") for path in self.output.iterdir()))

        bridge = ConverterBridge()
        bridge.addFiles(str(source))
        job_id = str(bridge.jobs[0]["id"])
        bridge.setTargetFormat(job_id, "pdf-encrypt")
        bridge.setJobPassword(job_id, password)
        self.assertEqual(bridge.passwordForJob(job_id), password)
        self.assertNotIn(password, repr(bridge.jobs))
        self.assertNotIn("password", bridge.jobs[0])
        bridge.convertReady()
        self.assertTrue(bridge._pool.waitForDone(10_000))
        self.app.processEvents()
        self.assertEqual(bridge.jobs[0]["status"], "done")
        self.assertEqual(bridge.passwordForJob(job_id), "")
        self.assertFalse(bridge._job_passwords)
        bridge.removeJob(job_id)
        self.assertEqual(bridge.passwordForJob(job_id), "")

    def test_queued_pdf_merge_preserves_queue_order_and_originals(self) -> None:
        from pypdf import PdfReader, PdfWriter

        sources = [self.root / "first.pdf", self.root / "second.pdf"]
        page_sizes = [(72, 72), (144, 144)]
        for source, (width, height) in zip(sources, page_sizes):
            writer = PdfWriter()
            writer.add_blank_page(width=width, height=height)
            with source.open("wb") as stream:
                writer.write(stream)
        originals = [source.read_bytes() for source in sources]

        bridge = ConverterBridge()
        bridge.addFiles([str(source) for source in sources])
        bridge.mergeQueuedPdfs()
        self.assertEqual(len(bridge.jobs), 3)
        self.assertEqual(bridge.jobs[2]["taskType"], "pdf-merge")
        self.assertTrue(bridge._pool.waitForDone(10_000))
        self.app.processEvents()
        self.assertEqual(bridge.saveAll(), {"ok": True, "scheduled": 1, "failed": 0})
        self._wait_for_saves(bridge)
        merge_job = bridge.jobs[2]
        self.assertEqual(merge_job["status"], "done")
        merged = PdfReader(merge_job["outputPath"])
        self.assertEqual(len(merged.pages), 2)
        self.assertEqual([float(page.mediabox.width) for page in merged.pages], [72, 144])
        self.assertEqual([source.read_bytes() for source in sources], originals)

    def test_explicit_pdf_merge_order_is_validated_and_preserved(self) -> None:
        from pypdf import PdfReader, PdfWriter

        sources = [self.root / "first.pdf", self.root / "second.pdf"]
        for source, width in zip(sources, (72, 144)):
            writer = PdfWriter()
            writer.add_blank_page(width=width, height=width)
            with source.open("wb") as stream:
                writer.write(stream)
        originals = [source.read_bytes() for source in sources]

        bridge = ConverterBridge()
        bridge.addFiles([str(source) for source in sources])
        job_ids = [str(job["id"]) for job in bridge.jobs]
        rejected = bridge.mergePdfJobs([job_ids[0], "stale-job-id"])
        self.assertFalse(rejected["ok"])
        self.assertEqual(len(bridge.jobs), 2)

        started = bridge.mergePdfJobs([job_ids[1], job_ids[0]])
        self.assertTrue(started["ok"], started)
        self.assertTrue(bridge._pool.waitForDone(10_000))
        self.app.processEvents()
        merge_job = bridge.jobs[2]
        self.assertEqual(merge_job["status"], "done", merge_job["error"])
        self.assertEqual(bridge.saveJob(merge_job["id"]), {"ok": True, "pending": True})
        self._wait_for_saves(bridge)
        merge_job = bridge.jobs[2]
        merged = PdfReader(merge_job["outputPath"])
        self.assertEqual([float(page.mediabox.width) for page in merged.pages], [144, 72])
        self.assertEqual([source.read_bytes() for source in sources], originals)

    def test_merge_candidates_report_excluded_pdf_reasons(self) -> None:
        from pypdf import PdfWriter

        valid_sources = [self.root / "valid-first.pdf", self.root / "valid-second.pdf"]
        for source in valid_sources:
            writer = PdfWriter()
            writer.add_blank_page(width=72, height=72)
            with source.open("wb") as stream:
                writer.write(stream)

        invalid = self.root / "invalid.pdf"
        invalid.write_bytes(b"not a PDF")
        encrypted = self.root / "encrypted.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.encrypt("sample-password")
        with encrypted.open("wb") as stream:
            writer.write(stream)

        bridge = ConverterBridge()
        bridge.addFiles([str(path) for path in (*valid_sources, invalid, encrypted)])
        preview = bridge.pdfMergeCandidates()

        self.assertEqual(
            {item["fileName"] for item in preview["items"]},
            {path.name for path in valid_sources},
        )
        self.assertEqual(preview["skipped"], 2)
        excluded = {item["fileName"]: item["reasonCode"] for item in preview["excluded"]}
        self.assertEqual(excluded[invalid.name], "invalid")
        self.assertEqual(excluded[encrypted.name], "decrypt_first")

    def test_docx_pdf_and_text_conversions_preserve_text_and_report_limits(self) -> None:
        from docx import Document
        from pypdf import PdfReader

        source = self.root / "notes.docx"
        document = Document()
        document.add_heading("合成标题", level=1)
        document.add_paragraph("保留的正文内容。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "姓名"
        table.cell(0, 1).text = "分数"
        table.cell(1, 0).text = "样本"
        table.cell(1, 1).text = "95"
        document.save(source)
        targets = {item["value"] for item in formats_for(source)}
        self.assertTrue({"txt", "md", "html", "pdf"}.issubset(targets))

        plain = convert_file(source, "txt", self.output).read_text(encoding="utf-8")
        self.assertIn("合成标题", plain)
        self.assertNotIn("# 合成标题", plain)
        self.assertIn("姓名\t分数", plain)
        markdown = convert_file(source, "md", self.output).read_text(encoding="utf-8")
        self.assertIn("# 合成标题", markdown)
        html = convert_file(source, "html", self.output).read_text(encoding="utf-8")
        self.assertIn("<table>", html)
        self.assertIn("<h1>合成标题</h1>", html)

        pdf = convert_file(source, "pdf", self.output)
        extracted = "\n".join(page.extract_text() or "" for page in PdfReader(str(pdf)).pages)
        self.assertIn("合成标题", extracted)
        round_trip = convert_file(pdf, "docx", self.output)
        self.assertIn("保留的正文内容", "\n".join(p.text for p in Document(round_trip).paragraphs))
        self.assertIn("png", {item["value"] for item in formats_for(pdf)})
        self.assertIn("epub", {item["value"] for item in formats_for(source)})
        ebook = convert_file(source, "epub", self.output)
        ebook_text = convert_file(ebook, "md", self.output).read_text(encoding="utf-8")
        self.assertIn("合成标题", ebook_text)
        self.assertIn("保留的正文内容", ebook_text)

        if "pdf" in {bytes(value).decode("ascii").lower() for value in QImageReader.supportedImageFormats()}:
            scanned = self.root / "image-only.png"
            picture = QImage(32, 18, QImage.Format.Format_RGB32)
            picture.fill(0xFF426A7D)
            self.assertTrue(picture.save(str(scanned), "PNG"))
            scanned_pdf = convert_file(scanned, "pdf", self.output)
            with self.assertRaisesRegex(ConversionError, "扫描件需要 OCR"):
                convert_file(scanned_pdf, "txt", self.output)

    def test_text_markdown_html_and_docx_conversions(self) -> None:
        from docx import Document

        source = self.root / "guide.md"
        source.write_text("# 合成指南\n\n正文 **加粗**，包含 `code`。\n\n- 一项\n- 二项\n", encoding="utf-8")
        html = convert_file(source, "html", self.output).read_text(encoding="utf-8")
        self.assertIn("<h1>合成指南</h1>", html)
        self.assertIn("<strong>加粗</strong>", html)
        self.assertIn("<code>code</code>", html)
        docx = convert_file(source, "docx", self.output)
        self.assertEqual(Document(docx).paragraphs[0].style.name, "Heading 1")
        plain = self.root / "page.html"
        plain.write_text("<h1>标题</h1><p>正文</p><script>alert('hidden')</script>", encoding="utf-8")
        text = convert_file(plain, "txt", self.output).read_text(encoding="utf-8")
        self.assertIn("标题", text)
        self.assertIn("正文", text)
        self.assertNotIn("alert", text)
        utf16 = self.root / "utf16.txt"
        utf16.write_bytes("\ufeffUTF-16 正文".encode("utf-16-le"))
        utf16_text = convert_file(utf16, "md", self.output).read_text(encoding="utf-8")
        self.assertEqual(utf16_text, "UTF-16 正文")

    def test_markdown_tables_nested_lists_code_and_inline_formatting_survive_exports(self) -> None:
        from docx import Document
        from pypdf import PdfReader

        source = self.root / "structured-guide.md"
        source.write_text(
            "# 结构指南\n\n"
            "正文 **加粗**、*斜体*、`inline`、[链接](https://example.com)、"
            "[危险链接](javascript:alert(1))、![远程图片](https://example.com/p.png) 和 $E=mc^2$。\n\n"
            "| 字段 | 说明 |\n| --- | --- |\n| 格式 | 表格内容 |\n| 公式 | $E=mc^2$ |\n\n"
            "1. 第一层\n   - 嵌套项\n2. 第二条\n\n"
            "> 引用内容\n\n"
            "```python\nprint(\"code\")\n```\n\n"
            "<script>alert('blocked')</script>\n",
            encoding="utf-8",
        )

        html = convert_file(source, "html", self.output).read_text(encoding="utf-8")
        self.assertIn("<table>", html)
        self.assertIn("<ol>", html)
        self.assertIn("<ul>", html)
        self.assertIn("<blockquote>", html)
        self.assertIn('<pre><code class="language-python">', html)
        self.assertIn("<strong>加粗</strong>", html)
        self.assertIn("<em>斜体</em>", html)
        self.assertIn("&lt;script&gt;", html)
        self.assertIn("远程图片", html)
        self.assertNotIn("<script>", html)
        self.assertNotIn('<img src=', html)
        self.assertNotIn('href="javascript:', html)

        docx_path = convert_file(source, "docx", self.output)
        document = Document(docx_path)
        self.assertEqual(document.paragraphs[0].style.name, "Heading 1")
        self.assertTrue(any(paragraph.style.name == "List Number" and "第一层" in paragraph.text for paragraph in document.paragraphs))
        self.assertTrue(any(paragraph.style.name == "List Bullet" and "嵌套项" in paragraph.text for paragraph in document.paragraphs))
        self.assertTrue(any(paragraph.style.name == "Quote" and "引用内容" in paragraph.text for paragraph in document.paragraphs))
        self.assertTrue(any(run.bold and run.text == "加粗" for paragraph in document.paragraphs for run in paragraph.runs))
        self.assertTrue(any(run.italic and run.text == "斜体" for paragraph in document.paragraphs for run in paragraph.runs))
        self.assertTrue(any(run.font.name == "Consolas" and "print" in run.text for paragraph in document.paragraphs for run in paragraph.runs))
        self.assertEqual(document.tables[0].cell(1, 0).text, "格式")
        self.assertEqual(document.tables[0].cell(1, 1).text, "表格内容")
        self.assertTrue(any(relation.reltype.endswith("/hyperlink") and relation.target_ref == "https://example.com" for relation in document.part.rels.values()))

        pdf_path = convert_file(source, "pdf", self.output)
        extracted = "\n".join(page.extract_text() or "" for page in PdfReader(str(pdf_path)).pages)
        for expected in ("结构指南", "格式", "表格内容", "第一层", "嵌套项", "第二条", "引用内容", "print", "E=mc^2"):
            self.assertIn(expected, extracted)

        from zipfile import ZipFile

        epub_path = convert_file(source, "epub", self.output)
        with ZipFile(epub_path) as archive:
            chapter_name = next(name for name in archive.namelist() if name.startswith("OEBPS/chapter-") and name.endswith(".xhtml"))
            chapter = archive.read(chapter_name).decode("utf-8")
        self.assertIn("<table>", chapter)
        self.assertIn("<ol>", chapter)
        self.assertIn("<ul>", chapter)
        self.assertIn("<pre><code", chapter)

    def test_log_xml_yaml_and_yml_inputs_use_safe_text_conversion(self) -> None:
        contents = {
            "log": "2026-09-28 ERROR: <script>alert('log')</script>\n",
            "xml": "<root><value>XML 内容</value></root>\n",
            "yaml": "name: YAML 内容\nvalue: <script>\n",
            "yml": "name: YML 内容\nvalue: <script>\n",
        }
        for extension, content in contents.items():
            with self.subTest(extension=extension):
                source = self.root / f"sample.{extension}"
                source.write_text(content, encoding="utf-8")
                targets = {item["value"] for item in formats_for(source)}
                self.assertTrue({"txt", "md", "html", "pdf"}.issubset(targets))
                self.assertEqual(
                    convert_file(source, "txt", self.output).read_text(encoding="utf-8"),
                    content,
                )
                html = convert_file(source, "html", self.output).read_text(encoding="utf-8")
                self.assertIn("&lt;", html)
                self.assertNotIn("<script>", html)

    def test_table_formats_export_to_html_docx_and_pdf(self) -> None:
        from docx import Document
        from pypdf import PdfReader

        source = self.root / "table.csv"
        source.write_text("名称,数量\n样本,3\n第二项,8\n", encoding="utf-8-sig")
        formats = {item["value"] for item in formats_for(source)}
        self.assertTrue({"html", "docx", "pdf"}.issubset(formats))
        html = convert_file(source, "html", self.output).read_text(encoding="utf-8")
        self.assertIn("<th>名称</th>", html)
        self.assertIn("<td>第二项</td>", html)
        docx = Document(convert_file(source, "docx", self.output))
        self.assertEqual(docx.tables[0].cell(1, 0).text, "样本")
        pdf = convert_file(source, "pdf", self.output)
        extracted = "\n".join(page.extract_text() or "" for page in PdfReader(str(pdf)).pages)
        self.assertIn("名称", extracted)
        self.assertIn("第二项", extracted)

        workbook_source = self._table_source("xlsx")
        from openpyxl import load_workbook

        workbook = load_workbook(workbook_source)
        second_sheet = workbook.create_sheet("第二工作表")
        second_sheet["A1"] = "第二工作表专用内容"
        workbook.save(workbook_source)
        workbook.close()
        excel_pdf = convert_file(workbook_source, "pdf", self.output)
        excel_text = "\n".join(page.extract_text() or "" for page in PdfReader(str(excel_pdf)).pages)
        if "excel" in _office_engines():
            self.assertIn("第二工作表专用内容", excel_text)

    def test_legacy_excel_and_wps_sheet_targets_are_discoverable(self) -> None:
        with patch("wanxiang.converter._office_engines", return_value=frozenset({"excel"})), patch(
            "wanxiang.converter._wps_engines", return_value=frozenset(),
        ):
            office_targets = formats_for(self.root / "legacy.xls")
        self.assertEqual([item["value"] for item in office_targets], ["xlsx", "csv", "ods", "pdf"])
        self.assertEqual([item["engine"] for item in office_targets], ["excel"] * 4)
        self.assertIn("Microsoft Excel", office_targets[0]["label"])
        self.assertIn("Microsoft Excel", office_targets[1]["label"])
        self.assertIn("Microsoft Excel", office_targets[2]["label"])

        with patch("wanxiang.converter._office_engines", return_value=frozenset({"excel"})):
            tabular_targets = {item["value"] for item in formats_for(self.root / "table.xlsx")}
        self.assertTrue({"xls", "ods"}.issubset(tabular_targets))
        self.assertNotIn("xls", {item["value"] for item in office_targets})
        self.assertEqual(_office_operation("xlsx", "xls"), "excel-xls")
        self.assertEqual(
            _office_operation("xlsx", "ods"),
            "excel-ods",
        )

        with patch("wanxiang.converter._office_engines", return_value=frozenset()), patch(
            "wanxiang.converter._wps_engines", return_value=frozenset({"excel"}),
        ):
            wps_targets = formats_for(self.root / "sheet.et")
        self.assertEqual([item["value"] for item in wps_targets], ["xlsx", "csv", "pdf"])
        self.assertEqual([item["engine"] for item in wps_targets], ["wps-excel"] * 3)
        self.assertIn("WPS 表格", wps_targets[0]["label"])
        self.assertIn("WPS 表格", wps_targets[1]["label"])

    def test_excel_can_import_and_export_ods_when_office_is_installed(self) -> None:
        from openpyxl import Workbook, load_workbook

        if "excel" not in _office_engines() or os.name != "nt":
            self.skipTest("Microsoft Excel COM is not registered on this device")
        if not shutil.which("powershell.exe"):
            self.skipTest("Windows PowerShell is unavailable")

        source = self.root / "ods-roundtrip.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "ODS 往返验证"
        workbook.active["B1"] = 42
        workbook.save(source)
        workbook.close()
        original = source.read_bytes()

        self.assertIn("ods", {item["value"] for item in formats_for(source)})
        self.assertIn("xls", {item["value"] for item in formats_for(source)})
        ods = convert_file(source, "ods", self.output)
        with zipfile.ZipFile(ods) as archive:
            self.assertEqual(
                archive.read("mimetype").decode("ascii"),
                "application/vnd.oasis.opendocument.spreadsheet",
            )
            self.assertIsNone(archive.testzip())

        self.assertIn("xlsx", {item["value"] for item in formats_for(ods)})
        self.assertNotIn("ods", {item["value"] for item in formats_for(ods)})
        restored = convert_file(ods, "xlsx", self.output)
        restored_workbook = load_workbook(restored, data_only=True)
        try:
            self.assertEqual(restored_workbook.active["A1"].value, "ODS 往返验证")
            self.assertEqual(restored_workbook.active["B1"].value, 42)
        finally:
            restored_workbook.close()
        self.assertEqual(source.read_bytes(), original)

        legacy = convert_file(source, "xls", self.output)
        self.assertEqual(legacy.read_bytes()[:8], bytes.fromhex("D0CF11E0A1B11AE1"))
        restored_legacy = convert_file(legacy, "xlsx", self.output)
        legacy_workbook = load_workbook(restored_legacy, data_only=True)
        try:
            self.assertEqual(legacy_workbook.active["A1"].value, "ODS 往返验证")
            self.assertEqual(legacy_workbook.active["B1"].value, 42)
        finally:
            legacy_workbook.close()

    def test_powerpoint_exports_pdf_when_office_is_installed(self) -> None:
        if "powerpoint" not in _office_engines():
            self.skipTest("Microsoft PowerPoint is not registered on this device")
        powershell = shutil.which("powershell.exe")
        if not powershell:
            self.skipTest("Windows PowerShell is unavailable")
        source = self.root / "synthetic.pptx"
        script = self.root / "make-synthetic-presentation.ps1"
        script.write_text(
            """param([string]$Destination)
$ErrorActionPreference = 'Stop'
$application = $null
$presentation = $null
try {
    $application = New-Object -ComObject PowerPoint.Application
    $application.DisplayAlerts = 1
    $application.AutomationSecurity = 3
    $presentation = $application.Presentations.Add(0)
    $slide = $presentation.Slides.Add(1, 12)
    $shape = $slide.Shapes.AddTextbox(1, 24, 24, 600, 100)
    $shape.TextFrame.TextRange.Text = 'Synthetic presentation fixture'
    $presentation.SaveAs($Destination, 24)
} finally {
    if ($presentation) {
        $presentation.Close()
        [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($presentation)
    }
    if ($application) {
        $application.Quit()
        [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($application)
    }
}
""",
            encoding="utf-8-sig",
        )
        result = subprocess.run(
            [powershell, "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(script), "-Destination", str(source)],
            capture_output=True,
            text=True,
            timeout=30,
            creationflags=subprocess.CREATE_NO_WINDOW,
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual({item.get("engine") for item in formats_for(source)}, {"powerpoint"})
        pdf = convert_file(source, "pdf", self.output)
        from pypdf import PdfReader

        text = "\n".join(page.extract_text() or "" for page in PdfReader(str(pdf)).pages)
        self.assertIn("Synthetic presentation fixture", text)

    def test_colliding_outputs_get_new_names_without_overwriting_existing_files(self) -> None:
        source = self._table_source("csv")
        first = convert_file(source, "json", self.output)
        first_before = first.read_bytes()
        second = convert_file(source, "json", self.output)
        self.assertEqual(first.name, "fixture_converted.json")
        self.assertEqual(second.name, "fixture_converted_2.json")
        self.assertEqual(first.read_bytes(), first_before)

    def test_malformed_empty_oversize_and_unsupported_inputs_fail_without_artifacts(self) -> None:
        malformed = self.root / "bad.json"
        malformed.write_text("{broken", encoding="utf-8")
        empty = self.root / "empty.csv"
        empty.write_bytes(b"")
        invalid_encoding = self.root / "invalid.csv"
        invalid_encoding.write_bytes(b"name,value\n\xff\xfe")
        fake_image = self.root / "broken.png"
        fake_image.write_bytes(b"not a png")
        source = self._table_source("csv")
        too_large = self.root / "large.csv"
        too_large.write_bytes(b"a,b\n1,2\n")
        unsupported = self.root / "unsupported.pdf"
        unsupported.write_bytes(b"synthetic unsupported format")

        checks = [
            (malformed, "csv"),
            (empty, "json"),
            (invalid_encoding, "tsv"),
            (fake_image, "jpg"),
        ]
        for path, target in checks:
            with self.subTest(path=path.name):
                with self.assertRaises(ConversionError):
                    convert_file(path, target, self.output)
        with patch("wanxiang.converter.MAX_TABULAR_BYTES", 2):
            with self.assertRaisesRegex(ConversionError, "大小上限"):
                convert_file(too_large, "json", self.output)
        with self.assertRaisesRegex(ConversionError, "不适用于"):
            convert_file(source, "csv", self.output)
        self.assertEqual(formats_for(unsupported), [])
        self.assertFalse(any(self.output.iterdir()), "failed conversions must leave no partial output")
        self.assertEqual(
            {path.name for path in self.root.iterdir()},
            {"converted", "bad.json", "empty.csv", "invalid.csv", "broken.png", "fixture.csv", "large.csv", "unsupported.pdf"},
        )

    def test_bridge_classifies_unsupported_deduplicates_and_normalizes_output_directory(self) -> None:
        source = self._table_source("csv")
        unsupported = self.root / "opaque.bin"
        unsupported.write_bytes(b"synthetic")
        bridge = ConverterBridge()
        bridge.addFiles([QUrl.fromLocalFile(str(source)), str(source), str(unsupported)])
        self.assertEqual(len(bridge.jobs), 2)
        self.assertEqual(bridge.jobs[0]["status"], "ready")
        self.assertEqual(bridge.jobs[1]["status"], "unsupported")
        self.assertFalse(bridge.jobs[1]["canConvert"])
        bridge.setTargetFormat(bridge.jobs[0]["id"], "csv")
        self.assertNotEqual(bridge.jobs[0]["targetFormat"], "csv")
        bridge.setOutputDirectory(QUrl.fromLocalFile(str(self.output)))
        self.assertEqual(bridge.outputDirectory, str(self.output.resolve()))
        bridge.setOutputDirectory(QUrl.fromLocalFile(str(self.root / "missing")))
        self.assertEqual(bridge.outputDirectory, "")

    def test_save_uses_the_task_directory_snapshot_not_a_later_default_change(self) -> None:
        source = self._table_source("csv")
        default_directory = self.root / "default-output"
        later_directory = self.root / "later-output"
        task_directory = self.root / "task-output"
        for directory in (default_directory, later_directory, task_directory):
            directory.mkdir()

        bridge = ConverterBridge()
        bridge.setOutputDirectory(str(default_directory))
        bridge.setTaskOutputDirectory(str(task_directory))
        bridge.addFiles(str(source))
        bridge.convertReady()
        self.assertTrue(bridge._pool.waitForDone(10_000))
        self.app.processEvents()

        job = bridge.jobs[0]
        self.assertEqual(Path(job["saveDirectory"]), task_directory.resolve())
        bridge.setOutputDirectory(str(later_directory))
        bridge.setTaskOutputDirectory(str(later_directory))
        self.assertEqual(bridge.saveJob(str(job["id"])), {"ok": True, "pending": True})
        jobs = self._wait_for_saves(bridge)
        saved = Path(str(jobs[0]["outputPath"]))
        self.assertEqual(saved.parent, task_directory.resolve())
        self.assertFalse((later_directory / saved.name).exists())

    def test_video_jobs_default_to_video_output_and_keep_audio_extraction_available(self) -> None:
        source = self.root / "sample.avi"
        source.write_bytes(b"synthetic video fixture")
        bridge = ConverterBridge()
        bridge._remembered_targets = {}
        with patch("wanxiang.converter.media_targets_for", return_value=["mp3", "mkv", "mp4"]):
            bridge.addFiles(str(source))
            job = bridge.jobs[0]
            self.assertIn(job["targetFormat"], VIDEO_TARGETS)
            self.assertEqual(job["targetFormat"], "mp4")
            bridge.setTargetFormat(str(job["id"]), "mp3")
            self.assertEqual(bridge.jobs[0]["targetFormat"], "mp3")

    def test_raw_camera_file_decodes_to_png_without_modifying_the_source(self) -> None:
        class FakeRaw:
            sizes = SimpleNamespace(width=2, height=2)

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def postprocess(self, output_bps: int):
                self_output_bps.append(output_bps)
                return np.full((2, 2, 3), 128, dtype=np.uint8)

        class FakeRawPy:
            @staticmethod
            def imread(_path: str):
                return FakeRaw()

        self_output_bps: list[int] = []
        source = self.root / "camera.dng"
        source.write_bytes(b"synthetic raw payload")
        source_before = source.read_bytes()

        with patch("wanxiang.converter._rawpy_module", return_value=FakeRawPy):
            targets = {item["value"] for item in formats_for(source)}
            self.assertIn("png", targets)
            output = convert_file(source, "png", self.output)

        image = QImageReader(str(output)).read()
        self.assertFalse(image.isNull())
        self.assertEqual((image.width(), image.height()), (2, 2))
        self.assertEqual(source.read_bytes(), source_before)
        self.assertEqual(self_output_bps, [8])

    def test_image_to_ico_writes_seven_png_sizes(self) -> None:
        source = self.root / "icon-source.png"
        image = QImage(80, 40, QImage.Format.Format_ARGB32)
        image.fill(0xFF426A7D)
        self.assertTrue(image.save(str(source), "PNG"))
        source_before = source.read_bytes()

        destination = convert_file(source, "ico", self.output)
        self.assertEqual(destination.suffix, ".ico")
        payload = destination.read_bytes()
        reserved, resource_type, count = struct.unpack_from("<HHH", payload, 0)
        self.assertEqual((reserved, resource_type, count), (0, 1, 7))

        actual_sizes = []
        for index, expected_size in enumerate((16, 24, 32, 48, 64, 128, 256)):
            entry_offset = 6 + index * 16
            width, height, colors, reserved, planes, bit_count, size, offset = struct.unpack_from(
                "<BBBBHHII", payload, entry_offset
            )
            actual_width = width or 256
            actual_height = height or 256
            self.assertEqual((actual_width, actual_height), (expected_size, expected_size))
            self.assertEqual((colors, reserved, planes, bit_count), (0, 0, 0, 0))
            self.assertGreater(size, 8)
            self.assertGreaterEqual(offset, 6 + count * 16)
            self.assertLessEqual(offset + size, len(payload))
            frame = payload[offset:offset + size]
            self.assertEqual(frame[:8], b"\x89PNG\r\n\x1a\n")
            decoded = QImage.fromData(frame, "PNG")
            self.assertEqual((decoded.width(), decoded.height()), (expected_size, expected_size))
            self.assertEqual(decoded.pixelColor(0, 0).alpha(), 0)
            self.assertGreater(decoded.pixelColor(expected_size // 2, expected_size // 2).alpha(), 0)
            actual_sizes.append((actual_width, actual_height))

        self.assertEqual(actual_sizes, [(size, size) for size in (16, 24, 32, 48, 64, 128, 256)])
        reader = QImageReader(str(destination))
        self.assertTrue(reader.canRead(), reader.errorString())
        self.assertEqual(source.read_bytes(), source_before)

    def test_animated_gif_frames_are_preserved_when_exporting_to_webp(self) -> None:
        from PIL import Image

        source = self.root / "animated-source.gif"
        frames = []
        for color in ((220, 30, 40, 255), (20, 190, 80, 255), (25, 70, 220, 255)):
            frame = Image.new("RGBA", (32, 24), color)
            frames.append(frame)
        frames[0].save(
            source,
            format="GIF",
            save_all=True,
            append_images=frames[1:],
            duration=[80, 130, 240],
            loop=3,
            disposal=2,
        )
        for frame in frames:
            frame.close()
        source_before = source.read_bytes()

        targets = {item["value"] for item in formats_for(source)}
        self.assertIn("webp", targets)
        destination = convert_file(source, "webp", self.output)
        with Image.open(destination) as converted:
            self.assertEqual(converted.n_frames, 3)
            self.assertEqual(converted.info.get("loop"), 3)
        reader = QImageReader(str(destination))
        delays = []
        for _index in range(reader.imageCount()):
            self.assertFalse(reader.read().isNull())
            delays.append(reader.nextImageDelay())
        self.assertEqual(delays, [80, 130, 240])
        del reader
        self.assertEqual(source.read_bytes(), source_before)

        still = self.root / "still.png"
        still_image = Image.new("RGB", (20, 12), (120, 90, 60))
        still_image.save(still)
        still_image.close()
        self.assertIn("gif-animation", {item["value"] for item in formats_for(still)})
        single_frame_gif = convert_file(still, "gif-animation", self.output)
        with Image.open(single_frame_gif) as converted:
            self.assertEqual(converted.n_frames, 1)

    def test_image_ocr_target_tracks_local_engine_and_writes_text(self) -> None:
        source = self.root / "ocr-source.png"
        image = QImage(40, 24, QImage.Format.Format_RGB32)
        image.fill(Qt.GlobalColor.white)
        self.assertTrue(image.save(str(source), "PNG"))
        source_before = source.read_bytes()
        engine = TesseractEngine("tesseract.exe", frozenset({"chi_sim", "eng"}), "eng")
        recognized_lines = [
            {"text": "FLUKE workbench", "confidence": 97.0},
            {"text": "local conversion", "confidence": 94.0},
        ]

        with patch("wanxiang.converter.searchable_pdf_ocr_available", return_value=False):
            self.assertNotIn("image-ocr-txt", {item["value"] for item in formats_for(source)})

        with (
            patch("wanxiang.converter.searchable_pdf_ocr_available", return_value=True),
            patch("wanxiang.converter.tesseract_engine", return_value=engine),
            patch("wanxiang.converter.recognize_tesseract_tsv", return_value=recognized_lines) as recognize,
        ):
            targets = {item["value"] for item in formats_for(source)}
            self.assertIn("image-ocr-txt", targets)
            progress: list[int] = []
            output = convert_file(source, "image-ocr-txt", self.output, ocr_language="chi_sim", progress_callback=progress.append)

        self.assertEqual(output.suffix, ".txt")
        self.assertEqual(output.read_text(encoding="utf-8"), "FLUKE workbench\nlocal conversion\n")
        self.assertEqual(recognize.call_args.args[0].language, "chi_sim")
        self.assertEqual(progress, sorted(progress))
        self.assertEqual(progress[-1], 100)
        self.assertEqual(source.read_bytes(), source_before)


class ConverterQmlIntegrationTests(unittest.TestCase):
    """Run the real converter controls with only generated local files and a temp DB."""

    @classmethod
    def setUpClass(cls) -> None:
        QQuickStyle.setStyle("Basic")
        cls.app = QApplication.instance() or QApplication([])
        cls.app.setQuitOnLastWindowClosed(False)
        cls.native_dir = Path(__file__).resolve().parents[1]

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory(prefix="wanxiang-converter-qml-")
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        self.database_path = self.root / "synthetic.sqlite3"
        self.engine = QQmlApplicationEngine()
        self.addCleanup(self._destroy_engine)

        migration = MigrationBridge(self.database_path)
        snapshot = migration.data
        finance = FinanceBridge(self.database_path, snapshot)
        self.converter = ConverterBridge()
        self.converter_engine_controller = ConverterEngineUpdateBridge(self.root / "engine-updates")
        controllers = {
            "migrationController": migration,
            "weatherController": WeatherBridge(self.database_path, snapshot),
            "newsController": NewsBridge(self.database_path, snapshot),
            "habitController": HabitBridge(self.database_path, snapshot),
            "preferencesController": IssuePreferencesBridge(self.database_path, snapshot),
            "readingController": ReadingBridge(self.database_path, snapshot),
            "dailyController": DailyBridge(self.database_path, snapshot),
            "financeController": finance,
            "backupController": BackupBridge(self.database_path, finance),
            "fitnessController": FitnessBridge(self.database_path, snapshot),
            "plannerController": PlannerBridge(self.database_path, snapshot),
            "shoppingController": ShoppingBridge(self.database_path, snapshot),
            "mediaController": MediaBridge(self.database_path, snapshot),
            "archiveController": ArchiveBridge(self.database_path, snapshot),
            "converterController": self.converter,
            "converterEngineController": self.converter_engine_controller,
            "brandController": BrandBridge(BrandRepository(self.database_path), self.engine),
        }
        self.engine.setInitialProperties(controllers)
        self.engine.load(QUrl.fromLocalFile(str(self.native_dir / "qml" / "Main.qml")))
        self.assertTrue(self.engine.rootObjects(), "production Main.qml failed to load")
        self.window = self.engine.rootObjects()[0]
        self.window.setProperty("currentSectionIndex", 8)
        self.window.resize(1480, 960)
        self.window.show()
        QTest.qWait(120)
        self.page = self.window.findChild(QObject, "converterPage")
        self.assertIsNotNone(self.page, "production converter page was not loaded")

    def _destroy_engine(self) -> None:
        if getattr(self, "engine", None) is not None:
            engine = self.engine
            self.engine = None
            engine.deleteLater()
            if QApplication.instance() is not None:
                QCoreApplication.sendPostedEvents(engine, QEvent.Type.DeferredDelete)

    def _click(self, item: QObject) -> None:
        self.assertTrue(QMetaObject.invokeMethod(item, "click", Qt.ConnectionType.DirectConnection))
        QTest.qWait(40)

    def _find_quick_item(self, object_name: str):
        pending = [self.window.contentItem()]
        while pending:
            item = pending.pop()
            if item.objectName() == object_name:
                return item
            pending.extend(item.childItems())
        return None

    def test_converter_engine_dialog_compact_layout_and_scroll(self) -> None:
        self.window.resize(760, 620)
        engine_rows = []
        for index in range(8):
            engine_rows.append({
                "id": f"engine-{index}",
                "label": f"转换引擎 {index + 1}：处理高分辨率图像、批量文档和长音频的完整工具名称",
                "description": "用于处理多种本地文件格式的转换组件。这段较长说明用于检查在较窄窗口里中文换行后，操作按钮是否仍留在同一条目内，状态和版本信息是否会互相挤压。",
                "installedVersion": "" if index == 0 else "1.2.3",
                "installedSource": "未内置" if index == 0 else "已更新",
                "availableVersion": "1.3.0",
                "updateAvailable": True,
                "downloadSize": 72 * 1024 * 1024,
                "canRollback": True,
                "previousVersion": "1.2.2",
                "licenseNoticeAvailable": True,
            })
        self.converter_engine_controller._snapshot = lambda: {
            "engines": engine_rows,
            "busy": False,
            "progress": 0,
            "statusMessage": "已检查更新。此处放置较长的状态说明，用于观察窄窗口中的文字换行和列表可用高度。",
            "releaseTag": "fluke-engine-bundle-2026.09-preview-with-a-long-release-channel-name",
        }
        self.converter_engine_controller._begin = lambda *_args, **_kwargs: None
        self.converter_engine_controller.stateChanged.emit()

        manage_button = self.window.findChild(QObject, "converterManageEnginesButton")
        self.assertIsNotNone(manage_button)
        self._click(manage_button)
        QTest.qWait(120)

        dialog = self.window.findChild(QObject, "converterEnginesDialog")
        engine_list = self.window.findChild(QObject, "converterEngineList")
        close_button = self.window.findChild(QObject, "converterEngineCloseButton")
        first_install_button = self._find_quick_item("converterEngineInstallButton_engine-0")
        first_rollback_button = self._find_quick_item("converterEngineRollbackButton_engine-0")
        self.assertIsNotNone(dialog)
        self.assertIsNotNone(engine_list)
        self.assertIsNotNone(close_button)
        self.assertIsNotNone(first_install_button)
        self.assertIsNotNone(first_rollback_button)
        self.assertTrue(dialog.property("visible"))
        self.assertTrue(self.window.property("bundledSansFontReady"))
        self.assertTrue(self.window.property("bundledSerifFontReady"))
        self.assertGreater(dialog.property("width"), 0)
        self.assertLessEqual(dialog.property("x") + dialog.property("width"), self.window.width())
        self.assertGreater(engine_list.property("height"), 100)
        self.assertEqual(engine_list.property("count"), len(engine_rows))
        self.assertGreater(engine_list.property("contentHeight"), engine_list.property("height"))
        self.assertTrue(close_button.property("visible"))
        self.assertEqual(first_install_button.property("text"), "安装")
        self.assertEqual(first_rollback_button.property("text"), "回滚到 1.2.2")

        engine_list.setProperty("contentY", engine_list.property("contentHeight"))
        QTest.qWait(60)
        self.assertGreater(engine_list.property("contentY"), 0)
        self.assertGreater(engine_list.property("contentY"), 0)

        engine_list.setProperty("contentY", 0)
        QTest.qWait(60)

        screenshot = self.window.grabWindow()
        self.assertFalse(screenshot.isNull(), "compact converter engine dialog screenshot was empty")
        screenshot_dir = self.native_dir / "qa-release" / "ui-runtime-review"
        screenshot_dir.mkdir(parents=True, exist_ok=True)
        self.assertTrue(screenshot.save(str(screenshot_dir / "converter-engines-dialog-compact-760x620-2026-09-29.png"), "PNG"))

    def test_output_format_selector_has_an_accessible_field_name(self) -> None:
        source = self.root / "accessible-format-fixture.csv"
        source.write_text("name,value\nsynthetic,1\n", encoding="utf-8")
        self.converter.addFiles(QUrl.fromLocalFile(str(source)))
        QTest.qWait(80)
        selector = self._quick_item("converterTargetFormat")
        self.assertIsNotNone(selector, "the queued job output format selector was not found")
        interface = QAccessible.queryAccessibleInterface(selector)
        self.assertIsNotNone(interface)
        self.assertEqual(str(interface.text(QAccessible.Text.Name)), "输出格式")

    def test_production_page_applies_a_common_target_to_the_queue(self) -> None:
        sources = []
        for name in ("batch-first.csv", "batch-second.csv"):
            source = self.root / name
            source.write_text("name,value\nsynthetic,1\n", encoding="utf-8")
            sources.append(source)
        self.converter.addFiles([QUrl.fromLocalFile(str(path)) for path in sources])
        QTest.qWait(80)

        row = self._quick_item("converterCommonTargetRow")
        selector = self._quick_item("converterCommonTargetFormat")
        apply_button = self._quick_item("converterApplyCommonTargetButton")
        remember = self._quick_item("converterRememberTargetFormats")
        self.assertIsNotNone(row)
        self.assertTrue(row.isVisible())
        self.assertIsNotNone(selector)
        self.assertIsNotNone(remember)
        self.assertFalse(remember.property("checked"))
        interface = QAccessible.queryAccessibleInterface(selector)
        self.assertIsNotNone(interface)
        self.assertEqual(str(interface.text(QAccessible.Text.Name)), "共同目标格式")
        self.assertIsNotNone(apply_button)
        self.assertTrue(apply_button.property("enabled"))

        options = self.converter.commonTargetFormats
        html_index = next(index for index, option in enumerate(options) if option["value"] == "html")
        self.assertTrue(selector.setProperty("currentIndex", html_index))
        QTest.qWait(40)
        self._click(apply_button)

        self.assertEqual([job["targetFormat"] for job in self.converter.jobs], ["html", "html"])
        message = str(self._quick_item("converterCommonTargetMessage").property("text"))
        self.assertIn("2", message)
        self.assertIn("未改变以后文件的默认格式", message)

        self._click(remember)
        self.assertTrue(remember.property("checked"))
        self.assertTrue(self.converter.rememberTargetFormats)
        self._click(apply_button)
        message = str(self._quick_item("converterCommonTargetMessage").property("text"))
        self.assertIn("已记为以后同类文件默认", message)

    def _quick_item(self, object_name: str) -> QObject | None:
        pending = [self.window.contentItem()]
        while pending:
            item = pending.pop()
            if item.objectName() == object_name:
                return item
            pending.extend(item.childItems())
        return None

    def _wait_for_statuses(self, expected: set[str], timeout_seconds: float = 8.0) -> list[dict[str, object]]:
        deadline = time.monotonic() + timeout_seconds
        while time.monotonic() < deadline:
            jobs = self.converter.jobs
            if jobs and all(job["status"] in expected for job in jobs):
                return jobs
            QTest.qWait(40)
        self.fail(f"conversion jobs did not reach {sorted(expected)}: {self.converter.jobs!r}")

    def _wait_for_saves(self, timeout_seconds: float = 8.0) -> list[dict[str, object]]:
        deadline = time.monotonic() + timeout_seconds
        while time.monotonic() < deadline:
            self.app.processEvents()
            jobs = self.converter.jobs
            if not any(job.get("savePending") for job in jobs):
                return jobs
            QTest.qWait(20)
        self.fail(f"save jobs did not finish: {self.converter.jobs!r}")

    def test_empty_queue_offers_a_visible_file_picker_action(self) -> None:
        drop_zone = self._quick_item("converterEmptyDropZone")
        add_files = self._quick_item("converterEmptyAddFilesButton")
        dialog = self.window.findChild(QObject, "converterInputFileDialog")
        self.assertIsNotNone(drop_zone)
        self.assertTrue(drop_zone.isVisible())
        self.assertIsNotNone(add_files)
        self.assertIsNotNone(dialog)

        self._click(add_files)
        self.assertTrue(dialog.property("visible"))
        dialog.close()

        fixture = self.root / "empty-state-action.csv"
        fixture.write_text("name,value\nsynthetic,1\n", encoding="utf-8")
        self.converter.addFiles(QUrl.fromLocalFile(str(fixture)))
        QTest.qWait(80)
        self.assertFalse(drop_zone.isVisible())

    def test_pdf_split_control_drives_grouped_conversion_and_hides_when_done(self) -> None:
        from pypdf import PdfWriter

        source = self.root / "ui-grouped.pdf"
        writer = PdfWriter()
        for _ in range(3):
            writer.add_blank_page(width=72, height=72)
        with source.open("wb") as stream:
            writer.write(stream)

        self.converter.addFiles(QUrl.fromLocalFile(str(source)))
        job_id = str(self.converter.jobs[0]["id"])
        self.converter.setTargetFormat(job_id, "zip")
        QTest.qWait(80)
        size_row = self._quick_item("converterPdfSplitGroupSizeRow")
        size_control = self._quick_item("converterPdfSplitGroupSize")
        self.assertIsNotNone(size_row)
        self.assertIsNotNone(size_control)
        self.assertTrue(size_row.isVisible())
        self.assertEqual(size_control.property("value"), 1)

        size_control.forceActiveFocus()
        QTest.keyClick(self.window, Qt.Key.Key_Up)
        QTest.qWait(80)
        self.assertEqual(size_control.property("value"), 2)
        self.assertEqual(self.converter.pdfSplitGroupSizeForJob(job_id), 2)

        self.converter.convertReady()
        jobs = self._wait_for_statuses({"done"})
        self._click(self.window.findChild(QObject, "converterSaveAllButton"))
        jobs = self._wait_for_saves()
        with zipfile.ZipFile(jobs[0]["outputPath"]) as archive:
            self.assertEqual(archive.namelist(), ["pages_0001-0002.pdf", "pages_0003-0003.pdf"])
        QTest.qWait(80)
        completed_row = self._quick_item("converterPdfSplitGroupSizeRow")
        self.assertTrue(completed_row is None or not completed_row.isVisible())

    def test_format_limits_are_grouped_and_expand_on_demand(self) -> None:
        source = self.root / "synthetic.csv"
        source.write_text("title,count\nsynthetic,2\n", encoding="utf-8")
        self.converter.addFiles(QUrl.fromLocalFile(str(source)))
        QTest.qWait(80)

        merge_button = self.window.findChild(QObject, "converterMergePdfButton")
        self.assertIsNotNone(merge_button)
        self.assertFalse(merge_button.property("visible"))

        disclosure = self.window.findChild(QObject, "converterFormatLimitsToggle")
        details = self.window.findChild(QObject, "converterFormatLimitsGrid")
        sections = self.window.findChild(QObject, "converterFormatLimitSections")
        self.assertIsNotNone(disclosure)
        self.assertIsNotNone(details)
        self.assertIsNotNone(sections)
        self.assertFalse(details.property("visible"))
        self.assertEqual(sections.property("count"), 6)

        locale = LocalizationBridge(LocalizationRepository(self.database_path), self.engine, self.app)
        self.addCleanup(locale.close)
        self.window.setProperty("localeController", locale)
        self.assertTrue(locale.setLanguage("en_US")["ok"])
        intro = self.window.findChild(QObject, "converterIntroDescription")
        self.assertIsNotNone(intro)
        self.assertEqual(
            intro.property("text"),
            "Supports common image, document, spreadsheet, e-book, and PDF tasks. "
            "The installer includes FFmpeg for audio/video, Tesseract OCR with Simplified Chinese and English models, "
            "and Calibre for e-book conversion. Each engine can be updated separately from Conversion engines.",
        )
        self.assertEqual(disclosure.property("text"), "View format notes")

        self._click(disclosure)
        self.assertTrue(details.property("visible"))
        self._click(disclosure)
        self.assertFalse(details.property("visible"))

    def test_production_file_picker_filters_accepts_fixture_and_cancel_preserves_queue(self) -> None:
        source = self.root / "picker-fixture.csv"
        source.write_text("title,count\nsynthetic,3\n", encoding="utf-8")

        dialog = self.window.findChild(QObject, "converterInputFileDialog")
        self.assertIsNotNone(dialog, "production converter OpenFiles dialog was not found")
        filters = [str(value) for value in dialog.property("nameFilters")]
        self.assertEqual(len(filters), 2)
        self.assertEqual(
            filters[0],
            f"支持的文件 ({self.converter.inputFilePatterns})",
        )
        self.assertEqual(filters[1], "所有文件 (*)")

        # selectedFiles is read-only in Qt; seed one selection through the
        # writable selectedFile property, then exercise the OpenFiles handler.
        dialog.setProperty("selectedFile", QUrl.fromLocalFile(str(source)))
        self.assertEqual(
            [Path(url.toLocalFile()).resolve() for url in dialog.property("selectedFiles")],
            [source.resolve()],
        )
        accepted = QMetaObject.invokeMethod(dialog, "accepted", Qt.ConnectionType.DirectConnection)
        self.assertTrue(accepted, "production converter FileDialog accepted handler was not invokable")
        self.assertEqual([job["sourcePath"] for job in self.converter.jobs], [str(source.resolve())])

        queued = self.converter.jobs
        rejected = QMetaObject.invokeMethod(dialog, "rejected", Qt.ConnectionType.DirectConnection)
        self.assertTrue(rejected, "production converter FileDialog rejected signal was not invokable")
        self.assertEqual(self.converter.jobs, queued)

    def test_production_folder_picker_opens_in_temp_directory_and_cancel_preserves_choice(self) -> None:
        selected = self.root / "selected-output"
        selected.mkdir()

        dialog = self.window.findChild(QObject, "converterOutputFolderDialog")
        self.assertIsNotNone(dialog, "production converter FolderDialog was not found")
        dialog.setProperty("currentFolder", QUrl.fromLocalFile(str(self.root)))
        current_folder = dialog.property("currentFolder")
        if isinstance(current_folder, QUrl):
            current_folder = current_folder.toLocalFile()
        self.assertEqual(Path(str(current_folder)).resolve(), self.root.resolve())

        self.converter.setOutputDirectory(QUrl.fromLocalFile(str(selected)))
        opened = QMetaObject.invokeMethod(dialog, "open", Qt.ConnectionType.DirectConnection)
        self.assertTrue(opened, "production converter FolderDialog could not be opened")
        self.assertTrue(dialog.property("visible"))
        rejected = QMetaObject.invokeMethod(dialog, "rejected", Qt.ConnectionType.DirectConnection)
        self.assertTrue(rejected, "production converter FolderDialog rejected signal was not invokable")
        self.assertEqual(self.converter.outputDirectory, str(selected.resolve()))

    def test_production_page_separates_default_and_current_task_output_locations(self) -> None:
        default_directory = self.root / "default-output"
        task_directory = self.root / "task-output"
        default_directory.mkdir()
        task_directory.mkdir()

        default_button = self._quick_item("converterChooseDefaultOutputButton")
        task_button = self._quick_item("converterChooseTaskOutputButton")
        clear_button = self._quick_item("converterClearTaskOutputButton")
        default_summary = self._quick_item("converterDefaultOutputSummary")
        task_summary = self._quick_item("converterTaskOutputSummary")
        self.assertIsNotNone(default_button)
        self.assertIsNotNone(task_button)
        self.assertIsNotNone(clear_button)
        self.assertIsNotNone(default_summary)
        self.assertIsNotNone(task_summary)

        self.converter.setOutputDirectory(QUrl.fromLocalFile(str(default_directory)))
        self.converter.setTaskOutputDirectory(QUrl.fromLocalFile(str(task_directory)))
        self.assertEqual(self.converter.outputDirectory, str(default_directory.resolve()))
        self.assertEqual(self.converter.taskOutputDirectory, str(task_directory.resolve()))
        self.assertTrue(clear_button.property("visible"))
        self.assertEqual(str(clear_button.property("text")), "跟随默认位置")
        self.assertIn("默认位置", str(default_summary.property("text")))
        self.assertIn(str(task_directory.resolve()), str(task_summary.property("text")))

        self._click(clear_button)
        self.assertFalse(self.converter.taskOutputDirectory)
        self.assertFalse(clear_button.property("visible"))

    def test_production_image_folder_picker_starts_a_folder_pdf_job(self) -> None:
        folder = self.root / "picked-images"
        folder.mkdir()
        for index in range(2):
            image = QImage(24, 32, QImage.Format.Format_ARGB32)
            image.fill(0xFF426A7D + index)
            self.assertTrue(image.save(str(folder / f"page-{index + 1}.png"), "PNG"))

        button = self._quick_item("converterImageFolderPdfButton")
        dialog = self.window.findChild(QObject, "converterImageFolderDialog")
        self.assertIsNotNone(button)
        self.assertIsNotNone(dialog)
        self.assertEqual(str(button.property("text")), "图片文件夹转 PDF")
        self._click(button)
        self.assertTrue(dialog.property("visible"))
        dialog.setProperty("selectedFolder", QUrl.fromLocalFile(str(folder)))
        accepted = QMetaObject.invokeMethod(dialog, "accepted", Qt.ConnectionType.DirectConnection)
        self.assertTrue(accepted)
        message = self._quick_item("converterImageFolderMessage")
        self.assertIn("已添加 2 张图片", str(message.property("text")))

        deadline = time.monotonic() + 8
        while time.monotonic() < deadline:
            self.app.processEvents()
            if self.converter.jobs and self.converter.jobs[0]["status"] in ("done", "failed"):
                break
            QTest.qWait(20)
        self.assertEqual(self.converter.jobs[0]["status"], "done")

    def test_production_page_batches_files_converts_and_clears_completed_jobs(self) -> None:
        source = self.root / "synthetic.csv"
        source.write_text("title,count\n合成项目,2\n", encoding="utf-8")
        source_before = source.read_bytes()
        unsupported = self.root / "synthetic.bin"
        unsupported.write_bytes(b"synthetic; not a supported document")
        output_dir = self.root / "chosen-output"
        output_dir.mkdir()

        self.converter.addFiles([QUrl.fromLocalFile(str(source)), QUrl.fromLocalFile(str(unsupported))])
        jobs = self.converter.jobs
        self.assertEqual([job["status"] for job in jobs], ["ready", "unsupported"])
        self.assertEqual(jobs[0]["targetFormat"], "json")
        self.assertEqual(self.page.property("visible"), True)
        self.converter.setOutputDirectory(QUrl.fromLocalFile(str(output_dir)))

        unsupported_help = self._quick_item("converterUnsupportedHelp_" + str(jobs[1]["id"]))
        self.assertIsNotNone(unsupported_help)
        self.assertTrue(unsupported_help.isVisible())
        self.assertEqual(unsupported_help.property("text"), "检查转换组件")
        self._click(unsupported_help)
        engines_dialog = self.window.findChild(QObject, "converterEnginesDialog")
        request_context = self.window.findChild(QObject, "converterEngineRequestContext")
        request_summary = self.window.findChild(QObject, "converterEngineRequestContextSummary")
        self.assertIsNotNone(engines_dialog)
        self.assertIsNotNone(request_context)
        self.assertIsNotNone(request_summary)
        self.assertTrue(engines_dialog.property("visible"))
        self.assertTrue(request_context.property("visible"))
        self.assertIn("synthetic.bin", str(request_summary.property("text")))
        engines_dialog.close()

        start = self.window.findChild(QObject, "converterStartButton")
        self.assertIsNotNone(start)
        self.assertTrue(start.property("enabled"))
        self._click(start)
        jobs = self._wait_for_statuses({"done", "unsupported"})
        self.assertEqual([job["status"] for job in jobs], ["done", "unsupported"])
        self._click(self.window.findChild(QObject, "converterSaveAllButton"))
        jobs = self.converter.jobs
        converted = Path(str(jobs[0]["outputPath"]))
        self.assertEqual(converted.parent, output_dir.resolve())
        self.assertEqual(json.loads(converted.read_text(encoding="utf-8")), [{"title": "合成项目", "count": "2"}])
        self.assertEqual(source.read_bytes(), source_before)
        self.assertTrue(jobs[1]["error"])

        clear = self.window.findChild(QObject, "converterClearFinishedButton")
        self.assertIsNotNone(clear)
        self.assertTrue(clear.property("enabled"))
        self._click(clear)
        self.assertEqual(self.converter.jobs, [])

    def test_image_ocr_target_exposes_language_choice_on_the_production_page(self) -> None:
        source = self.root / "ocr-preview.png"
        image = QImage(48, 32, QImage.Format.Format_RGB32)
        image.fill(Qt.GlobalColor.white)
        self.assertTrue(image.save(str(source), "PNG"))
        engine = TesseractEngine("tesseract.exe", frozenset({"chi_sim", "eng"}), "chi_sim+eng")
        languages = [
            {"label": "简体中文 + English", "value": "chi_sim+eng"},
            {"label": "eng", "value": "eng"},
            {"label": "chi_sim", "value": "chi_sim"},
        ]

        with (
            patch("wanxiang.converter.searchable_pdf_ocr_available", return_value=True),
            patch("wanxiang.converter.tesseract_engine", return_value=engine),
            patch("wanxiang.converter.available_ocr_languages", return_value=languages),
        ):
            self.converter.addFiles(QUrl.fromLocalFile(str(source)))
            job_id = str(self.converter.jobs[0]["id"])
            options = self.converter.formatsForPath(str(source))
            self.assertIn("image-ocr-txt", {item["value"] for item in options})
            self.converter.setTargetFormat(job_id, "image-ocr-txt")
            self.converter.setJobOcrLanguage(job_id, "eng")

            language_selector = self._quick_item("converterOcrLanguage")
            self.assertIsNotNone(language_selector)
            self.assertTrue(language_selector.isVisible())
            self.assertEqual(self.converter.ocrLanguageForJob(job_id), "eng")

    def test_production_page_previews_then_saves_a_text_conversion(self) -> None:
        source = self.root / "preview-fixture.csv"
        source.write_text("name,value\n预览项目,7\n", encoding="utf-8")
        source_before = source.read_bytes()
        output_dir = self.root / "preview-output"
        output_dir.mkdir()

        self.converter.addFiles(QUrl.fromLocalFile(str(source)))
        job_id = str(self.converter.jobs[0]["id"])
        self.converter.setOutputDirectory(QUrl.fromLocalFile(str(output_dir)))
        self._click(self.window.findChild(QObject, "converterStartButton"))
        jobs = self._wait_for_statuses({"done"})
        self.assertEqual(jobs[0]["progress"], 100)
        self.assertTrue(Path(str(jobs[0]["stagedPath"])).is_file())
        planned = self._quick_item(f"converterPlannedSaveDirectory_{job_id}")
        status = self._quick_item(f"converterJobStatus_{job_id}")
        save_summary = self._quick_item("converterSaveAllSummary")
        self.assertIsNotNone(planned)
        self.assertIsNotNone(status)
        self.assertIsNotNone(save_summary)
        self.assertTrue(planned.isVisible())
        self.assertIn(str(output_dir.resolve()), str(planned.property("text")))
        self.assertEqual(status.property("text"), "待保存")
        self.assertIn(str(output_dir.resolve()), str(save_summary.property("text")))

        preview_button = self._quick_item(f"converterPreviewButton_{job_id}")
        self.assertIsNotNone(preview_button)
        self.assertTrue(preview_button.isVisible())
        self._click(preview_button)
        dialog = self.window.findChild(QObject, "converterResultPreviewDialog")
        preview_text = self.window.findChild(QObject, "converterResultPreviewText")
        self.assertIsNotNone(dialog)
        self.assertIsNotNone(preview_text)
        self.assertTrue(dialog.property("visible"))
        self.assertIn("预览项目", str(preview_text.property("text")))
        self.assertIn('"value": "7"', str(preview_text.property("text")))
        self.assertTrue(QMetaObject.invokeMethod(dialog, "close", Qt.ConnectionType.DirectConnection))

        self._click(self.window.findChild(QObject, "converterSaveAllButton"))
        jobs = self._wait_for_saves()
        saved = Path(str(jobs[0]["outputPath"]))
        self.assertEqual(saved.parent, output_dir.resolve())
        self.assertEqual(json.loads(saved.read_text(encoding="utf-8")), [{"name": "预览项目", "value": "7"}])
        self.assertEqual(source.read_bytes(), source_before)

    def test_production_page_reports_corrupt_input_and_retry_resets_job(self) -> None:
        source = self.root / "broken.json"
        source.write_text("not valid json", encoding="utf-8")
        self.converter.addFiles(QUrl.fromLocalFile(str(source)))
        start = self.window.findChild(QObject, "converterStartButton")
        self.assertIsNotNone(start)
        self._click(start)
        jobs = self._wait_for_statuses({"failed"})
        self.assertEqual(jobs[0]["status"], "failed")
        self.assertIn("无法读取 JSON", jobs[0]["error"])
        self.assertEqual(jobs[0]["outputPath"], "")
        self.assertFalse(any(path.name.endswith(".tmp") or "_converted" in path.name for path in self.root.iterdir()))

        self.converter.retryJob(str(jobs[0]["id"]))
        retried = self.converter.jobs[0]
        self.assertEqual(retried["status"], "ready")
        self.assertEqual(retried["error"], "")
        self.assertEqual(retried["progress"], 0)
        self.assertTrue(start.property("enabled"))

    def test_production_page_shows_password_entry_for_pdf_security_targets(self) -> None:
        from pypdf import PdfWriter

        source = self.root / "security.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        with source.open("wb") as stream:
            writer.write(stream)
        self.converter.addFiles(QUrl.fromLocalFile(str(source)))
        job = self.converter.jobs[0]
        self.converter.setTargetFormat(str(job["id"]), "pdf-encrypt")
        QTest.qWait(80)

        repeater = self.page.findChild(QObject, "converterJobsRepeater")
        password_group = self._quick_item("converterPdfPasswordGroup")
        password_label = self._quick_item("converterPdfPasswordLabel")
        password_field = self._quick_item("converterPdfPasswordField")
        self.assertIsNotNone(repeater)
        self.assertEqual(repeater.property("count"), 1)
        self.assertIsNotNone(password_group)
        self.assertTrue(password_group.isVisible())
        self.assertIsNotNone(password_label)
        self.assertEqual(password_label.property("text"), "PDF 加密密码")
        self.assertIsNotNone(password_field)
        password_accessible = QAccessible.queryAccessibleInterface(password_field)
        self.assertIsNotNone(password_accessible)
        self.assertEqual(str(password_accessible.text(QAccessible.Text.Name)), "PDF 加密密码")
        self.assertEqual(password_field.property("placeholderText"), "至少 8 个字符")
        self.converter.setJobPassword(str(job["id"]), "ui-only-password")
        self.assertNotIn("ui-only-password", repr(self.converter.jobs))
        self.assertNotIn("password", self.converter.jobs[0])

    def test_production_page_merges_queued_pdfs(self) -> None:
        from pypdf import PdfReader, PdfWriter

        sources = [self.root / "first.pdf", self.root / "second.pdf"]
        originals = []
        for source, width in zip(sources, (72, 144)):
            writer = PdfWriter()
            writer.add_blank_page(width=width, height=width)
            with source.open("wb") as stream:
                writer.write(stream)
            originals.append(source.read_bytes())
        invalid = self.root / "invalid.pdf"
        invalid.write_bytes(b"not a PDF")
        self.converter.addFiles([QUrl.fromLocalFile(str(source)) for source in (*sources, invalid)])

        merge_button = self.window.findChild(QObject, "converterMergePdfButton")
        self.assertIsNotNone(merge_button)
        self.assertTrue(merge_button.property("enabled"))
        self._click(merge_button)
        dialog = self.window.findChild(QObject, "converterPdfMergeDialog")
        model = self.window.findChild(QObject, "converterMergeOrderModel")
        self.assertIsNotNone(dialog)
        self.assertIsNotNone(model)
        self.assertTrue(dialog.property("visible"))
        self.assertEqual(model.property("count"), 2)
        excluded_toggle = self._quick_item("converterMergeExcludedToggle")
        excluded_list = self._quick_item("converterMergeExcludedList")
        self.assertIsNotNone(excluded_toggle)
        self.assertIsNotNone(excluded_list)
        self.assertFalse(excluded_list.property("visible"))
        self._click(excluded_toggle)
        self.assertTrue(excluded_list.property("visible"))
        self.assertEqual(excluded_list.property("count"), 1)

        locale = LocalizationBridge(LocalizationRepository(self.database_path), self.engine, self.app)
        self.addCleanup(locale.close)
        self.window.setProperty("localeController", locale)
        self.assertTrue(locale.setLanguage("en_US")["ok"])
        self.assertEqual(dialog.property("title"), "Merge PDFs")
        confirm = self.window.findChild(QObject, "converterMergeConfirmButton")
        self.assertIsNotNone(confirm)
        self.assertEqual(confirm.property("text"), "Merge 2 PDFs")

        move_up = self._quick_item("converterMergeMoveUp_1")
        self.assertIsNotNone(move_up)
        self._click(move_up)
        self.assertTrue(confirm.property("enabled"))
        self._click(confirm)

        deadline = time.monotonic() + 8.0
        while time.monotonic() < deadline:
            jobs = self.converter.jobs
            if len(jobs) == 4 and jobs[3]["status"] in ("done", "failed"):
                break
            QTest.qWait(40)
        self.assertEqual(len(self.converter.jobs), 4)
        merge_job = self.converter.jobs[3]
        self.assertEqual(merge_job["status"], "done", merge_job["error"])
        self._click(self.window.findChild(QObject, "converterSaveAllButton"))
        self._wait_for_saves()
        merge_job = self.converter.jobs[3]
        merged = PdfReader(merge_job["outputPath"])
        self.assertEqual(len(merged.pages), 2)
        self.assertEqual([float(page.mediabox.width) for page in merged.pages], [144, 72])
        self.assertEqual([source.read_bytes() for source in sources], originals)

    def test_audio_and_video_previews_classify_completed_outputs(self) -> None:
        for job_id, extension, expected_kind in (
            ("audio-preview", "mp3", "audio"),
            ("video-preview", "mp4", "video"),
        ):
            output = self.root / f"preview.{extension}"
            output.write_bytes(b"synthetic preview fixture")
            self.converter._jobs.append({
                "id": job_id,
                "status": "done",
                "outputPath": str(output),
            })
            preview = self.converter.previewInfo(job_id)
            self.assertEqual(preview["kind"], expected_kind)
            self.assertEqual(Path(QUrl(preview["url"]).toLocalFile()).resolve(), output.resolve())

if __name__ == "__main__":
    unittest.main()
